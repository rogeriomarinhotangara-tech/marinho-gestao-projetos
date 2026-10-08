"""Exportação dos painéis da Parceria de Tricologia: relatório PDF (HTML de impressão -> Chromium) e Excel formatado.
Uso: python3 -I build_export.py <scratch>   (lê out/dash2.json e out/out26.json; grava out/exports/)"""
import sys, json, os, re, html, subprocess
SP = sys.argv[1]; sys.path.insert(0, SP + '/scripts')
D = json.load(open(SP + '/out/dash2.json', encoding='utf-8'))
O = json.load(open(SP + '/out/out26.json', encoding='utf-8'))
OUTD = SP + '/out/exports'; os.makedirs(OUTD, exist_ok=True)
NAME = 'Tricologia_Paineis_Out2026'
DATA_REF = '08/10/2026'

# ------------------------------------------------------------------ helpers
def nfmt(v, d=2):
    s = f'{abs(v):,.{d}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return ('−' if v < -0.0000001 else '') + s
def brl(v): return ('−' if v < -0.004 else '') + 'R$ ' + nfmt(abs(v))
def brl0(v): return ('−' if v < -0.5 else '') + 'R$ ' + nfmt(abs(v), 0)
def pc(v, d=1): return nfmt(v * 100, d) + '%'
def esc(s): return html.escape(str(s))
sm = sum
MES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set']; ML = [m + '/26' for m in MES]
L = {x['k']: x for x in D['dre']}; ROB = L['rob']['v']
J = [6, 7, 8]; avgJ = lambda f: sm(f(i) for i in J) / 3
EST_MES = avgJ(lambda i: D['est_tot'][i]); ROB_MES = avgJ(lambda i: ROB[i]); EST_PCT = EST_MES / ROB_MES
d = O['d']; R = O['R']; P = D['parc']['dre']; LINES = O['lines']; ADJN = O.get('adjn', ['', ''])
TSRV = {'PESCOÇO DRA PATRICIA': 'Pescoço', 'TRICOLOGIA': 'Tricologia', 'FACIAL DRA PATRICIA': 'Facial',
        'CORPORAL DRA PATRICIA': 'Corporal', 'CONSULTA DRA PATRICIA FABRINI': 'Consulta'}
def parc(l):
    m = re.search(r'\((\d+/\d+)\)\s*$', l['desc'])
    s = ('Consulta' if l['t'] == 'Consulta' else TSRV.get(l['srv'], l['srv'])) + (' · parcela ' + m.group(1) if m else '')
    if l.get('man'): s += ' · ajuste da diretoria'
    return s
RAT = O['rat']; CZS_P = d['pass_']; RB_P = round(O['p_base'] - d['pass_'], 2)

# política (calculadora / base de cálculo)
IXP = dict(pis=.0065, cof=.03, iss=0.0, cart=.022, est=EST_PCT, ins=.0349, viag=.0069, out=.0043, amort=61139.78 / 24, irc=.0768, irp=.0228, pat=.5)
def calc(rec, cons=0.0, ix=IXP):
    r = dict(rec=rec, pis=rec * ix['pis'], cof=rec * ix['cof'], iss=rec * ix['iss'], cart=rec * ix['cart'], ins=rec * ix['ins'], viag=rec * ix['viag'],
             out=rec * ix['out'], est=rec * ix['est'], am=ix['amort'], irc=cons * ix['irc'] + (rec - cons) * ix['irp'])
    r['res'] = rec - sm(r[k] for k in ('pis', 'cof', 'iss', 'cart', 'ins', 'viag', 'out', 'est', 'am', 'irc'))
    r['base'] = max(0, r['res']); r['pat'] = r['base'] * ix['pat']; r['cli'] = r['base'] - r['pat']; r['clitot'] = r['cli'] + r['est'] + r['am']
    return r
EX = calc(100000.0)
assert abs(EX['pat'] - 33124.82) < 0.01

# ------------------------------------------------------------------ conteúdo compartilhado (mesmos textos do painel)
DRE_OUT = [  # (rótulo, valor, tipo) t=total p=patrícia c=clínica
    ('Receita — consultas da Dra. Patrícia (valor total)', R['R1'], ''), ('Receita — procedimentos de tricologia', R['R2'], ''),
    ('Receita — facial, pescoço e corporal da Dra. Patrícia', R['R3'], ''), ('(=) Faturamento da parceria', d['rec'], 't'),
    ('(−) PIS (0,65%)', -d['pis'], ''), ('(−) COFINS (3%)', -d['cof'], ''),
    (f'(−) Antecipação do cartão de crédito (8,74% de {brl(d["cred"])})', -d['cart'], ''), ('(=) Receita líquida', d['rl'], 't'),
    ('(−) Insumos e ativos (consumo estimado, 3,49%)', -d['ins'], ''), ('(−) Passagens aéreas (75% da fatura)', -d['pass_'], ''),
    ('(−) Hotel, alimentação e transporte', -d['loc'], ''), (f'(−) Taxa de estrutura da clínica ({pc(O["ix"]["est"], 2)})', -d['est'], ''),
    ('(=) Resultado operacional', d['ebitda'], 't'), ('(−) Amortização do investimento (parcela 1 de 24 × 75%)', -d['am'], ''),
    ('(−) IRPJ e CSLL', -d['irc'], ''), ('(=) Resultado de Cruzeiro do Sul', d['res'], 't'),
    ('Dra. Patrícia — 50%', d['pat'], 'p'), ('Clínica Núcleo S — 50%', d['cli'], 'c'),
    ('Clínica Núcleo S — total (50% + estrutura + amortização)', d['clitot'], 'c')]
WHY = {'Pessoal': ['Equipe: salários, encargos, benefícios e pró-labore', 'Entra a equipe de apoio (recepção, enfermagem, limpeza) e seus encargos. Fica de fora: pró-labore, bônus, rescisões e cursos.'],
       'Ocupação e Infraestrutura': ['Aluguel, energia, internet, limpeza, segurança e manutenção', 'Entra tudo: são as salas e a estrutura física que a parceria usa.'],
       'Administrativo e Profissionais': ['Contabilidade, jurídico, escritório e consultorias de gestão', 'Entra contabilidade, jurídico, escritório e copa. Fica de fora: consultorias e mentorias de gestão da clínica.'],
       'Tecnologia': ['Sistema de gestão e informática', 'Entra tudo: agenda, prontuário e financeiro rodam no mesmo sistema.'],
       'Taxas e Obrigações Regulatórias': ['Alvarás e taxas de funcionamento', 'Entra tudo: a clínica só funciona com eles.'],
       'Equipamentos e Compliance Clínico': ['Biossegurança, resíduos e manutenção de aparelhos', 'Entra coleta de resíduos e biossegurança. Fica de fora: manutenção de aparelhos de outros serviços.'],
       'Despesas Financeiras': ['Tarifas bancárias e maquineta', 'Entra tarifa e aluguel de maquineta. Fica de fora: juros e IOF.'],
       'Custos Diretos': ['Medicamentos, insumos, honorários e comissões dos outros serviços', 'Não entra: são custos do emagrecimento, implantes e injetáveis. A tricologia paga só os insumos dela.'],
       'Despesas sem identificação': ['Faturas de cartão sem detalhamento', 'Não entra: não se cobra do parceiro despesa sem comprovante.'],
       'Deduções da Receita': ['Impostos sobre o faturamento da clínica', 'Não entra aqui: a tricologia paga os impostos sobre o próprio faturamento (passo 1).'],
       'Marketing': ['Marketing institucional', 'Não entra. Campanha própria da tricologia entra 100% nela, se aprovada pelos dois.'],
       'Rateios e Overhead Compartilhado': ['Royalties de outro procedimento', 'Não entra.'], 'Tributos sobre Lucro': ['IRPJ e CSLL da clínica', 'Não entra aqui: a tricologia paga os dela (passo 8).']}
EST_ROWS = []
for r in D['razao']:
    if r['tipo'] != 'Despesa' or not r['dre']: continue
    tot = avgJ(lambda i: r['v'][i]); ent = avgJ(lambda i: sm(c['est'][i] for c in r['contas']))
    if tot > 0.5: EST_ROWS.append((r['g'], tot, ent))
EST_ROWS.sort(key=lambda x: (-x[2], -x[1]))
T_DESP = sm(x[1] for x in EST_ROWS); E_DESP = sm(x[2] for x in EST_ROWS)
assert abs(E_DESP - EST_MES) < 0.05
STEPS = [
    ('1', 'PIS e COFINS', pc(IXP['pis'] + IXP['cof'], 2) + ' do faturamento', EX['pis'] + EX['cof'], 'Impostos federais que a clínica recolhe sobre todo o faturamento, inclusive o da tricologia.', 'Lei do Lucro Presumido: PIS 0,65% + COFINS 3%. Aparecem na DRE da clínica como deduções da receita.'),
    ('2', 'Taxas de cartão e antecipação', pc(IXP['cart'], 2) + ' do faturamento (média)', EX['cart'], 'O que a maquininha e o banco cobram para a clínica receber as vendas no cartão.', 'Média da clínica: R$ 60,2 mil ÷ R$ 2,72 mi faturados em jan–jul/26. Em out/26 todas as vendas no crédito foram antecipadas a 8,74%: no fechamento usa-se a taxa real.'),
    ('3', 'Insumos e ativos', 'valor real das notas', EX['ins'], 'Ativos capilares, seringas, luvas, gazes e pentes usados nos atendimentos dela.', 'Compras marcadas TRICOLOGIA. Em ago–set foram R$ 6.281,73 (3,49% do faturamento): é o índice do exemplo.'),
    ('4', 'Viagens', 'valor real dos comprovantes', EX['viag'], 'Passagem aérea, hotel, alimentação e transporte de cada visita.', 'A clínica paga e desconta da receita da tricologia (regra de 08/10). Contas 2.02.06.001 a 004.'),
    ('5', 'Outras despesas diretas', 'valor real', EX['out'], 'Recepção e ambientação dos dias de atendimento da tricologia.', 'Despesas marcadas TRICOLOGIA ou PATRICIA na descrição.'),
    ('6', 'Taxa de estrutura da clínica', pc(EST_PCT, 2) + ' do faturamento', EX['est'], 'A parte do custo fixo que todo atendimento usa: sala, equipe de apoio, recepção, energia, sistema, limpeza, contabilidade.', f'Estrutura compartilhada média de jul–set ({brl(EST_MES)}) ÷ faturamento médio da clínica ({brl(ROB_MES)}).'),
    ('7', 'Amortização do investimento', brl(IXP['amort']) + ' por mês', EX['am'], 'Devolve à clínica os aparelhos e a consultoria de implantação, que ela pagou sozinha.', 'R$ 61.139,78 ÷ 24 meses, de out/26 a set/28. Os aparelhos continuam da clínica.'),
    ('8', 'IRPJ e CSLL', pc(IXP['irp'], 2) + ' (procedimentos) · ' + pc(IXP['irc'], 2) + ' (consultas)', EX['irc'], 'Impostos sobre o lucro, calculados sobre o faturamento.', 'Lucro Presumido: 8% × 15% + 12% × 9% nos procedimentos; 32% × 24% nas consultas.'),
    ('9', 'Partilha 50/50', '50% para cada', EX['pat'], 'O que sobra depois dos passos 1 a 8 é dividido meio a meio.', 'Mês negativo não tem partilha: o valor é compensado no mês seguinte. Pagamento à PJ dela, contra nota fiscal, até o dia 15.')]
ENTRA = ['Impostos sobre o faturamento da parceria (PIS, COFINS, ISS se houver)', 'Taxas de cartão e antecipação das vendas da parceria', 'Insumos e ativos dos procedimentos',
         'Passagem aérea, hotel, alimentação e transporte da Dra. Patrícia: a clínica paga e desconta da receita da tricologia', 'Marketing específico aprovado pelos dois',
         'Manutenção dos aparelhos da tricologia', 'Taxa de estrutura: a parte da despesa fixa da clínica usada pela parceria', 'Amortização dos aparelhos e da consultoria de implantação',
         'IRPJ e CSLL sobre o faturamento da parceria', 'Reembolsos e devoluções a pacientes da parceria']
NAO = ['Medicamentos, implantes e honorários de outros serviços', 'Marketing institucional da clínica', 'Pró-labore e retiradas dos sócios', 'Empréstimos e parcelamentos de impostos antigos',
       'Equipamentos e obras de outras áreas', 'Faturas de cartão sem detalhamento e despesa sem comprovante', 'Consultorias e mentorias de gestão da clínica',
       'Transferências entre contas e dinheiro enviado para a conta do sócio', 'Multas e juros por atraso da clínica', 'Comissões da equipe comercial de outros serviços']
OPER = ['Receita da parceria: serviço de tricologia ou "DRA PATRICIA". Despesa: "TRICOLOGIA" ou "PATRICIA" na descrição. O que não estiver marcado não entra.',
        'Fechamento até o dia 10; a Sra. Viviane confere os lançamentos da tricologia no sistema.', 'Repasse até o dia 15 do mês seguinte, contra nota fiscal da empresa da Dra. Patrícia (conta 2.02.04.002).',
        'Mês negativo não tem partilha: o valor é compensado nos meses seguintes.', 'Taxa de estrutura revisada a cada 6 meses com a DRE da clínica.',
        'Prestação de contas mensal: DRE da parceria e razão com comprovantes, aberta às duas partes.']
PEN = [(f'Outubro · Cruzeiro do Sul (05–07/10): fechado, {brl(d["rec"])} de faturamento (consultas pelo valor total + ajustes)', 'OK', 'Fechado'),
       (f'Pescoço de R$ 8.100 (PIX em 05/10) fora do relatório: lançar no sistema · {ADJN[0]}', 'ATENÇÃO', 'Lançar'),
       (f'Saldo de R$ 5.000 do corporal considerado pago: conferir a entrada do PIX · {ADJN[1]}', 'ATENÇÃO', 'Conferir'),
       ('Outubro · Rio Branco (08/10): aguardando o relatório de transações', 'PENDENTE', 'Pendente'),
       ('Consultas pelo valor total em outubro (R$ 800 / R$ 900): tirar de ago/set no sistema a 1ª parte de R$ 400 de cada uma (R$ 2.400)', 'ATENÇÃO', 'Sra. Viviane'),
       ('Hotel, alimentação e transporte em Cruzeiro do Sul: sem comprovante no relatório', 'PENDENTE', 'Localizar'),
       ('2ª parcela da passagem (R$ 2.160, fatura de nov/26): entra no fechamento de novembro', 'PENDENTE', 'A confirmar'),
       ('MDR do cartão (R$ 646,14) já está dentro da antecipação de 8,74%?', 'PENDENTE', 'A confirmar'),
       ('Pacotes vendidos em ago/set com sessões feitas a partir de outubro: partilhar ou não', 'PENDENTE', 'A validar'),
       ('Insumos sem compra no mês: consumo estimado de 3,49% até a ficha técnica ter os preços das notas', 'PENDENTE', 'A validar'),
       ('Taxa de estrutura pelo método A (18,46%), revisão semestral', 'PENDENTE', 'A validar'),
       ('Facial e corporal da Dra. Patrícia entram na parceria?', 'PENDENTE', 'A validar'),
       ('Investimento: amortizar R$ 61,1 mil em 24× a partir de out/26', 'PENDENTE', 'A validar'),
       ('Início da partilha: receitas a partir de out/26; jul–set 100% da clínica', 'OK', 'Decidido'),
       ('Viagens da Dra. Patrícia: a clínica paga e desconta da receita da tricologia', 'OK', 'Decidido')]
HIST = [('9', 'Receita bruta — consultas da Dra. Patrícia', ''), ('10', 'Receita bruta — procedimentos de tricologia', ''), ('11', 'Receita bruta — facial e corporal da Dra. Patrícia', ''),
        ('12', '(=) Receita bruta da parceria', 't'), ('13', '(−) PIS e COFINS', ''), ('15', '(−) Taxas de cartão e antecipação', ''), ('17', '(=) Receita líquida', 't'),
        ('19', '(−) Insumos e ativos', ''), ('20', '(=) Margem de contribuição', 't'), ('22', '(−) Viagens e estadia da Dra. Patrícia', ''), ('24', '(−) Outras despesas diretas', ''),
        ('25', '(−) Taxa de estrutura da clínica', ''), ('26', '(=) Resultado operacional', 't'), ('28', '(−) IRPJ e CSLL presumidos', ''),
        ('29', '(=) Resultado líquido do mês', 't'), ('35', 'Resultado que ficou com a clínica (100%)', 'c')]
# pacientes
PAC = {}
for l in LINES:
    p = PAC.setdefault(l['p'], dict(cons=0, tri=0, fac=0, pes=0, cor=0, n=0, met=set(), al=0))
    k = 'cons' if l['t'] == 'Consulta' else {'TRICOLOGIA': 'tri', 'FACIAL DRA PATRICIA': 'fac', 'PESCOÇO DRA PATRICIA': 'pes', 'CORPORAL DRA PATRICIA': 'cor'}[l['srv']]
    p[k] += l['v']; p['n'] += 1; p['al'] += 1 if l['a'] else 0
    m = l['met'].lower(); p['met'].add('cartão parcelado' if 'parcelado' in m else 'cartão à vista' if 'vista' in m else m)
for p in PAC.values(): p['tot'] = p['cons'] + p['tri'] + p['fac'] + p['pes'] + p['cor']
assert abs(sm(p['tot'] for p in PAC.values()) - d['rec']) < 0.01
PAC_ORD = sorted(PAC.items(), key=lambda x: -x[1]['tot'])

# ================================================================== PDF (HTML de impressão)
CSS = open(SP + '/fonts/local.css', encoding='utf-8').read() + r'''
@page{size:A4;margin:16mm 14mm 18mm 14mm;
  @bottom-left{content:"Clínica Núcleo S · Parceria de Tricologia · relatório dos painéis · ''' + DATA_REF + r'''";font:8pt "IBM Plex Sans",sans-serif;color:#6b7e82}
  @bottom-right{content:"Página " counter(page) " de " counter(pages);font:8pt "IBM Plex Sans",sans-serif;color:#6b7e82}}
@page:first{margin:0;@bottom-left{content:none}@bottom-right{content:none}}
@page land{size:A4 landscape;margin:14mm 14mm 16mm 14mm}
.land{page:land}
.keep{break-inside:avoid}
:root{--ink:#10272b;--ink2:#3d5357;--mut:#6b7e82;--line:#d9e3e3;--soft:#f3f7f7;--acc:#0a6a70;--accs:#e0eff0;--s1:#2a78d6;--s2:#eb6834;--s3:#1baf7a;--n1:#7d8f92;--n2:#c3cdce;--good:#11795a;--goodb:#e2f3eb;--bad:#bf3a2b;--badb:#fbe8e5;--warn:#94600a;--warnb:#fcefd6}
*{box-sizing:border-box}html,body{margin:0;padding:0}
body{font:9.5pt/1.45 "IBM Plex Sans",sans-serif;color:var(--ink);-webkit-print-color-adjust:exact;print-color-adjust:exact;font-variant-numeric:tabular-nums}
h1,h2,h3{font-family:"Manrope",sans-serif;margin:0;color:var(--ink)}
h2{font-size:15pt;font-weight:800;margin:0 0 2mm}
h3{font-size:11pt;font-weight:700;margin:5mm 0 2mm}
p{margin:0 0 2mm}
.sec{break-before:page}
.eyebrow{font:700 7.5pt "Manrope",sans-serif;letter-spacing:.12em;text-transform:uppercase;color:var(--acc);margin-bottom:1mm}
.lead{color:var(--ink2);margin-bottom:4mm;max-width:165mm}
.note{font-size:8pt;color:var(--mut)}
.cover{height:297mm;width:210mm;background:#0c2a2f;color:#e6f0f0;padding:28mm 20mm 18mm;display:flex;flex-direction:column}
.cover .mark{width:16mm;height:16mm;border-radius:4mm;background:#0a6a70;display:grid;place-items:center;font:800 15pt "Manrope";color:#fff}
.cover h1{color:#fff;font-size:28pt;font-weight:800;line-height:1.1;margin:14mm 0 4mm;letter-spacing:-.01em}
.cover .sub{font-size:12pt;color:#a9c4c6;max-width:150mm}
.cover .kp{display:grid;grid-template-columns:1fr 1fr;gap:5mm;margin-top:16mm}
.cover .kp div{border:1px solid #1d4a51;border-radius:3mm;padding:5mm 6mm;background:#10353b}
.cover .kp small{display:block;color:#8fb0b3;font-size:8.5pt}.cover .kp b{font:800 17pt "Manrope";color:#fff}
.cover .toc{margin-top:auto;border-top:1px solid #1d4a51;padding-top:6mm;columns:2;column-gap:10mm;font-size:9pt;color:#cfe0e0}
.cover .toc div{break-inside:avoid;margin-bottom:1.5mm}.cover .toc b{color:#fff}
.cover .foot{margin-top:6mm;font-size:8pt;color:#7f9ea1}
.kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:3mm;margin:3mm 0 4mm}
.kpi{border:1px solid var(--line);border-radius:2.5mm;padding:3mm 3.5mm;background:#fff}
.kpi small{display:block;color:var(--mut);font-size:7.5pt}.kpi b{font:800 13pt "Manrope";display:block;white-space:nowrap}.kpi span{font-size:7.5pt;color:var(--ink2)}
.banner{background:var(--accs);border:1px solid #b9d9da;border-radius:2.5mm;padding:3mm 4mm;margin:2mm 0 4mm;font-size:9pt}
table{width:100%;border-collapse:collapse;font-size:8.6pt;margin:1mm 0 3mm}
th{font:600 7.6pt "IBM Plex Sans";color:var(--ink2);text-align:right;background:var(--soft);border-bottom:1px solid var(--line);padding:1.6mm 2mm;text-transform:none}
td{padding:1.4mm 2mm;border-bottom:1px solid var(--line);text-align:right;vertical-align:top}
th:first-child,td:first-child{text-align:left}
td.l,th.l{text-align:left}
tr.t td{font-weight:700;background:var(--soft)}
tr.p td{color:var(--s1);font-weight:700}tr.c td{color:#0f7f59;font-weight:700}
tr{break-inside:avoid}thead{display:table-header-group}
.sm table,table.sm{font-size:7.6pt}table.sm td,table.sm th{padding:1.1mm 1.6mm}
.grid2{display:grid;grid-template-columns:1fr 1fr;gap:6mm}
.chip{display:inline-block;font:600 7pt "IBM Plex Sans";border-radius:9pt;padding:.4mm 2mm;white-space:nowrap}
.ok{background:var(--goodb);color:var(--good)}.at{background:var(--badb);color:var(--bad)}.pe{background:var(--warnb);color:var(--warn)}
.stack{display:flex;height:9mm;border-radius:1.5mm;overflow:hidden;gap:.6mm;margin:2mm 0}
.stack div{display:flex;align-items:center;justify-content:center;color:#fff;font:600 7.5pt "IBM Plex Sans";white-space:nowrap;overflow:hidden}
.legend{display:flex;flex-wrap:wrap;gap:2mm 5mm;font-size:7.8pt;color:var(--ink2)}
.legend i{display:inline-block;width:2.6mm;height:2.6mm;border-radius:.6mm;margin-right:1.2mm;vertical-align:-.3mm}
.formula{display:flex;align-items:center;gap:3mm;margin:3mm 0}
.formula div{border:1px solid var(--line);border-radius:2.5mm;padding:3mm 4mm;background:var(--soft)}
.formula small{display:block;font-size:7.5pt;color:var(--mut)}.formula b{font:800 13pt "Manrope"}
.formula .eq{border:0;background:none;font:800 14pt "Manrope";color:var(--mut);padding:0}
.formula .res{background:var(--acc);border-color:var(--acc)}.formula .res b,.formula .res small{color:#fff}
ul.ck{list-style:none;padding:0;margin:0}ul.ck li{padding:1mm 0 1mm 5mm;position:relative;border-bottom:1px solid var(--line);font-size:8.6pt}
ul.ck.ok li::before{content:"✓";position:absolute;left:0;color:var(--good);font-weight:800}ul.ck.no li::before{content:"✕";position:absolute;left:0;color:var(--bad);font-weight:800}
.box{border:1px solid var(--line);border-radius:2.5mm;padding:3mm 4mm;break-inside:avoid}
.al{color:var(--bad);font-size:7.2pt}
.pat{font-weight:600}
svg text{font-family:"IBM Plex Sans",sans-serif}
'''
def tbl(head, rows, cls='', aligns=None):
    h = '<table class="%s"><thead><tr>%s</tr></thead><tbody>' % (cls, ''.join(f'<th class="{"l" if (aligns and aligns[i]=="l") else ""}">{esc(x)}</th>' for i, x in enumerate(head)))
    for r in rows:
        cells, rc = (r[0], r[1]) if isinstance(r, tuple) and len(r) == 2 and isinstance(r[1], str) and isinstance(r[0], list) else (r, '')
        h += f'<tr class="{rc}">' + ''.join(f'<td class="{"l" if (aligns and aligns[i]=="l") else ""}">{c}</td>' for i, c in enumerate(cells)) + '</tr>'
    return h + '</tbody></table>'
def chip(st):
    return {'OK': '<span class="chip ok">✓ OK</span>', 'ATENÇÃO': '<span class="chip at">! Atenção</span>', 'PENDENTE': '<span class="chip pe">! Pendente</span>'}[st]
def waterfall_svg(rows, W=640, rowh=19):
    """rows: (rótulo, valor, tipo) tipo t=total n=dedução e=estrutura p=patrícia c=clínica; escala = primeiro total."""
    mx = rows[0][1]; lw, vw = 190, 130; bw = W - lw - vw - 10; H = rowh * len(rows) + 6
    s = f'<svg width="100%" viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg">'
    s += f'<line x1="{lw}" x2="{lw}" y1="0" y2="{H}" stroke="#d9e3e3"/><line x1="{lw+bw}" x2="{lw+bw}" y1="0" y2="{H}" stroke="#d9e3e3"/>'
    run = 0; patv = 0
    for i, (lab, v, k) in enumerate(rows):
        y = 3 + i * rowh
        if k == 't': a, b = 0, v; run = v; col = '#7d8f92'
        elif k == 'p': a, b = 0, v; patv = v; col = '#2a78d6'
        elif k == 'c': a, b = patv, patv + v; col = '#1baf7a'
        else: a, b = run + v, run; run += v; col = '#eb6834' if k == 'e' else '#c3cdce'
        x0 = lw + max(0, min(a, b)) / mx * bw; w = max(1.2, abs(b - a) / mx * bw)
        bold = k in ('t', 'p', 'c')
        s += f'<text x="0" y="{y+12.5}" font-size="10.5" fill="#10272b" font-weight="{700 if bold else 400}">{esc(lab)}</text>'
        s += f'<rect x="{x0:.1f}" y="{y+2}" width="{w:.1f}" height="{rowh-6}" rx="2" fill="{col}"/>'
        s += f'<text x="{W}" y="{y+12.5}" font-size="10.5" text-anchor="end" fill="#10272b" font-weight="{700 if bold else 500}">{esc(brl0(v))} <tspan fill="#6b7e82" font-size="9">{esc(pc(v/mx))}</tspan></text>'
    return s + '</svg>'
def stack_html(r):
    parts = [('Impostos e taxas', r['pis'] + r['cof'] + r['iss'] + r['cart'] + r['irc'], '#7d8f92'), ('Custos diretos da tricologia', r['ins'] + r['viag'] + r['out'], '#c3cdce'),
             ('Estrutura e investimento (repõem a clínica)', r['est'] + r['am'], '#eb6834'), ('Partilha da clínica', r['cli'], '#1baf7a'), ('Partilha da Dra. Patrícia', r['pat'], '#2a78d6')]
    T = sm(p[1] for p in parts)
    h = '<div class="stack">' + ''.join(f'<div style="width:{p[1]/T*100:.2f}%;background:{p[2]}">{("R$ " + nfmt(p[1]/T*100,1)) if p[1]/T > .07 else ""}</div>' for p in parts) + '</div>'
    h += '<div class="legend">' + ''.join(f'<span><i style="background:{p[2]}"></i>{esc(p[0])}: <b>R$ {nfmt(p[1]/T*100,1)}</b></span>' for p in parts) + '</div>'
    return h
def wf_rows(r, viag_label='Viagens'):
    rows = [('Faturamento', r['rec'], 't'), ('PIS e COFINS', -(r['pis'] + r['cof'] + r['iss']), 'n'), ('Taxas de cartão', -r['cart'], 'n'), ('Insumos e ativos', -r['ins'], 'n'),
            (viag_label, -r['viag'], 'n'), ('Outras despesas diretas', -r['out'], 'n'), ('Taxa de estrutura', -r['est'], 'e'), ('Amortização do investimento', -r['am'], 'e'),
            ('IRPJ e CSLL', -r['irc'], 'n'), ('Resultado a dividir', r['base'], 't'), ('Dra. Patrícia — 50%', r['pat'], 'p'), ('Clínica Núcleo S — 50%', r['cli'], 'c')]
    return rows
ROUT = dict(rec=d['rec'], pis=d['pis'], cof=d['cof'], iss=d['iss'], cart=d['cart'], ins=d['ins'], viag=d['pass_'] + d['loc'], out=d['out'], est=d['est'], am=d['am'], irc=d['irc'], base=d['base'], pat=d['pat'], cli=d['cli'])
def est_chart_svg(W=640, H=200):
    v = [D['est_tot'][i] / ROB[i] for i in range(9)]; mx = 0.30; ml, mb, mt = 34, 20, 8; bw = (W - ml - 6) / 9
    y = lambda x: mt + (H - mt - mb) * (1 - x / mx)
    s = f'<svg width="100%" viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg">'
    for t in (0, .1, .2, .3):
        s += f'<line x1="{ml}" x2="{W}" y1="{y(t):.1f}" y2="{y(t):.1f}" stroke="#e3eaea"/><text x="{ml-5}" y="{y(t)+3:.1f}" font-size="9" text-anchor="end" fill="#6b7e82">{int(t*100)}%</text>'
    for i, x in enumerate(v):
        cx = ml + i * bw + bw / 2; s += f'<rect x="{cx-11:.1f}" y="{y(x):.1f}" width="22" height="{y(0)-y(x):.1f}" rx="2" fill="#2a78d6"/>'
        s += f'<text x="{cx:.1f}" y="{y(x)-3:.1f}" font-size="8.5" text-anchor="middle" fill="#10272b">{nfmt(x*100,1)}%</text><text x="{cx:.1f}" y="{H-6}" font-size="9" text-anchor="middle" fill="#6b7e82">{ML[i]}</text>'
    s += f'<line x1="{ml}" x2="{W}" y1="{y(EST_PCT):.1f}" y2="{y(EST_PCT):.1f}" stroke="#3d5357" stroke-dasharray="5 4" stroke-width="1.3"/>'
    s += f'<text x="{W}" y="{y(EST_PCT)-4:.1f}" font-size="9" text-anchor="end" fill="#3d5357" font-weight="600">média jul–set {pc(EST_PCT,2)}</text>'
    return s + '</svg>'

H = ['<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>Parceria de Tricologia — relatório dos painéis</title><style>' + CSS + '</style></head><body>']
# capa
H.append(f'''<section class="cover"><div class="mark">NS</div>
<h1>Parceria de Tricologia<br>Relatório dos painéis</h1>
<div class="sub">Clínica Núcleo S × Dra. Patrícia Fabrini · fechamento de outubro/2026 (Cruzeiro do Sul, 05 a 07/10), base de cálculo, taxa de estrutura, calculadora, histórico e regras.</div>
<div class="kp"><div><small>Faturamento · Cruzeiro do Sul</small><b>{brl(d['rec'])}</b></div><div><small>Resultado · Cruzeiro do Sul</small><b>{brl(d['res'])}</b></div>
<div><small>Dra. Patrícia · 50%</small><b>{brl(d['pat'])}</b></div><div><small>Clínica Núcleo S · total</small><b>{brl(d['clitot'])}</b></div></div>
<div class="toc"><div><b>1</b> · Fechamento de outubro/2026</div><div><b>2</b> · Consultas, procedimentos e pacientes</div><div><b>3</b> · Lançamentos de Cruzeiro do Sul</div>
<div><b>4</b> · Viagem e Rio Branco</div><div><b>5</b> · Auditoria</div><div><b>6</b> · Base de cálculo</div><div><b>7</b> · Taxa de estrutura</div><div><b>8</b> · Calculadora</div>
<div><b>9</b> · Histórico e investimento</div><div><b>10</b> · Regras e pendências</div></div>
<div class="foot">Controladoria · Rogério Marinho · gerado em {DATA_REF} · fonte: relatórios de transações do sistema, plano de contas oficial e ajustes da diretoria de 08/10.</div></section>''')
# 1 fechamento
H.append(f'''<section class="sec"><div class="eyebrow">1 · Fechamento de outubro/2026</div><h2>DRE da parceria · Cruzeiro do Sul · 05 a 07/10</h2>
<p class="lead">Primeiro mês com partilha 50/50. {O['n_sis']} lançamentos do relatório de transações, todos da Dra. Patrícia, mais {O['n_adj']} ajustes da diretoria. Consultas pelo valor total (R$ 800 / R$ 900). Todas as vendas no cartão de crédito foram antecipadas a 8,74%.</p>
<div class="kpis"><div class="kpi"><small>Faturamento</small><b>{brl(d['rec'])}</b><span>{O['n_sis']} do relatório + {O['n_adj']} ajustes</span></div>
<div class="kpi"><small>Resultado</small><b>{brl(d['res'])}</b><span>{pc(d['res']/d['rec'])} do faturamento</span></div>
<div class="kpi"><small>Dra. Patrícia · 50%</small><b style="color:#2a78d6">{brl(d['pat'])}</b><span>{pc(d['pat']/d['rec'])} do faturamento</span></div>
<div class="kpi"><small>Clínica Núcleo S · total</small><b style="color:#0f7f59">{brl(d['clitot'])}</b><span>50% + estrutura + amortização</span></div></div>
{tbl(['Linha', 'R$', '% do faturamento'], [([esc(l), brl(v), pc(v/d['rec'], 2)], {'t': 't', 'p': 'p', 'c': 'c'}.get(k, '')) for l, v, k in DRE_OUT])}
<div class="keep"><h3>Do faturamento ao valor de cada um</h3>{waterfall_svg(wf_rows(ROUT, 'Passagens e viagem'))}
<h3>De cada R$ 100 faturados em Cruzeiro do Sul</h3>{stack_html(ROUT)}</div></section>''')
# 2 consultas e pacientes
cons = [l for l in LINES if l['t'] == 'Consulta']; proc = [l for l in LINES if l['t'] == 'Procedimento']
svsum = lambda s: sm(l['v'] for l in proc if l['srv'] == s)
H.append(f'''<section class="sec"><div class="eyebrow">2 · Consultas, procedimentos e pacientes</div><h2>Consultas e procedimentos</h2>
<p class="lead">Consultas pelo valor total: R$ 400 no sistema = consulta de R$ 800; R$ 500 = consulta de R$ 900 (a 1ª parte de R$ 400 foi recebida em ago/set). Procedimentos pelo valor do sistema mais os ajustes da diretoria de 08/10.</p>
{tbl(['Tipo / serviço', 'Lançamentos', 'Valor', '% do total'], [
  (['<b>Consultas (valor total)</b>', str(len(cons)), brl(sm(l['v'] for l in cons)), pc(sm(l['v'] for l in cons)/d['rec'])], ''),
  (['<b>Procedimentos</b>', str(len(proc)), brl(sm(l['v'] for l in proc)), pc(sm(l['v'] for l in proc)/d['rec'])], ''),
  (['&nbsp;&nbsp;Tricologia', str(sum(1 for l in proc if l['srv']=='TRICOLOGIA')), brl(svsum('TRICOLOGIA')), pc(svsum('TRICOLOGIA')/d['rec'])], ''),
  (['&nbsp;&nbsp;Facial', str(sum(1 for l in proc if l['srv']=='FACIAL DRA PATRICIA')), brl(svsum('FACIAL DRA PATRICIA')), pc(svsum('FACIAL DRA PATRICIA')/d['rec'])], ''),
  (['&nbsp;&nbsp;Pescoço', str(sum(1 for l in proc if l['srv']=='PESCOÇO DRA PATRICIA')), brl(svsum('PESCOÇO DRA PATRICIA')), pc(svsum('PESCOÇO DRA PATRICIA')/d['rec'])], ''),
  (['&nbsp;&nbsp;Corporal', str(sum(1 for l in proc if l['srv']=='CORPORAL DRA PATRICIA')), brl(svsum('CORPORAL DRA PATRICIA')), pc(svsum('CORPORAL DRA PATRICIA')/d['rec'])], ''),
  (['Total', str(len(LINES)), brl(d['rec']), '100,0%'], 't')])}
<h3>Pacientes atendidas</h3>
{tbl(['Paciente', 'Consulta', 'Tricologia', 'Facial', 'Pescoço', 'Corporal', 'Total', 'Lanç.', 'Forma de pagamento'],
   [([f'<span class="pat">{esc(n)}</span>' + (f' <span class="chip at">! {p["al"]}</span>' if p['al'] else ''), nfmt(p['cons'],0) if p['cons'] else '–', nfmt(p['tri'],0) if p['tri'] else '–',
      nfmt(p['fac'],0) if p['fac'] else '–', nfmt(p['pes'],0) if p['pes'] else '–', nfmt(p['cor'],0) if p['cor'] else '–', '<b>' + nfmt(p['tot']) + '</b>', str(p['n']), esc(' + '.join(sorted(p['met'])))], '') for n, p in PAC_ORD]
   + [(['Total', nfmt(sm(p['cons'] for p in PAC.values()),0), nfmt(sm(p['tri'] for p in PAC.values()),0), nfmt(sm(p['fac'] for p in PAC.values()),0), nfmt(sm(p['pes'] for p in PAC.values()),0),
        nfmt(sm(p['cor'] for p in PAC.values()),0), nfmt(d['rec']), str(len(LINES)), ''], 't')], 'sm', ['l','r','r','r','r','r','r','r','l'])}
<p class="note">Valores em R$.</p>
<p class="note">! = lançamentos com alerta da auditoria (detalhe na seção 3).</p></section>''')
# 3 lançamentos (parcelas do mesmo pacote agrupadas; detalhe completo no Excel)
def met_s(m):
    ml = m.lower()
    return 'Cartão de crédito parcelado' if 'parcelado' in ml else 'Cartão de crédito à vista (ELO)' if 'vista' in ml else m
rows = []
for n, pinfo in PAC_ORD:
    rows.append(([f'<b>{esc(n)}</b>', '', '', '', '', '', f'<b>{nfmt(pinfo["tot"])}</b>'], 't'))
    groups = {}
    for l in LINES:
        if l['p'] != n: continue
        m = re.search(r'\((\d+)/(\d+)\)\s*$', l['desc'])
        key = (l['t'], l['srv'], l['met'], l['cta'], bool(l.get('man')), m.group(2) if m else 'x' + l['desc'] + l['ve'])
        groups.setdefault(key, []).append((l, int(m.group(1)) if m else 0, m.group(2) if m else None))
    for key, ls in sorted(groups.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        l0 = ls[0][0]; tot_parc = ls[0][2]
        lab = ('Consulta' if l0['t'] == 'Consulta' else TSRV.get(l0['srv'], l0['srv']))
        if tot_parc: lab += f' · {len(ls)} parcelas ({min(x[1] for x in ls)}/{tot_parc} a {max(x[1] for x in ls)}/{tot_parc}) de {brl(ls[0][0]["vs"] if len(ls)==1 else sorted(x[0]["vs"] for x in ls)[len(ls)//2])}'
        if l0.get('man'): lab += ' · ajuste da diretoria'
        alerts = sorted(set(x[0]['a'] for x in ls if x[0]['a']))
        ves = sorted(set(x[0]['ve'] for x in ls))
        rows.append(([esc(lab) + ''.join(f'<div class="al">{esc(a)}</div>' for a in alerts), ', '.join(ves), esc(met_s(l0['met'])), esc(l0['cta']),
                      nfmt(sm(x[0]['vs'] for x in ls)) if sm(x[0]['vs'] for x in ls) else '–', nfmt(sm(x[0]['cmp'] for x in ls)) if sm(x[0]['cmp'] for x in ls) else '–',
                      nfmt(sm(x[0]['v'] for x in ls))], ''))
rows.append((['Total geral', '', '', '', nfmt(d['sis']), nfmt(d['cmp'] + d['adj']), nfmt(d['rec'])], 't'))
H.append(f'''<section class="sec land"><div class="eyebrow">3 · Lançamentos</div><h2>Lançamentos de Cruzeiro do Sul · 05 a 07/10</h2>
<p class="lead" style="max-width:none">Os {O['n_sis']} lançamentos do relatório de transações e os {O['n_adj']} ajustes da diretoria, por paciente. As parcelas do mesmo pacote no cartão aparecem somadas (cada parcela está no Excel). "Sistema" é o valor do relatório; "Ajuste" é a 1ª parte das consultas ou o ajuste da diretoria; "Considerado" entra na DRE. Valores em R$.</p>
{tbl(['Lançamento', 'Vencimento', 'Forma de pagamento', 'Conta', 'Sistema', 'Ajuste', 'Considerado'], rows, 'sm', ['l','l','l','l','r','r','r'])}</section>''')
# 4 viagem e Rio Branco
pass_rows = [([x[0], esc(x[1]) + (' · ' + x[3] if x[3] else ''), brl(x[4])], '') for x in O['pass_']]
pass_rows += [(['', 'Total das passagens na fatura de 05/10', brl(O['p_base'])], 't'), (['', f'→ Cruzeiro do Sul: {pc(RAT,0)} (3 de 4 dias: 05, 06 e 07/10)', brl(CZS_P)], ''),
              (['', f'→ Rio Branco: {pc(1-RAT,0)} (1 de 4 dias: 08/10)', brl(RB_P)], ''), (['', 'Soma do rateio (= fatura)', brl(CZS_P + RB_P)], 't'),
              (['', esc(O['pass2'][1].replace(', vem na fatura de nov/26', '')) + ' — fica para o fechamento de novembro', '<span class="note">' + brl(O['pass2'][4]) + '</span>'], '')]
H.append(f'''<section class="sec"><div class="eyebrow">4 · Viagem e Rio Branco</div><h2>Viagem SP → Cruzeiro do Sul → Rio Branco → SP</h2>
<p class="lead">Passagens pagas no cartão da clínica (Mastercard final 5917, fatura de 05/10/2026). A clínica paga e desconta da receita da tricologia. Rateio por dias de atendimento.</p>
{tbl(['Data', 'Descrição na fatura', 'Valor'], pass_rows, '', ['l','l','r'])}
<div class="box" style="margin-top:5mm"><h3 style="margin-top:0">Rio Branco · 08/10 <span class="chip pe">! pendente</span></h3>
{tbl(['', 'R$'], [(['Passagens (25% da fatura)', brl(RB_P)], ''), (['Amortização (25% da parcela de outubro)', brl(O['ix']['amort'] * (1 - RAT))], ''), (['Receitas de 08/10', 'a lançar'], '')])}
<p class="note">A partilha de outubro soma Cruzeiro do Sul e Rio Branco. Se Rio Branco der negativo, o valor é compensado no total do mês antes do repasse.</p></div></section>''')
# 5 auditoria
H.append('<section class="sec"><div class="eyebrow">5 · Auditoria</div><h2>Auditoria feita antes da entrega</h2><p class="lead">O que foi conferido e o que ainda depende de confirmação.</p>'
         + tbl(['#', 'Verificação', 'Status'], [([str(i + 1), f'<b>{esc(a[0])}</b><div class="note">{esc(a[2])}</div>', chip(a[1])], '') for i, a in enumerate(O['audit'])], '', ['l','l','r']) + '</section>')
# 6 base de cálculo
H.append(f'''<section class="sec"><div class="eyebrow">6 · Base de cálculo</div><h2>Do faturamento ao valor de cada um</h2>
<p class="lead">A partir das receitas de outubro/2026, tudo o que a parceria fatura, menos tudo o que ela custa, é dividido meio a meio. Exemplo com R$ 100.000 de procedimentos e custos diretos pelos índices médios (no mês real, insumos e viagens entram pelo valor das notas).</p>
{tbl(['Passo', 'Índice', 'No exemplo', 'O que é e de onde vem'], [([f'<b>{s[0]}. {esc(s[1])}</b>', esc(s[2]), ('' if s[0]=='9' else '−') + brl0(s[3]) + (' cada' if s[0]=='9' else ''), f'{esc(s[4])}<div class="note">{esc(s[5])}</div>'], '') for s in STEPS], '', ['l','l','r','l'])}
<div class="grid2"><div class="box"><small class="note">Dra. Patrícia recebe</small><div style="font:800 15pt Manrope;color:#2a78d6">{brl(EX['pat'])}</div><span class="note">{pc(EX['pat']/EX['rec'])} do faturamento</span></div>
<div class="box"><small class="note">Clínica fica com (50% + estrutura + investimento)</small><div style="font:800 15pt Manrope;color:#0f7f59">{brl(EX['clitot'])}</div><span class="note">{pc(EX['clitot']/EX['rec'])} do faturamento</span></div></div>
<div class="keep"><h3>De cada R$ 100 faturados</h3>{stack_html(EX)}</div></section>''')
# 7 taxa de estrutura
H.append(f'''<section class="sec"><div class="eyebrow">7 · Taxa de estrutura</div><h2>Como os {pc(EST_PCT,2)} são calculados</h2>
<p class="lead">Média mensal de jul–set/2026, o período com a estrutura atual das três unidades.</p>
<div class="formula"><div><small>Estrutura compartilhada por mês</small><b>{brl(EST_MES)}</b></div><div class="eq">÷</div><div><small>Faturamento médio da clínica</small><b>{brl(ROB_MES)}</b></div><div class="eq">=</div><div class="res"><small>Taxa de estrutura</small><b>{pc(EST_PCT,2)}</b></div></div>
<h3>A clínica gasta {brl0(T_DESP)} por mês para funcionar ({pc(T_DESP/ROB_MES)} do faturamento)</h3>
<div class="stack"><div style="width:{E_DESP/T_DESP*100:.2f}%;background:#eb6834">Entra na taxa: {brl0(E_DESP)}</div><div style="width:{(T_DESP-E_DESP)/T_DESP*100:.2f}%;background:#c3cdce;color:#10272b">Fica 100% com a clínica: {brl0(T_DESP-E_DESP)}</div></div>
{tbl(['Despesa da clínica (média por mês)', 'R$ por mês', '% fat.', 'Entra na taxa', 'Fica com a clínica', 'Por quê'],
  [([esc(WHY.get(g, [g, ''])[0]), nfmt(t, 0), pc(t/ROB_MES), f'<b style="color:#c4521f">{nfmt(e,0)}</b>' if e > .5 else '–', nfmt(t-e, 0) if t-e > .5 else '–', f'<span class="note">{esc(WHY.get(g, ["", ""])[1])}</span>'], '') for g, t, e in EST_ROWS]
  + [(['Total', nfmt(T_DESP, 0), pc(T_DESP/ROB_MES), nfmt(E_DESP, 0), nfmt(T_DESP-E_DESP, 0), ''], 't'), (['% do faturamento', '', '', pc(E_DESP/ROB_MES, 2), pc((T_DESP-E_DESP)/ROB_MES, 2), '← a taxa de estrutura'], 't')], 'sm', ['l','r','r','r','r','l'])}
<div class="keep"><h3>Taxa de estrutura mês a mês</h3>{est_chart_svg()}
<p class="note">Estrutura do mês ÷ receita bruta do mês. Oscila com o faturamento: por isso a parceria usa a média de jul–set, revisada a cada 6 meses.</p></div></section>''')
# 8 calculadora
IXROWS = [('PIS', pc(IXP['pis'], 2), 'Lei, Lucro Presumido'), ('COFINS', pc(IXP['cof'], 2), 'Lei, Lucro Presumido'), ('ISS', pc(IXP['iss'], 2), 'ISS fixo anual: zero no mês'),
          ('Cartão e antecipação (média)', pc(IXP['cart'], 2), 'Média jan–jul/26; em out/26 a antecipação real foi 8,74% das vendas no crédito'), ('Taxa de estrutura', pc(EST_PCT, 2), 'Estrutura ÷ faturamento, jul–set/26'),
          ('Insumos (sem nota no mês)', pc(IXP['ins'], 2), 'Compras de ago–set ÷ receita'), ('Viagens (sem comprovante)', pc(IXP['viag'], 2), 'Ago–set, subestimado'),
          ('Outras despesas diretas', pc(IXP['out'], 2), 'Ago–set'), ('IRPJ + CSLL — procedimentos', pc(IXP['irp'], 2), '8% × 15% + 12% × 9%'), ('IRPJ + CSLL — consultas', pc(IXP['irc'], 2), '32% × 24%'),
          ('Amortização mensal', brl(IXP['amort']), 'R$ 61.139,78 ÷ 24'), ('Parte da Dra. Patrícia', pc(IXP['pat'], 0), 'Partilha 50/50')]
CST = [('Faturamento da parceria', EX['rec'], 't'), ('(−) PIS', -EX['pis'], ''), ('(−) COFINS', -EX['cof'], ''), ('(−) ISS', -EX['iss'], ''), ('(−) Taxas de cartão e antecipação', -EX['cart'], ''),
       ('(=) Receita líquida', EX['rec'] - EX['pis'] - EX['cof'] - EX['iss'] - EX['cart'], 't'), ('(−) Insumos e ativos', -EX['ins'], ''), ('(−) Viagens e estadia', -EX['viag'], ''),
       ('(−) Outras despesas diretas', -EX['out'], ''), ('(−) Taxa de estrutura da clínica', -EX['est'], ''), ('(−) Amortização do investimento', -EX['am'], ''), ('(−) IRPJ e CSLL', -EX['irc'], ''),
       ('(=) Resultado do mês', EX['res'], 't'), ('Dra. Patrícia', EX['pat'], 'p'), ('Clínica Núcleo S', EX['cli'], 'c'), ('Clínica — total (50% + estrutura + amortização)', EX['clitot'], 'c')]
H.append(f'''<section class="sec"><div class="eyebrow">8 · Calculadora</div><h2>Calculadora da partilha</h2>
<p class="lead">Índices padrão da política e o passo a passo com R$ 100.000 de faturamento (só procedimentos). A planilha Excel desta exportação traz a calculadora com fórmulas: basta digitar o faturamento.</p>
<div class="grid2"><div>{tbl(['Índice', 'Valor', 'Origem'], [([esc(a), b, f'<span class="note">{esc(c)}</span>'], '') for a, b, c in IXROWS], 'sm', ['l','r','l'])}</div>
<div>{tbl(['Passo (exemplo R$ 100.000)', 'R$', '%'], [([esc(a), brl(v), pc(v/EX['rec'], 2)], {'t':'t','p':'p','c':'c'}.get(k, '')) for a, v, k in CST], 'sm')}</div></div></section>''')
# 9 histórico e investimento
H.append(f'''<section class="sec"><div class="eyebrow">9 · Histórico e investimento</div><h2>A conta aplicada aos meses de teste · jul a set/2026</h2>
<p class="lead">Até setembro/2026 a receita da tricologia é 100% da clínica; a partilha 50/50 vale a partir de outubro. Números reais da base, sem partilha (amortização começa em out/26).</p>
{tbl(['Linha', 'jul/26', 'ago/26', 'set/26', 'Acum. jul–set', '% da receita'], [([esc(lab)] + [nfmt(x, 0) if abs(x) > .5 else '–' for x in P[k][1:5]] + [pc(P[k][4]/P['12'][4])], {'t':'t','c':'c'}.get(t, '')) for k, lab, t in HIST], 'sm')}
<div class="grid2" style="margin-top:3mm"><div class="box"><h3 style="margin-top:0">Investimento de implantação</h3>
{tbl(['', 'R$'], [(['Aparelhos (eletroporação, LED, câmera, microscópio), líquidos do estorno', brl(21139.82)], ''), (['Consultoria de implantação: 6 parcelas pagas (abr–set)', brl(34285.68)], ''),
   (['Consultoria de implantação: 7ª parcela (out/26)', brl(5714.28)], ''), (['Investimento total', brl(61139.78)], 't'), (['Amortização: 24 × de out/26 a set/28', brl(61139.78/24) + ' / mês'], '')], 'sm')}
<p class="note">Como a amortização entra antes da divisão, cada parte arca com metade ({brl(61139.78/2)}). Os aparelhos continuam patrimônio da clínica.</p></div>
<div class="box"><h3 style="margin-top:0">Linha do tempo</h3>{tbl(['Quando', 'O quê'], [(['abr–jun/26', 'Consultoria de implantação e compra dos aparelhos'], ''), (['jul/26', 'Material inicial (R$ 13,6 mil)'], ''),
   (['11/08/26', 'Primeiro atendimento da Dra. Patrícia'], ''), (['ago–set/26', 'Período de teste: 100% da clínica'], ''), (['out/26', 'Começa a partilha 50/50: Cruzeiro do Sul 05–07/10, Rio Branco 08/10'], 't'),
   (['set/28', 'Última parcela da amortização'], '')], 'sm', ['l','l'])}</div></div></section>''')
# 10 regras e pendências
H.append('<section class="sec"><div class="eyebrow">10 · Regras e pendências</div><h2>Regras da parceria</h2><p class="lead">A partir das receitas de outubro/2026, tudo o que a parceria fatura, menos tudo o que ela custa (impostos, cartão, insumos, viagens, estrutura da clínica e o investimento), é dividido meio a meio: 50% Clínica Núcleo S, 50% Dra. Patrícia.</p>'
         + '<div class="grid2"><div class="box"><h3 style="margin-top:0">Entra na conta</h3><ul class="ck ok">' + ''.join(f'<li>{esc(x)}</li>' for x in ENTRA) + '</ul></div>'
         + '<div class="box"><h3 style="margin-top:0">Não entra (fica com a clínica)</h3><ul class="ck no">' + ''.join(f'<li>{esc(x)}</li>' for x in NAO) + '</ul></div></div>'
         + '<h3>Como funciona todo mês</h3>' + tbl(['#', 'Regra'], [([str(i + 1), esc(t)], '') for i, t in enumerate(OPER)], '', ['l','l'])
         + '<h3>Pendências antes de apresentar</h3>' + tbl(['Pendência', 'Situação'], [([esc(a), chip(b) + f' <span class="note">{esc(c)}</span>'], '') for a, b, c in PEN], '', ['l','r']) + '</section>')
H.append('</body></html>')
open(OUTD + '/report.html', 'w', encoding='utf-8').write('\n'.join(H))
pdf = OUTD + f'/{NAME}.pdf'
subprocess.run(['/opt/pw-browsers/chromium', '--headless', '--no-sandbox', '--disable-gpu', '--no-pdf-header-footer', '--virtual-time-budget=8000',
                f'--print-to-pdf={pdf}', 'file://' + OUTD + '/report.html'], check=True, capture_output=True, timeout=180)
print('pdf ok', os.path.getsize(pdf))

# ================================================================== EXCEL
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
wb = openpyxl.Workbook()
PET = '0A6A70'; INK = '10272B'
F_H = PatternFill('solid', fgColor=PET); F_T = PatternFill('solid', fgColor='EEF3F3'); F_IN = PatternFill('solid', fgColor='FFF2CC'); F_P = PatternFill('solid', fgColor='E3EEF9'); F_C = PatternFill('solid', fgColor='E2F3EB')
F_AT = PatternFill('solid', fgColor='FBE8E5'); F_PE = PatternFill('solid', fgColor='FCEFD6'); F_OK = PatternFill('solid', fgColor='E2F3EB'); F_BAND = PatternFill('solid', fgColor='0C2A2F')
th = Side(style='thin', color='D9E3E3'); BD = Border(bottom=th)
FN = Font(name='Calibri', size=10, color=INK); FB = Font(name='Calibri', size=10, bold=True, color=INK); FH = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
FT = Font(name='Calibri', size=16, bold=True, color='FFFFFF'); FS = Font(name='Calibri', size=10, color='CFE0E0'); FNOTE = Font(name='Calibri', size=9, italic=True, color='5B6E72')
FBLUE = Font(name='Calibri', size=10, bold=True, color='1F5FAF'); FGRN = Font(name='Calibri', size=10, bold=True, color='0F7F59'); FIN = Font(name='Calibri', size=10, color='0000FF')
NUM = '#,##0.00;[Red]-#,##0.00;"–"'; NUM0 = '#,##0;[Red]-#,##0;"–"'; PCT = '0.0%;[Red]-0.0%;"–"'; PCT2 = '0.00%'
def sheet(name, title, sub, widths, tab=PET, landscape=False):
    ws = wb.create_sheet(name); ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = tab
    for i, w in enumerate(widths, start=1): ws.column_dimensions[CL(i)].width = w
    n = len(widths)
    for c in range(1, n + 1):
        ws.cell(1, c).fill = F_BAND; ws.cell(2, c).fill = F_BAND
    ws.cell(1, 1, title).font = FT; ws.cell(2, 1, sub).font = FS; ws.row_dimensions[1].height = 26; ws.row_dimensions[2].height = 18
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True; ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.left.text = 'Clínica Núcleo S · Parceria de Tricologia · ' + DATA_REF; ws.oddFooter.right.text = 'Página &P de &N'
    return ws
def hdr(ws, r, vals, c0=1):
    for i, v in enumerate(vals):
        c = ws.cell(r, c0 + i, v); c.fill = F_H; c.font = FH; c.alignment = Alignment(horizontal='center' if i else 'left', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 30
def put(ws, r, c, v, fmt=None, font=FN, fill=None, wrap=False, al=None):
    x = ws.cell(r, c, v); x.font = font; x.border = BD
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    if wrap or al: x.alignment = Alignment(wrap_text=wrap, vertical='top', horizontal=al)
    return x
def section(ws, r, text):
    x = ws.cell(r, 1, text); x.font = Font(name='Calibri', size=12, bold=True, color=PET); return r + 1
KF = {'t': (FB, F_T), 'p': (FBLUE, F_P), 'c': (FGRN, F_C), '': (FN, None)}
# Resumo
ws = wb.active; ws.title = 'Resumo'; ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = '0C2A2F'
for i, w in enumerate([4, 46, 22, 60], start=1): ws.column_dimensions[CL(i)].width = w
for r in range(1, 5):
    for c in range(1, 5): ws.cell(r, c).fill = F_BAND
ws.cell(2, 2, 'Parceria de Tricologia — relatório dos painéis').font = Font(name='Calibri', size=20, bold=True, color='FFFFFF')
ws.cell(3, 2, f'Clínica Núcleo S × Dra. Patrícia Fabrini · exportado do painel em {DATA_REF}').font = FS; ws.row_dimensions[2].height = 32
r = 6; r = section(ws, r, 'Outubro/2026 · Cruzeiro do Sul (05 a 07/10)')
for lab, v, f in [('Faturamento', d['rec'], FB), ('Resultado', d['res'], FB), ('Dra. Patrícia · 50%', d['pat'], FBLUE), ('Clínica Núcleo S · total (50% + estrutura + amortização)', d['clitot'], FGRN)]:
    put(ws, r, 2, lab); put(ws, r, 3, v, NUM, f); r += 1
r += 1; r = section(ws, r, 'Abas desta pasta')
for a, b in [('Out26 · DRE CZS', 'DRE da parceria em Cruzeiro do Sul, linha a linha, com % do faturamento'), ('Out26 · Pacientes', 'Consultas e procedimentos por paciente'),
             ('Out26 · Lançamentos', 'Os 40 lançamentos do relatório + 2 ajustes da diretoria, com alertas'), ('Out26 · Viagem', 'Passagens da fatura e rateio Cruzeiro do Sul / Rio Branco'),
             ('Out26 · Auditoria', 'Verificações feitas antes da entrega'), ('Base de cálculo', 'Os 9 passos, de onde vem cada dedução'), ('Taxa de estrutura', 'Cálculo dos 18,46%, composição e mês a mês'),
             ('Calculadora', 'Digite o faturamento e veja o valor de cada um (fórmulas)'), ('Histórico jul–set', 'A conta aplicada aos meses de teste, sem partilha'),
             ('Investimento', 'Aparelhos, consultoria e amortização'), ('Regras e pendências', 'O que entra, o que não entra e o que falta decidir')]:
    put(ws, r, 2, a, font=FB); put(ws, r, 4, b); r += 1
r += 1; ws.cell(r, 2, 'Valores exportados do painel (fotografia de ' + DATA_REF + '). A planilha de fechamento com fórmulas e conferência é a "Fechamento_Parceria_Tricologia_Out2026_CZS.xlsx".').font = FNOTE
# DRE CZS
ws = sheet('Out26 · DRE CZS', 'DRE da parceria · Cruzeiro do Sul · 05 a 07/10/2026', f'{O["n_sis"]} lançamentos do relatório + {O["n_adj"]} ajustes da diretoria · consultas pelo valor total · antecipação de 8,74% no crédito', [62, 18, 16])
hdr(ws, 4, ['Linha', 'R$', '% do faturamento']); r = 5
for lab, v, k in DRE_OUT:
    f, fl = KF.get(k, (FN, None)); put(ws, r, 1, lab, font=f, fill=fl); put(ws, r, 2, v, NUM, f, fl); put(ws, r, 3, v / d['rec'], PCT2, f, fl); r += 1
r += 1; r = section(ws, r, 'Como o dinheiro entra')
for lab, v in [('PIX e dinheiro', d['vista']), ('Cartão de crédito à vista (antecipado)', d['cartvista']), ('Cartão de crédito parcelado (antecipado)', d['parc']),
               ('1ª parte das consultas (recebida em ago/set)', d['cmp']), ('Ajustes da diretoria (pescoço + saldo do corporal)', d['adj'])]:
    put(ws, r, 1, lab); put(ws, r, 2, v, NUM); put(ws, r, 3, v / d['rec'], PCT); r += 1
put(ws, r, 1, 'Total', font=FB, fill=F_T); put(ws, r, 2, d['vista'] + d['cartvista'] + d['parc'] + d['cmp'] + d['adj'], NUM, FB, F_T); put(ws, r, 3, 1, PCT, FB, F_T)
ws.freeze_panes = 'A5'
# Pacientes
ws = sheet('Out26 · Pacientes', 'Pacientes atendidas em Cruzeiro do Sul · 05 a 07/10/2026', 'Consultas pelo valor total (R$ 800 / R$ 900) separadas dos procedimentos', [38, 14, 14, 14, 14, 14, 15, 10, 34], landscape=True)
hdr(ws, 4, ['Paciente', 'Consulta', 'Tricologia', 'Facial', 'Pescoço', 'Corporal', 'Total', 'Lançamentos', 'Forma de pagamento']); r = 5
for n, p in PAC_ORD:
    put(ws, r, 1, n, font=FB)
    for c, k in enumerate(['cons', 'tri', 'fac', 'pes', 'cor', 'tot'], start=2): put(ws, r, c, p[k], NUM, FB if k == 'tot' else FN)
    put(ws, r, 8, p['n'], NUM0); put(ws, r, 9, ' + '.join(sorted(p['met']))); r += 1
put(ws, r, 1, 'Total', font=FB, fill=F_T)
for c in range(2, 9): put(ws, r, c, f'=SUM({CL(c)}5:{CL(c)}{r-1})', NUM if c < 8 else NUM0, FB, F_T)
ws.freeze_panes = 'B5'
# Lançamentos
ws = sheet('Out26 · Lançamentos', 'Lançamentos de Cruzeiro do Sul · 05 a 07/10/2026', '40 do relatório de transações + 2 ajustes da diretoria (em amarelo)', [34, 34, 12, 12, 30, 18, 13, 13, 14, 13, 70], landscape=True)
hdr(ws, 4, ['Paciente', 'Lançamento', 'Vencimento', 'Baixa', 'Forma de pagamento', 'Conta', 'Sistema', 'Ajuste', 'Considerado', 'Líquido no sistema', 'Alerta da auditoria']); r = 5
for n, _ in PAC_ORD:
    for l in sorted([x for x in LINES if x['p'] == n], key=lambda x: (x['t'], x['srv'], x['desc'])):
        fl = F_IN if l.get('man') else (F_P if l['t'] == 'Consulta' else None)
        vals = [n, parc(l), l['ve'], l['bx'], l['met'], l['cta'], l['vs'], l['cmp'], l['v'], l['liq'], l['a'] or None]
        for c, v in enumerate(vals, start=1): put(ws, r, c, v, NUM if c in (7, 8, 9, 10) else None, FB if c == 9 else FN, fl, wrap=(c == 11))
        r += 1
put(ws, r, 1, 'Total', font=FB, fill=F_T)
for c in (7, 8, 9, 10): put(ws, r, c, f'=SUM({CL(c)}5:{CL(c)}{r-1})', NUM, FB, F_T)
ws.freeze_panes = 'C5'; ws.auto_filter.ref = f'A4:K{r-1}'
# Viagem
ws = sheet('Out26 · Viagem', 'Viagem SP → Cruzeiro do Sul → Rio Branco → SP', 'Passagens pagas no cartão da clínica (fatura de 05/10/2026) · rateio por dias de atendimento', [14, 62, 16])
hdr(ws, 4, ['Data', 'Descrição na fatura', 'Valor']); r = 5; p0 = r
for x in O['pass_']: put(ws, r, 1, x[0]); put(ws, r, 2, x[1] + (' · ' + x[3] if x[3] else '')); put(ws, r, 3, x[4], NUM); r += 1
put(ws, r, 2, 'Total das passagens na fatura de 05/10', font=FB, fill=F_T); put(ws, r, 3, f'=SUM(C{p0}:C{r-1})', NUM, FB, F_T); rt = r; r += 1
put(ws, r, 2, f'→ Cruzeiro do Sul: {pc(RAT,0)} (3 de 4 dias: 05, 06 e 07/10)'); put(ws, r, 3, f'=ROUND(C{rt}*{RAT},2)', NUM); r += 1
put(ws, r, 2, f'→ Rio Branco: {pc(1-RAT,0)} (1 de 4 dias: 08/10)'); put(ws, r, 3, f'=C{rt}-C{r-1}', NUM); r += 1
put(ws, r, 2, 'Conferência: Cruzeiro do Sul + Rio Branco − fatura (deve ser zero)', font=FB); put(ws, r, 3, f'=ROUND(C{r-2}+C{r-1}-C{rt},2)', NUM, FB); r += 2
put(ws, r, 2, O['pass2'][1].replace(', vem na fatura de nov/26', '') + ' — fica para o fechamento de novembro', font=FNOTE); put(ws, r, 3, O['pass2'][4], NUM, FNOTE)
# Auditoria
ws = sheet('Out26 · Auditoria', 'Auditoria feita antes da entrega', 'O que foi conferido e o que ainda depende de confirmação', [5, 58, 13, 90], landscape=True)
hdr(ws, 4, ['#', 'Verificação', 'Status', 'Detalhe']); r = 5
for i, a in enumerate(O['audit'], start=1):
    put(ws, r, 1, i); put(ws, r, 2, a[0], font=FB, wrap=True); put(ws, r, 3, a[1], font=FB, fill={'OK': F_OK, 'ATENÇÃO': F_AT, 'PENDENTE': F_PE}[a[1]], al='center'); put(ws, r, 4, a[2], wrap=True)
    ws.row_dimensions[r].height = 15 * max(1, len(a[2]) // 95 + 1); r += 1
# Base de cálculo
ws = sheet('Base de cálculo', 'Base de cálculo da parceria', 'Do faturamento ao valor de cada um · exemplo com R$ 100.000 de procedimentos e índices médios', [34, 30, 16, 60, 60], landscape=True)
hdr(ws, 4, ['Passo', 'Índice', 'No exemplo (R$)', 'O que é', 'De onde vem']); r = 5
for s in STEPS:
    put(ws, r, 1, f'{s[0]}. {s[1]}', font=FB, wrap=True); put(ws, r, 2, s[2], wrap=True); put(ws, r, 3, s[3] if s[0] == '9' else -s[3], NUM, FB)
    put(ws, r, 4, s[4], wrap=True); put(ws, r, 5, s[5], wrap=True); ws.row_dimensions[r].height = 42; r += 1
r += 1
for lab, v, f in [('Dra. Patrícia recebe', EX['pat'], FBLUE), ('Clínica fica com (50% + estrutura + investimento)', EX['clitot'], FGRN)]:
    put(ws, r, 1, lab, font=f); put(ws, r, 3, v, NUM, f); put(ws, r, 4, v / EX['rec'], PCT, f); r += 1
# Taxa de estrutura
ws = sheet('Taxa de estrutura', f'Taxa de estrutura · {pc(EST_PCT,2)}', f'{brl(EST_MES)} de estrutura por mês ÷ {brl(ROB_MES)} de faturamento médio (jul–set/2026)', [52, 15, 12, 15, 17, 70] + [12] * 4, landscape=True)
hdr(ws, 4, ['Despesa da clínica (média por mês)', 'R$ por mês', '% fat.', 'Entra na taxa', 'Fica com a clínica', 'Por quê']); r = 5; e0 = r
for g, t, e in EST_ROWS:
    put(ws, r, 1, WHY.get(g, [g])[0], wrap=True); put(ws, r, 2, round(t, 2), NUM); put(ws, r, 3, f'=B{r}/{ROB_MES:.6f}', PCT); put(ws, r, 4, round(e, 2), NUM, FB if e > .5 else FN)
    put(ws, r, 5, f'=B{r}-D{r}', NUM); put(ws, r, 6, WHY.get(g, ['', ''])[1], wrap=True); ws.row_dimensions[r].height = 30; r += 1
put(ws, r, 1, 'Total', font=FB, fill=F_T)
for c in (2, 4, 5): put(ws, r, c, f'=SUM({CL(c)}{e0}:{CL(c)}{r-1})', NUM, FB, F_T)
put(ws, r, 3, f'=B{r}/{ROB_MES:.6f}', PCT, FB, F_T); rt = r; r += 1
put(ws, r, 1, 'Taxa de estrutura (entra na taxa ÷ faturamento médio)', font=FB, fill=F_T); put(ws, r, 4, f'=D{rt}/{ROB_MES:.6f}', PCT2, FB, F_T); r += 2
r = section(ws, r, 'Mês a mês'); hdr(ws, r, ['Mês', 'Estrutura (R$)', '', 'Receita bruta (R$)', 'Taxa do mês']); r += 1
for i in range(9):
    put(ws, r, 1, ML[i]); put(ws, r, 2, D['est_tot'][i], NUM); put(ws, r, 4, ROB[i], NUM); put(ws, r, 5, f'=B{r}/D{r}', PCT2); r += 1
# Calculadora (fórmulas)
ws = sheet('Calculadora', 'Calculadora da partilha', 'Amarelo = você digita. Os índices também podem ser alterados.', [52, 18, 14, 56])
r = 4; r = section(ws, r, 'Entradas')
put(ws, r, 1, 'Faturamento da parceria no mês (R$)', font=FB); c = put(ws, r, 2, 100000, NUM, FIN, F_IN); rIN = r; r += 1
put(ws, r, 1, 'dos quais: consultas (R$)'); put(ws, r, 2, 0, NUM, FIN, F_IN); rCO = r; r += 1
put(ws, r, 1, 'Vendas no cartão de crédito (R$) — vazio = taxa média sobre tudo'); put(ws, r, 2, None, NUM, FIN, F_IN); rCR = r; r += 2
r = section(ws, r, 'Índices'); IXR = {}
for k, lab, v, fmt in [('pis', 'PIS', IXP['pis'], PCT2), ('cof', 'COFINS', IXP['cof'], PCT2), ('iss', 'ISS', IXP['iss'], PCT2), ('cart', 'Cartão — média sobre o faturamento', IXP['cart'], PCT2),
                       ('ant', 'Antecipação — % sobre as vendas no crédito', 0.0874, PCT2), ('est', 'Taxa de estrutura', EST_PCT, PCT2), ('ins', 'Insumos (índice)', IXP['ins'], PCT2),
                       ('viag', 'Viagens (índice)', IXP['viag'], PCT2), ('out', 'Outras despesas diretas (índice)', IXP['out'], PCT2), ('irp', 'IRPJ + CSLL — procedimentos', IXP['irp'], PCT2),
                       ('irc', 'IRPJ + CSLL — consultas', IXP['irc'], PCT2), ('amort', 'Amortização mensal (R$)', IXP['amort'], NUM), ('pat', 'Parte da Dra. Patrícia', IXP['pat'], PCT)]:
    put(ws, r, 1, lab); put(ws, r, 2, v, fmt, FIN, F_IN); IXR[k] = f'$B${r}'; r += 1
r += 1; r = section(ws, r, 'Cálculo'); hdr(ws, r, ['Passo', 'R$', '% fat.', 'Como é calculado']); r += 1
REC_ = f'$B${rIN}'; CO_ = f'$B${rCO}'; CR_ = f'$B${rCR}'
CSTX = [('Faturamento da parceria', f'={REC_}', 't', 'Valor digitado'), ('(−) PIS', f'=-{REC_}*{IXR["pis"]}', '', '0,65% do faturamento'), ('(−) COFINS', f'=-{REC_}*{IXR["cof"]}', '', '3% do faturamento'),
        ('(−) ISS', f'=-{REC_}*{IXR["iss"]}', '', 'ISS fixo anual'), ('(−) Cartão e antecipação', f'=-IF({CR_}="",{REC_}*{IXR["cart"]},{CR_}*{IXR["ant"]})', '', 'Média de 2,2% sobre tudo, ou 8,74% sobre as vendas no crédito'),
        ('(=) Receita líquida', None, 't', ''), ('(−) Insumos e ativos', f'=-{REC_}*{IXR["ins"]}', '', 'Índice (no mês real: valor das notas)'), ('(−) Viagens e estadia', f'=-{REC_}*{IXR["viag"]}', '', 'Índice (no mês real: comprovantes)'),
        ('(−) Outras despesas diretas', f'=-{REC_}*{IXR["out"]}', '', 'Índice'), ('(−) Taxa de estrutura da clínica', f'=-{REC_}*{IXR["est"]}', '', '18,46% do faturamento'),
        ('(−) Amortização do investimento', f'=-IF({REC_}>0,{IXR["amort"]},0)', '', 'Parcela fixa mensal'), ('(−) IRPJ e CSLL', f'=-({CO_}*{IXR["irc"]}+({REC_}-{CO_})*{IXR["irp"]})', '', '7,68% das consultas + 2,28% dos procedimentos'),
        ('(=) Resultado do mês', None, 't', ''), ('Dra. Patrícia', None, 'p', '50% do resultado (zero se negativo)'), ('Clínica Núcleo S', None, 'c', '50% do resultado'),
        ('Clínica — total (50% + estrutura + amortização)', None, 'c', 'O que fica com a clínica')]
rows_at = {}
for lab, f, k, nt in CSTX:
    rows_at[lab] = r; ff, fl = KF.get(k, (FN, None)); put(ws, r, 1, lab, font=ff, fill=fl); put(ws, r, 4, nt, font=FNOTE)
    if lab == '(=) Receita líquida': f = f'=SUM(B{rows_at["Faturamento da parceria"]}:B{r-1})'
    if lab == '(=) Resultado do mês': f = f'=B{rows_at["(=) Receita líquida"]}+SUM(B{rows_at["(−) Insumos e ativos"]}:B{r-1})'
    if lab == 'Dra. Patrícia': f = f'=MAX(0,B{rows_at["(=) Resultado do mês"]})*{IXR["pat"]}'
    if lab == 'Clínica Núcleo S': f = f'=MAX(0,B{rows_at["(=) Resultado do mês"]})*(1-{IXR["pat"]})'
    if lab.startswith('Clínica — total'): f = f'=B{rows_at["Clínica Núcleo S"]}-B{rows_at["(−) Taxa de estrutura da clínica"]}-B{rows_at["(−) Amortização do investimento"]}'
    put(ws, r, 2, f, NUM, ff, fl); put(ws, r, 3, f'=IFERROR(B{r}/{REC_},0)', PCT2, ff, fl); r += 1
# Histórico
ws = sheet('Histórico jul–set', 'A conta aplicada aos meses de teste · jul a set/2026', 'Até setembro a receita da tricologia é 100% da clínica; a partilha vale a partir de outubro', [52, 14, 14, 14, 16, 14])
hdr(ws, 4, ['Linha', 'jul/26', 'ago/26', 'set/26', 'Acum. jul–set', '% da receita']); r = 5
for k, lab, t in HIST:
    ff, fl = KF.get(t, (FN, None)); put(ws, r, 1, lab, font=ff, fill=fl)
    for c, v in enumerate(P[k][1:5], start=2): put(ws, r, c, v, NUM, ff, fl)
    put(ws, r, 6, P[k][4] / P['12'][4], PCT, ff, fl); r += 1
# Investimento
ws = sheet('Investimento', 'Investimento de implantação', 'Pago 100% pela clínica antes do primeiro atendimento; volta em 24 parcelas dentro da conta da parceria', [62, 18])
hdr(ws, 4, ['Item', 'R$']); r = 5
for lab, v, k in [('Aparelhos (eletroporação, LED, câmera, microscópio), líquidos do estorno', 21139.82, ''), ('Consultoria de implantação: 6 parcelas pagas (abr–set)', 34285.68, ''),
                  ('Consultoria de implantação: 7ª parcela (out/26)', 5714.28, '')]:
    put(ws, r, 1, lab); put(ws, r, 2, v, NUM); r += 1
put(ws, r, 1, 'Investimento total', font=FB, fill=F_T); put(ws, r, 2, '=SUM(B5:B7)', NUM, FB, F_T); r += 1
put(ws, r, 1, 'Amortização mensal (24 × de out/26 a set/28)'); put(ws, r, 2, f'=B{r-1}/24', NUM); r += 2
r = section(ws, r, 'Cronograma'); hdr(ws, r, ['Parcela · competência', 'Saldo a recuperar (R$)']); r += 1
yy, mm = 2026, 10
for n in range(1, 25):
    put(ws, r, 1, f'{n}/24 · {["jan","fev","mar","abr","mai","jun","jul","ago","set","out","nov","dez"][mm-1]}/{str(yy)[2:]}'); put(ws, r, 2, round(61139.78 - 61139.78 / 24 * n, 2), NUM); r += 1
    mm += 1
    if mm > 12: mm = 1; yy += 1
# Regras e pendências
ws = sheet('Regras e pendências', 'Regras e pendências da parceria', 'A partir das receitas de outubro/2026, o resultado da parceria é dividido 50/50', [70, 70])
hdr(ws, 4, ['Entra na conta (desconta antes da divisão)', 'Não entra (fica com a clínica)']); r = 5
for a, b in zip(ENTRA, NAO): put(ws, r, 1, '✓ ' + a, wrap=True); put(ws, r, 2, '✕ ' + b, wrap=True); ws.row_dimensions[r].height = 28; r += 1
r += 1; r = section(ws, r, 'Como funciona todo mês')
for i, t in enumerate(OPER, start=1): put(ws, r, 1, f'{i}. {t}', wrap=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2); ws.row_dimensions[r].height = 28; r += 1
r += 1; r = section(ws, r, 'Pendências antes de apresentar'); hdr(ws, r, ['Pendência', 'Situação']); r += 1
for a, b, c in PEN: put(ws, r, 1, a, wrap=True); put(ws, r, 2, c, font=FB, fill={'OK': F_OK, 'ATENÇÃO': F_AT, 'PENDENTE': F_PE}[b]); ws.row_dimensions[r].height = 28; r += 1
xl = OUTD + f'/{NAME}.xlsx'; wb.save(xl); print('xlsx ok', os.path.getsize(xl))
