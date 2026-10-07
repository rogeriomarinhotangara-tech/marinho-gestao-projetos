import sys, json, datetime as dt, re
import pandas as pd
sys.path.insert(0, sys.argv[2])
import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
from xl_style import *
from rules2 import PALAVRAS, DEPARA_RECEITA, DEPARA_CATEGORIA, norm_desc
SP = sys.argv[1]; OUT = sys.argv[3]
base = pd.read_pickle(SP+'/ref/base_v2.pkl').reset_index(drop=True)
plano = pd.read_pickle(SP+'/ref/plano_v2.pkl')
mem = json.load(open(SP+'/ref/memoria_v2.json'))
NROW = 6001
wb = openpyxl.Workbook()
SHEETS = ['00 Leia-me','01 Política Parceria','02 Premissas','03 DRE Tricologia','04 Rateio Estrutura','05 Custo Procedimento',
 '06 Investimento Tricologia','07 Razão Tricologia','08 DRE Clínica','09 Conciliação Datas','10 Jul Receitas','11 Jul Despesas',
 '12 Ago Receitas','13 Ago Despesas','14 Set Receitas','15 Set Despesas','16 Auditoria e Alertas','17 Decisões Diretoria',
 '18 De-Para','19 Plano de Contas','20 Base']
wb.active.title = SHEETS[0]
for s in SHEETS[1:]: wb.create_sheet(s)
W = {s: wb[s] for s in SHEETS}
def q(s): return "'"+s+"'"
def name(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)
MESES = [202601,202602,202603,202604,202605,202606,202607,202608,202609]
MLAB = ['Jan/26','Fev/26','Mar/26','Abr/26','Mai/26','Jun/26','Jul/26','Ago/26','Set/26']

# ============================================================ 19 Plano de Contas
ws = W['19 Plano de Contas']; title(ws,'Plano de Contas Financeiro-Gerencial — Clínica Núcleo (oficial)',
 'Fonte: "Plano de Contas Nucleo.xlsx" (Google Drive, 18/08/2026). Em amarelo: contas criadas nesta reestruturação. Colunas F e G: como cada conta entra na DRE da Tricologia e na base de rateio.')
hdr(ws,4,1,['Código','Conta','Grupo DRE','Natureza','Subgrupo','Linha na DRE da Tricologia','Entra na base de estrutura?','Tipo','Origem','Regra de uso (plano oficial)'],
    [13,50,34,16,40,34,14,11,24,70])
r=5
for _,p in plano.iterrows():
    vals=[p.cod,p.nome,p.grupo,p.natureza,p.subgrupo,p.linha_trico,p.estrutura,p.registro,p.origem,p.obs]
    for j,v in enumerate(vals,start=1):
        c=put(ws,r,j,v,border=B_ALL)
        if p.criada: c.fill=F_IN
        elif p.registro!='Analítica': c.font=FT_B
    r+=1
name('PC_TAB', f"{q('19 Plano de Contas')}!$A$5:$J$1500"); ws.freeze_panes='C5'; ws.auto_filter.ref=f"A4:J{r-1}"

# ============================================================ 18 De-Para
ws = W['18 De-Para']; title(ws,'De-Para de classificação',
 'Ordem: (1) código manual → (2) transferência = passagem → (3) receita: consulta da Dra. Patrícia, serviço, tricologia → (4) Dr. Marcos: MOD = honorário; sem MOD = retirada (pró-labore fica em Pessoal) → (5) memória → (6) palavra-chave → (7) categoria → (8) não classificada.')
hdr(ws,4,1,['Serviço (1º da lista)','Código'],[30,13]); r=5
for k,v in DEPARA_RECEITA.items(): put(ws,r,1,k,border=B_ALL); put(ws,r,2,v,border=B_ALL); r+=1
name('DP_REC', f"{q('18 De-Para')}!$A$5:$B$300")
hdr(ws,4,4,['Categoria do sistema (despesa)','Código'],[46,13]); r=5
for k,v in DEPARA_CATEGORIA.items(): put(ws,r,4,k,border=B_ALL); put(ws,r,5,v,border=B_ALL); r+=1
name('DP_CAT', f"{q('18 De-Para')}!$D$5:$E$300")
hdr(ws,4,7,['Palavra-chave na descrição','Código','Motivo'],[26,13,58]); r=5
for k,v,m in PALAVRAS: put(ws,r,7,k,border=B_ALL); put(ws,r,8,v,border=B_ALL); put(ws,r,9,m,border=B_ALL); r+=1
KW_LAST=r-1
put(ws,r+1,7,'A última palavra encontrada vence. Para incluir, INSIRA uma linha dentro da tabela (acima da última).',font=FT_NOTE)
name('DP_KW_K', f"{q('18 De-Para')}!$G$5:$G${KW_LAST}"); name('DP_KW_C', f"{q('18 De-Para')}!$H$5:$H${KW_LAST}")
hdr(ws,4,11,['Memória: descrição já classificada (jan–jun)','Código'],[52,13]); r=5
for k in sorted(mem): put(ws,r,11,k,border=B_ALL); put(ws,r,12,mem[k],border=B_ALL); r+=1
name('DP_MEM', f"{q('18 De-Para')}!$K$5:$L$3000"); ws.freeze_panes='A5'

# ============================================================ 20 Base
ws = W['20 Base']
COLS = ['R/D','Data de emissão','Data de vencimento','Data de baixa','Responsável','Paciente','Descrição','Serviços','Categorias','Nota fiscal',
        'Convênio','Método','Conta','Valor','Valor líquido','Agendado','Pago','Observações']
CALC = ['Origem','Código manual','Mês (venc.)','Código auto','Código final','Conta analítica','Grupo DRE','Natureza','Unidade',
        'Tricologia (auto)','Tricologia (manual)','Tricologia','Linha DRE Tricologia','Base de estrutura','Alertas','Mês (baixa)','Nota de ajuste']
hdr(ws,1,1,COLS+CALC,None)
for i,w in enumerate([9,11,11,11,16,18,46,24,30,8,10,20,30,12,12,8,8,30, 30,13,9,13,13,40,30,14,15,10,10,10,30,10,30,9,60],start=1): ws.column_dimensions[L(i)].width=w
for c in range(19,36): ws.cell(1,c).fill=PatternFill('solid',fgColor='48545E')
for c in (20,29): ws.cell(1,c).fill=PatternFill('solid',fgColor='B08900')
for i,row in base.iterrows():
    rr=i+2
    for j,cn in enumerate(COLS,start=1):
        v=row[cn]
        if v is None or (isinstance(v,float) and v!=v): continue
        if isinstance(v,pd.Timestamp): v=v.date()
        c=ws.cell(rr,j,v)
        if j in (2,3,4): c.number_format=DATE
        if j in (14,15): c.number_format=NUM
    ws.cell(rr,19,row['Origem'])
    if isinstance(row['Código manual'],str) and row['Código manual']: ws.cell(rr,20,row['Código manual'])
    if isinstance(row['Tricologia manual'],str) and row['Tricologia manual']: ws.cell(rr,29,row['Tricologia manual'])
    if isinstance(row['Nota de ajuste'],str): ws.cell(rr,35,row['Nota de ajuste'])
for rr in range(2,NROW+1):
    kw=f'UPPER(TRIM(SUBSTITUTE($G{rr},"Transação Recorrente:","")))'
    first=f'TRIM(IFERROR(LEFT($H{rr},FIND(",",$H{rr})-1),$H{rr}))'
    F={
    21:f'=IF(OR($A{rr}="",$C{rr}=""),"",YEAR($C{rr})*100+MONTH($C{rr}))',
    22:(f'=IF($A{rr}="","",IF(OR(LEFT($G{rr},13)="Transferência",ISNUMBER(SEARCH("TRANSFERÊNCIA",$G{rr})),ISNUMBER(SEARCH("TRANSFERENCIA",$G{rr}))),"3.01.01.003",'
        f'IF($A{rr}="Receita",IF(AND($AB{rr}="SIM",OR({first}="CONSULTA",ISNUMBER(SEARCH("CONSULTA DRA PATRICIA",$H{rr})))),"1.01.03.001",'
        f'IFERROR(VLOOKUP({first},DP_REC,2,0),IF($AB{rr}="SIM","1.01.03.010","1.01.99.000"))),'
        f'IF(AND(OR(ISNUMBER(SEARCH("DR MARCOS",$G{rr})),ISNUMBER(SEARCH("MARCOS SANTANA",$G{rr})),TRIM($I{rr})="DR MARCOS SANTANA"),NOT(ISNUMBER(SEARCH("LABORE",$G{rr})))),'
        f'IF(ISNUMBER(SEARCH("MOD",$G{rr})),"2.02.04.001","2.18.01.003"),'
        f'IFERROR(VLOOKUP({kw},DP_MEM,2,0),IFERROR(LOOKUP(2^15,SEARCH(DP_KW_K,{kw}),DP_KW_C),IFERROR(VLOOKUP(TRIM($I{rr}),DP_CAT,2,0),"2.13.01.090")))))))'),
    23:f'=IF($A{rr}="","",IF($T{rr}<>"",$T{rr},$V{rr}))',
    24:f'=IF($W{rr}="","",IFERROR(VLOOKUP($W{rr},PC_TAB,2,0),"CÓDIGO FORA DO PLANO"))',
    25:f'=IF($W{rr}="","",IFERROR(VLOOKUP($W{rr},PC_TAB,3,0),"CÓDIGO FORA DO PLANO"))',
    26:f'=IF($W{rr}="","",IFERROR(VLOOKUP($W{rr},PC_TAB,4,0),""))',
    27:(f'=IF($A{rr}="","",IF(OR(ISNUMBER(SEARCH("CZS",$G{rr})),ISNUMBER(SEARCH("CRZ",$G{rr})),ISNUMBER(SEARCH("CRUZEIRO DO SUL",$G{rr})),ISNUMBER(SEARCH("CRZ",$H{rr}))),"Cruzeiro do Sul",'
        f'IF(OR(ISNUMBER(SEARCH("EPITACIO",$G{rr})),ISNUMBER(SEARCH("EPITACIO",$H{rr}))),"Epitaciolândia",IF(ISNUMBER(SEARCH("CRUZEIRO",$M{rr})),"Cruzeiro do Sul",'
        f'IF(ISNUMBER(SEARCH("EPITACIO",$M{rr})),"Epitaciolândia",IF($M{rr}="","Não alocado","Rio Branco"))))))'),
    28:(f'=IF($A{rr}="","",IF(OR(ISNUMBER(SEARCH("TRICOLOG",$G{rr})),ISNUMBER(SEARCH("TROCOLOG",$G{rr})),ISNUMBER(SEARCH("CAPILAR",$G{rr})),'
        f'ISNUMBER(SEARCH("PATRICIA",$E{rr})),ISNUMBER(SEARCH("PATRICIA",$H{rr})),ISNUMBER(SEARCH("TRICOLOG",$H{rr})),ISNUMBER(SEARCH("CAPILAR",$H{rr})),'
        f'ISNUMBER(SEARCH("BARBA",$H{rr})),ISNUMBER(SEARCH("FOTOTERAPIA",$H{rr})),ISNUMBER(SEARCH("ELETROPORA",$H{rr})),'
        f'AND($A{rr}="Despesa",ISNUMBER(SEARCH("PATRICIA",$G{rr})))),"SIM","NÃO"))'),
    30:f'=IF($A{rr}="","",IF($AC{rr}<>"",$AC{rr},$AB{rr}))',
    31:f'=IF($AD{rr}="SIM",IFERROR(VLOOKUP($W{rr},PC_TAB,6,0),"D3 Outras despesas diretas"),"")',
    32:f'=IF($A{rr}="","",IF(AND($A{rr}="Despesa",$AD{rr}<>"SIM"),IF(IFERROR(VLOOKUP($W{rr},PC_TAB,7,0),"NÃO")="SIM","SIM","NÃO"),"NÃO"))',
    33:(f'=IF($A{rr}="","",TRIM(IF($W{rr}="2.13.01.099","CARTÃO SEM DETALHE; ","")&IF(AND($A{rr}="Despesa",$D{rr}=""),"SEM BAIXA; ","")'
        f'&IF(AND($A{rr}="Despesa",$N{rr}>=5000,$R{rr}=""),"≥5 MIL SEM OBS.; ","")&IF(LEFT($W{rr},4)="2.17","JUROS NÃO SEGREGADOS; ","")'
        f'&IF(OR($W{rr}="2.13.01.090",$W{rr}="1.01.99.000"),"NÃO CLASSIFICADA; ","")&IF($W{rr}="2.18.01.003","RETIRADA DE SÓCIO; ","")&IF(LEFT($W{rr},2)="3.","PASSAGEM (FORA DA DRE); ","")))'),
    34:f'=IF(OR($A{rr}="",$D{rr}=""),"",YEAR($D{rr})*100+MONTH($D{rr}))',
    }
    for c,v in F.items(): ws.cell(rr,c,v)
    ws.cell(rr,20).fill=F_IN; ws.cell(rr,29).fill=F_IN
ws.freeze_panes='H2'; ws.auto_filter.ref=f"A1:AI{NROW}"
B=q('20 Base')
for n,col in [('B_TIPO','A'),('B_BX','D'),('B_MET','L'),('B_CTA','M'),('B_VAL','N'),('B_MES','U'),('B_COD','W'),('B_GRP','Y'),('B_UNI','AA'),
              ('B_TRI','AD'),('B_TLIN','AE'),('B_EST','AF'),('B_MESB','AH')]:
    name(n, f"{B}!${col}$2:${col}${NROW}")
NDATA=len(base)+1
exec(open(sys.argv[2]+'/build2_part_b.py',encoding='utf-8').read())
exec(open(sys.argv[2]+'/build2_part_c.py',encoding='utf-8').read())
wb.active=0
wb.save(OUT)
base.to_pickle(SP+'/ref/base_v2_out.pkl')
print('ok', NDATA)
