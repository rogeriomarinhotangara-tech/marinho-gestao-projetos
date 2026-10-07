from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
NAVY='1F3A5F'; BLUE='2A78D6'; ORANGE='EB6834'; GREEN='1BAF7A'
F_HEAD = PatternFill('solid', fgColor=NAVY)
F_SUB = PatternFill('solid', fgColor='DCE6F2')
F_TOT = PatternFill('solid', fgColor='EEF1F4')
F_IN = PatternFill('solid', fgColor='FFF2CC')
F_WARN = PatternFill('solid', fgColor='FBE7E5')
F_OK = PatternFill('solid', fgColor='E2F2EB')
F_PART = PatternFill('solid', fgColor='FDF3DF')
FT_HEAD = Font(name='Calibri', bold=True, color='FFFFFF', size=10)
FT_B = Font(name='Calibri', bold=True, size=10)
FT_N = Font(name='Calibri', size=10)
FT_IN = Font(name='Calibri', size=10, color='0000FF')
FT_T1 = Font(name='Calibri', bold=True, size=16, color=NAVY)
FT_T2 = Font(name='Calibri', bold=True, size=12, color=NAVY)
FT_NOTE = Font(name='Calibri', italic=True, size=9, color='5B6770')
FT_RED = Font(name='Calibri', bold=True, size=10, color='B3382F')
FT_GRN = Font(name='Calibri', bold=True, size=10, color='127A55')
THIN = Side(style='thin', color='C8D1D8'); MED = Side(style='medium', color=NAVY)
B_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
B_TOP = Border(top=MED)
NUM = '#,##0.00;[Red]-#,##0.00;"–"'
NUM0 = '#,##0;[Red]-#,##0;"–"'
PCT = '0.0%;[Red]-0.0%;"–"'
PCT2 = '0.00%'
DATE = 'DD/MM/YYYY'
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

def hdr(ws, row, col, values, widths=None, fill=F_HEAD, font=FT_HEAD, height=30):
    for i,v in enumerate(values):
        c = ws.cell(row, col+i, v); c.fill = fill; c.font = font; c.alignment = CENTER; c.border = B_ALL
    ws.row_dimensions[row].height = height
    if widths:
        for i,w in enumerate(widths): ws.column_dimensions[get_column_letter(col+i)].width = w

def put(ws, row, col, v, fmt=None, font=FT_N, fill=None, align=None, border=None):
    c = ws.cell(row, col, v)
    c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align: c.alignment = align
    if border: c.border = border
    return c

def inp(ws, row, col, v, fmt=None):
    return put(ws, row, col, v, fmt=fmt, font=FT_IN, fill=F_IN, border=B_ALL)

def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = FT_T1
    if sub: ws['A2'] = sub; ws['A2'].font = FT_NOTE
    ws.sheet_view.showGridLines = False
