"""Planilha simples para a Dra. Patrícia: fechamento de Rio Branco (08/10/2026), no mesmo modelo de Cruzeiro do Sul.
Uso: python3 -I build_simples_rb.py <scratch> <saida.xlsx>"""
import sys, json
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.datavalidation import DataValidation
SP, OUT = sys.argv[1], sys.argv[2]
up = lambda s: s.astype(str).str.upper()
FMTBR = lambda v: f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

# ---------------- receitas de 08/10/2026 (relatório de transações de Rio Branco)
TIPOS = [('cons', 'Consultas', 'FAT_CONS'), ('tri', 'Tricologia', 'FAT_TRI'), ('fac', 'Facial', 'FAT_FAC'), ('pes', 'Pescoço', 'FAT_PES'), ('cor', 'Corporal', 'FAT_COR')]
SRVK = {'CONSULTA DRA PATRICIA FABRINI': 'cons', 'TRICOLOGIA': 'tri', 'FACIAL DRA PATRICIA': 'fac', 'PESCOÇO - DR PATRICIA FABRINI': 'pes', 'PESCOÇO DRA PATRICIA': 'pes', 'CORPORAL DRA PATRICIA': 'cor'}
# lançamentos com mais de um serviço: separados pelos detalhes da transação de 08/10 (facial 8.000 + tricologia 10.000; corporal, facial e pescoço 7.200 cada)
SPLIT = {frozenset({'fac', 'tri'}): {'fac': 8000 / 18000, 'tri': 10000 / 18000}, frozenset({'cor', 'fac', 'pes'}): {'cor': 1 / 3, 'fac': 1 / 3, 'pes': 1 / 3}}
RX = pd.read_excel(SP + '/data7/rec_rb_0810.xlsx'); RX = RX[RX['Valor'].notna() & RX['Paciente'].notna()].copy()
PAC = {}; V_ABERTO = 0.0; V_ANT = 0.0; N_LANC = len(RX)
for _, x in RX.iterrows():
    srv = [s.strip() for s in str(x['Serviços']).split(',')]
    ob = str(x['Observações']) if pd.notna(x['Observações']) else ''
    ks = (['pes'] if 'pesco' in ob.lower() else []) if srv == ['ELETROPORAÇÃO'] else [SRVK[s] for s in srv]
    assert ks, (srv, ob)
    v = float(x['Valor']); sh = {ks[0]: 1.0} if len(ks) == 1 else SPLIT[frozenset(ks)]
    p = PAC.setdefault(x['Paciente'], dict.fromkeys([t[0] for t in TIPOS] + ['cred'], 0.0))
    for k, f in sh.items(): p[k] += v * f
    met = str(x['Método']).lower()
    if 'cr' in met and 'dito' in met: p['cred'] += v
    if str(x['Pago']) == 'Não': V_ABERTO += v
    if pd.Timestamp(x['Data de vencimento']).month == 9: V_ANT += v
for p in PAC.values():
    for k in p: p[k] = round(p[k], 2)
TOTP = lambda p: sum(p[t[0]] for t in TIPOS)
PAC_ORD = sorted(PAC.items(), key=lambda x: -TOTP(x[1]))
SUMK = lambda k: round(sum(p[k] for p in PAC.values()), 2)
assert abs(sum(map(TOTP, PAC.values())) - 61300) < 0.01 and SUMK('cred') == 51700 and SUMK('cons') == 2500 and SUMK('tri') == 21200
assert SUMK('fac') == 15200 and SUMK('pes') == 15200 and SUMK('cor') == 7200 and V_ABERTO == 8000 and V_ANT == 800 and N_LANC == 19

# ---------------- custo fixo de Rio Branco: relatório de transações de setembro/2026 (mesmo relatório já classificado na base)
B = pd.read_pickle(SP + '/ref/base_v2_out.pkl'); SET = B[(B.Mes == 202609) & (B['R/D'] == 'Despesa')]
RAW = pd.read_excel(SP + '/data7/desp_set_rb.xlsx')
assert len(RAW) == len(SET) and abs(RAW.Valor.sum() - SET.Valor.sum()) < 0.01
S = SET[SET.Unidade == 'Rio Branco'].copy(); ca = S['Conta analítica']; dsc = up(S['Descrição'])
S['entra'] = (S.Estrutura == 'SIM') | ca.isin(['Combustível', 'Aviso-prévio e verbas rescisórias']) | ((ca == 'Manutenção corretiva de equipamento clínico') & dsc.str.contains('AR CONDICIONADO'))
S.loc[dsc.str.contains('DEPOSITO'), 'entra'] = False
FORA = S[~S.entra]; S = S[S.entra].copy()
RULES = [  # (conta analítica, palavra na descrição ou None, rótulo, grupo)
    ('Aluguel de imóvel', 'CONDOMINIO', 'Condomínio', 'Aluguel e condomínio'), ('Aluguel de imóvel', None, 'Aluguel do imóvel da clínica', 'Aluguel e condomínio'),
    ('Energia elétrica', 'CX 16', 'Energia elétrica (caixa 16)', 'Energia elétrica'), ('Energia elétrica', 'CX 17', 'Energia elétrica (caixa 17)', 'Energia elétrica'),
    ('Internet e links de dados', 'FRONTEIRAS', 'Internet (Sem Fronteiras)', 'Internet e telefone'), ('Internet e links de dados', 'TELEONE', 'Telefone (Vivo)', 'Internet e telefone'),
    ('Internet e links de dados', None, 'Internet e telefone (Vivo)', 'Internet e telefone'), ('Vigilância e segurança', None, 'Mensalidade do alarme', 'Vigilância e segurança'),
    ('Salários — equipe clínica e enfermagem', 'DIAS TRABALHADOS', 'Diárias — equipe', 'Equipe — salários'), ('Salários — equipe clínica e enfermagem', 'PAGAMENTO', 'Diárias — equipe', 'Equipe — salários'),
    ('Salários — equipe clínica e enfermagem', None, 'Salário — equipe', 'Equipe — salários'), ('Aviso-prévio e verbas rescisórias', None, 'Verbas rescisórias (despesa trabalhista)', 'Equipe — salários'),
    ('INSS patronal', None, 'INSS da equipe', 'Equipe — INSS e FGTS'), ('FGTS', None, 'FGTS da equipe', 'Equipe — INSS e FGTS'),
    ('Vale-refeição ou alimentação', 'AUX', 'Auxílio alimentação', 'Equipe — benefícios e saúde ocupacional'), ('Vale-refeição ou alimentação', None, 'Vale-alimentação da equipe', 'Equipe — benefícios e saúde ocupacional'),
    ('Seguro de vida', 'SANTANDER', 'Seguro (parcela Santander)', 'Equipe — benefícios e saúde ocupacional'), ('Seguro de vida', None, 'Seguro de vida da equipe', 'Equipe — benefícios e saúde ocupacional'),
    ('Combustível', None, 'Auxílio combustível', 'Equipe — benefícios e saúde ocupacional'), ('PCMSO, PGR, LTCAT e saúde ocupacional', None, 'Medicina do trabalho', 'Equipe — benefícios e saúde ocupacional'),
    ('Refeições administrativas', None, 'Refeição de reunião', 'Equipe — benefícios e saúde ocupacional'),
    ('Contabilidade', None, 'Contabilidade', 'Contabilidade, jurídico e sistemas'), ('Assessoria jurídica', None, 'Assessoria jurídica', 'Contabilidade, jurídico e sistemas'),
    ('Sistema de gestão clínica ou ERP', 'SEGUNDA PARTE', 'Sistema de gestão (2ª parte da implantação)', 'Contabilidade, jurídico e sistemas'),
    ('Sistema de gestão clínica ou ERP', None, 'Ponto eletrônico (mensalidade)', 'Contabilidade, jurídico e sistemas'),
    ('Alvará municipal e taxas de funcionamento', 'IPTU', 'IPTU', 'Taxas e IPTU'), ('Alvará municipal e taxas de funcionamento', None, 'Taxa de cancelamento de nota fiscal', 'Taxas e IPTU'),
    ('Coleta e destinação de resíduos de saúde', None, 'Coleta de lixo de saúde', 'Coleta de lixo de saúde'),
    ('Pequenos reparos e materiais de manutenção', 'TINTA', 'Tinta (manutenção)', 'Manutenção e reparos'), ('Pequenos reparos e materiais de manutenção', 'MATERIAL', 'Material de manutenção elétrica', 'Manutenção e reparos'),
    ('Pequenos reparos e materiais de manutenção', 'ELETRICA', 'Manutenção elétrica', 'Manutenção e reparos'), ('Pequenos reparos e materiais de manutenção', None, 'Mão de obra de manutenção', 'Manutenção e reparos'),
    ('Manutenção corretiva de equipamento clínico', None, 'Manutenção do ar-condicionado', 'Manutenção e reparos'), ('Manutenção de computadores e rede', None, 'Recarga de cartucho (bioimpedância)', 'Manutenção e reparos'),
    ('Material de escritório', None, 'Papelaria', 'Material de escritório e tarifas'), ('Tarifas bancárias', None, 'Tarifa bancária', 'Material de escritório e tarifas')]
GORD = list(dict.fromkeys(g for *_, g in RULES))
def rotulo(c, d):
    for cc, kw, lab, g in RULES:
        if cc == c and (kw is None or kw in d): return lab, g
    raise KeyError((c, d))
OBS = {'Aluguel do imóvel da clínica': 'Aluguel pago em setembro/2026', 'Condomínio': 'Pago em setembro/2026', 'Verbas rescisórias (despesa trabalhista)': 'Despesa trabalhista do mês',
       'Sistema de gestão (2ª parte da implantação)': 'Pagamento de setembro'}
RBL = []
for _, x in S.iterrows():
    lab, g = rotulo(x['Conta analítica'], str(x['Descrição']).upper())
    RBL.append(dict(d=pd.Timestamp(x['Data de vencimento']).strftime('%d/%m/%Y'), lab=lab, g=g, v=round(float(x['Valor']), 2), obs=OBS.get(lab, '')))
RBL.sort(key=lambda z: (GORD.index(z['g']), z['d'][3:5] + z['d'][:2], z['lab']))
CF_TOT = round(sum(z['v'] for z in RBL), 2)
assert len(RBL) == 44 and abs(CF_TOT - 64135.95) < 0.005, (len(RBL), CF_TOT)
assert any(z['lab'] == 'Aluguel do imóvel da clínica' and z['v'] == 11000.60 for z in RBL)
GRP_USED = [g for g in GORD if any(z['g'] == g for z in RBL)]
CF_NOTE = {'Aluguel e condomínio': 'Aluguel R$ 11.000,60 + condomínio R$ 1.460,66, pagos em setembro', 'Equipe — salários': 'Salários, diárias e verbas rescisórias da equipe de Rio Branco',
           'Equipe — benefícios e saúde ocupacional': 'Vale-alimentação, seguro de vida, combustível, medicina do trabalho e refeição',
           'Contabilidade, jurídico e sistemas': 'Contabilidade, assessoria jurídica e sistemas de gestão', 'Manutenção e reparos': 'Pintura, elétrica, ar-condicionado e cartucho da bioimpedância'}
# ---------------- salas produtivas (inventário patrimonial, aba Salas)
IV = openpyxl.load_workbook(SP + '/data7/inventario.xlsx', data_only=True)['Salas']
SALAS = [(r[0], r[1], float(r[3]), r[13]) for r in IV.iter_rows(min_row=7, values_only=True) if r and r[4] == 'Produtiva']
assert len(SALAS) == 6 and abs(sum(s[2] for s in SALAS) - 77.18) < 0.001

# ---------------- estilo (simples, letra grande)
PET = '0A6A70'; INK = '10272B'
F_HEAD = PatternFill('solid', fgColor=PET); F_IN = PatternFill('solid', fgColor='FFF2CC'); F_TOT = PatternFill('solid', fgColor='EEF3F3')
F_PAT = PatternFill('solid', fgColor='DCEBFA'); F_CLI = PatternFill('solid', fgColor='DDF1E7'); F_NOTE = PatternFill('solid', fgColor='F7F9F9')
thin = Side(style='thin', color='C9D4D4'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def F(size=13, bold=False, color=INK, italic=False): return Font(name='Calibri', size=size, bold=bold, color=color, italic=italic)
BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00;"R$" 0.00'
PCT = '0.00%'
wb = openpyxl.Workbook()
def name(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)
def q(s): return "'" + s + "'"
def setup(ws, widths, title, sub, landscape=False):
    ws.sheet_view.showGridLines = False; ws.sheet_view.zoomScale = 110
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws['A1'] = title; ws['A1'].font = F(20, True, PET); ws.row_dimensions[1].height = 34
    ws['A2'] = sub; ws['A2'].font = F(12, False, '4A5F63', True); ws.row_dimensions[2].height = 22
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = 'Clínica Núcleo S · Parceria de Tricologia'; ws.oddFooter.right.text = 'Página &P de &N'
def head(ws, r, vals, h=34):
    for i, v in enumerate(vals):
        c = ws.cell(r, 1 + i, v); c.fill = F_HEAD; c.font = F(13, True, 'FFFFFF'); c.border = BOX
        c.alignment = Alignment(horizontal='left' if i == 0 else 'center', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = h
def cell(ws, r, c, v, fmt=None, bold=False, fill=None, size=13, color=INK, wrap=False, h=None, italic=False):
    x = ws.cell(r, c, v); x.font = F(size, bold, color, italic); x.border = BOX
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    x.alignment = Alignment(vertical='center', wrap_text=wrap, horizontal=h)
    return x
def edit(ws, r, c, v, fmt=None):  # campo editável (amarelo, azul)
    return cell(ws, r, c, v, fmt, fill=F_IN, color='1F4FBF', bold=True)
def note(ws, r, text, cols=4, size=11):
    x = ws.cell(r, 1, text); x.font = F(size, False, '4A5F63', True); x.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=cols); ws.row_dimensions[r].height = 18 * max(1, len(text) // 95 + 1)

SH = ['Resumo', 'Receitas', 'Despesas', 'Insumos', 'Custo da sala', 'Despesas de RB', 'Como foi feito']
wb.active.title = SH[0]
for s in SH[1:]: wb.create_sheet(s)
L = openpyxl.utils.get_column_letter
BRL0 = '"R$" #,##0.00;[Red]-"R$" #,##0.00;"–"'

# ================= Receitas
ws = wb['Receitas']
setup(ws, [40, 16, 16, 14, 14, 14, 17, 20], 'Receitas — Rio Branco, 08/10/2026', 'Quanto cada paciente pagou, separado por tipo de serviço. Os campos amarelos podem ser alterados.', landscape=True)
head(ws, 4, ['Paciente'] + [t[1] + ' (R$)' for t in TIPOS] + ['Total (R$)', 'Pago no cartão de crédito (R$)'], 40)
r = 5; f0 = r
for n, p in PAC_ORD:
    cell(ws, r, 1, n.title(), bold=True)
    for i, t in enumerate(TIPOS): edit(ws, r, 2 + i, round(p[t[0]], 2), BRL0)
    cell(ws, r, 7, f'=SUM(B{r}:F{r})', BRL, bold=True); edit(ws, r, 8, round(p['cred'], 2), BRL0); ws.row_dimensions[r].height = 24; r += 1
f1 = r - 1; REC_F0, REC_F1 = f0, f1
cell(ws, r, 1, 'TOTAL', bold=True, fill=F_TOT)
for c in range(2, 9): cell(ws, r, c, f'=SUM({L(c)}{f0}:{L(c)}{f1})', BRL, True, F_TOT)
ws.row_dimensions[r].height = 28; tr = r
for i, t in enumerate(TIPOS): name(t[2], f"Receitas!${L(2 + i)}${tr}")
name('FAT_TOT', f"Receitas!$G${tr}"); name('FAT_CRED', f"Receitas!$H${tr}")
r += 2
ws.cell(r, 1, 'Receita por tipo de serviço').font = F(15, True, PET); ws.row_dimensions[r].height = 26; r += 1
head(ws, r, ['Tipo de receita', 'Valor (R$)', '% do total'], 28); r += 1; s0 = r
for t in TIPOS:
    cell(ws, r, 1, t[1]); cell(ws, r, 2, '=' + t[2], BRL); cell(ws, r, 3, f'=IFERROR(B{r}/FAT_TOT,0)', PCT); ws.row_dimensions[r].height = 24; r += 1
cell(ws, r, 1, 'RECEITA BRUTA TOTAL', bold=True, fill=F_TOT); cell(ws, r, 2, f'=SUM(B{s0}:B{r-1})', BRL, True, F_TOT); cell(ws, r, 3, f'=IFERROR(B{r}/FAT_TOT,0)', PCT, True, F_TOT)
ws.row_dimensions[r].height = 26; r += 2
ws.cell(r, 1, 'Receita por forma de pagamento').font = F(15, True, PET); ws.row_dimensions[r].height = 26; r += 1
head(ws, r, ['Forma de pagamento', 'Valor (R$)', '% do total'], 28); r += 1; s0 = r
cell(ws, r, 1, 'Cartão de crédito'); cell(ws, r, 2, '=FAT_CRED', BRL); cell(ws, r, 3, f'=IFERROR(B{r}/FAT_TOT,0)', PCT); ws.row_dimensions[r].height = 24; r += 1
cell(ws, r, 1, 'PIX, dinheiro e em aberto'); cell(ws, r, 2, '=FAT_TOT-FAT_CRED', BRL); cell(ws, r, 3, f'=IFERROR(B{r}/FAT_TOT,0)', PCT); ws.row_dimensions[r].height = 24; r += 1
cell(ws, r, 1, 'TOTAL', bold=True, fill=F_TOT); cell(ws, r, 2, f'=SUM(B{s0}:B{r-1})', BRL, True, F_TOT); cell(ws, r, 3, f'=IFERROR(B{r}/FAT_TOT,0)', PCT, True, F_TOT); ws.row_dimensions[r].height = 26; r += 2
note(ws, r, f'Consultas pelo valor total (R$ 800 ou R$ 900); a consulta de R$ {FMTBR(V_ANT)} paga em 23/09 via PIX entra em outubro, como em Cruzeiro do Sul. Procedimentos com mais de um serviço separados pelos detalhes da transação de 08/10. O pescoço de R$ {FMTBR(V_ABERTO)} está com pagamento em aberto (a cada 15 dias). Sobre o valor pago no cartão de crédito incidem a taxa da maquininha e a antecipação.', 8)
ws.freeze_panes = 'B5'

# ================= Insumos (por tratamento completo, por paciente)
ws = wb['Insumos']
setup(ws, [40, 12, 9, 14, 15, 11, 17, 42], 'Insumos de cada tratamento', 'Ativos e materiais usados em cada sessão, conforme as folhas de procedimento. Quantidades, preços e sessões em amarelo podem ser alterados.', landscape=True)
QT = 'General'
def ins_table(ws, r, rows, vezes):
    head(ws, r, ['Item', 'Quantidade', 'Unid.', 'Preço (R$)', 'Custo de cada vez (R$)', 'Vezes', 'Total (R$)', 'Observação'], 36); r += 1; t0 = r
    for it, qtd, un, pr, modo, ob in rows:
        cell(ws, r, 1, it); edit(ws, r, 2, qtd, QT); ws.cell(r, 2).alignment = Alignment(horizontal='center', vertical='center')
        cell(ws, r, 3, un, h='center', size=12); edit(ws, r, 4, pr, BRL); cell(ws, r, 5, f'=ROUND(B{r}*D{r},2)', BRL)
        if modo == 'S': cell(ws, r, 6, f'={vezes}', '0', h='center')
        else: edit(ws, r, 6, 1, '0'); ws.cell(r, 6).alignment = Alignment(horizontal='center', vertical='center')
        cell(ws, r, 7, f'=ROUND(E{r}*F{r},2)', BRL, True); cell(ws, r, 8, ob or None, size=11, color='4A5F63', wrap=True)
        ws.row_dimensions[r].height = 22 if len(ob or '') < 46 else 34; r += 1
    return r, t0
def ins_sub(ws, r, lab, t0, nm, fill=F_TOT, size=13):
    cell(ws, r, 1, lab, bold=True, fill=fill, size=size)
    for c in range(2, 7): cell(ws, r, c, None, fill=fill)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    cell(ws, r, 7, f'=SUM(G{t0}:G{r-1})' if t0 else None, BRL, True, fill, size); cell(ws, r, 8, None, fill=fill)
    if nm: name(nm, f"Insumos!$G${r}")
    ws.row_dimensions[r].height = 28; return r + 1
def kv(ws, r, lab, v, fmt, nm, ob=None, formula=False):
    cell(ws, r, 1, lab, bold=True)
    (cell(ws, r, 2, v, fmt, True, h='center') if formula else edit(ws, r, 2, v, fmt)); ws.cell(r, 2).alignment = Alignment(horizontal='center', vertical='center')
    for c in range(3, 9): cell(ws, r, c, ob if c == 3 else None, size=11, color='4A5F63')
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    if nm: name(nm, f"Insumos!$B${r}")
    ws.row_dimensions[r].height = 24; return r + 1
r = 4
x = ws.cell(r, 1, 'Resumo — insumos por paciente (tratamento fechado)'); x.font = F(15, True, PET); ws.row_dimensions[r].height = 26; r += 1
head(ws, r, ['Tratamento', 'Sessões', 'Custo por paciente (R$)', None, 'Pacientes', 'Total (R$)', None, 'De onde vem'], 36)
for a_, b_ in ((3, 4), (6, 7)): ws.merge_cells(start_row=r, start_column=a_, end_row=r, end_column=b_)
r += 1; rs0 = r
RES_I = [('Tricologia', '=TRI_SESS', '=INS_TRI', 'C', 'Tratamento completo: eletroporação + fototerapia + kit + sala (quadro 1)'), ('Pescoço (colo)', '=CFP_S', '=INS_CFP', 'E', 'Protocolo de eletroporação (quadro 2)'),
         ('Facial', '=CFP_S', '=INS_CFP', 'D', 'Protocolo de eletroporação (quadro 2)'), ('Corporal', '=CFP_S', '=INS_CFP', 'F', 'Protocolo de eletroporação (quadro 2)')]
for lab, ss, cp, col, ob in RES_I:
    cell(ws, r, 1, lab, bold=True); cell(ws, r, 2, ss, '0', h='center'); cell(ws, r, 3, cp, BRL); cell(ws, r, 4, None); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    cell(ws, r, 5, f'=COUNTIF(Receitas!{col}{REC_F0}:{col}{REC_F1},">0")', '0', h='center'); cell(ws, r, 6, f'=ROUND(C{r}*E{r},2)', BRL, True); cell(ws, r, 7, None)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7); cell(ws, r, 8, ob, size=11, color='4A5F63'); ws.row_dimensions[r].height = 26; r += 1
cell(ws, r, 1, 'TOTAL DE INSUMOS (vai para as Despesas)', bold=True, fill=F_CLI, size=14)
for c in (2, 3, 4, 5): cell(ws, r, c, None, fill=F_CLI)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
cell(ws, r, 6, f'=SUM(F{rs0}:F{r-1})', BRL, True, F_CLI, 14); cell(ws, r, 7, None, fill=F_CLI); ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7); cell(ws, r, 8, None, fill=F_CLI)
name('INS_TOTAL', f"Insumos!$F${r}"); ws.row_dimensions[r].height = 30; r += 1
note(ws, r, 'Pacientes: contados sozinhos na aba Receitas. Quem fez dois tratamentos (ex.: facial e pescoço) tem os insumos dos dois.', 8, 11); r += 2
# --- quadro 1: tricologia
ws.row_breaks.append(Break(id=r - 1))
x = ws.cell(r, 1, 'Quadro 1 — Tricologia (tratamento fechado)'); x.font = F(15, True, PET); ws.row_dimensions[r].height = 26; r += 1
r = kv(ws, r, 'Sessões de eletroporação', 3, '0', 'TRI_ELE_S', 'Aparelho de eletroporação')
r = kv(ws, r, 'Sessões de fototerapia', 8, '0', 'TRI_FOT_S', 'Cascata de LED')
r = kv(ws, r, 'Total de sessões', '=TRI_ELE_S+TRI_FOT_S', '0', 'TRI_SESS', 'Cada sessão é uma visita do paciente à clínica', True); r += 1
SER = 'A folha fala em até 5 seringas; a conta da folha usou 1 de 3 mL + 1 de 5 mL por sessão'
PENTE = 'Não entrou na conta final da folha (R$ 8,00). Para cobrar, ponha 1'
ELE = [('Hair cocktail', 2, 'mL', 23.80, 'S', '0 a 2 mL por sessão (máximo)'), ('Densi hair', 2, 'mL', 35.00, 'S', '0 a 2 mL'), ('Exossomos', 2, 'mL', 0.0, 'S', 'Preço não informado na folha: preencher'),
       ('Lumicen', 2, 'mL', 38.31, 'S', '0 a 2 mL'), ('A-OX', 2, 'mL', 24.00, 'S', '0 a 2 mL'), ('Luvas P', 1, 'par', 0.70, 'S', None),
       ('Seringa 3 mL com rosca', 1, 'un', 0.33, 'S', SER), ('Seringa 5 mL com rosca', 1, 'un', 0.50, 'S', None), ('Gaze', 4, 'un', 0.07, 'S', 'Pode gastar ou não (máximo 4)'),
       ('Lençol de maca', 1, 'un', 2.90, 'S', None), ('Máscara', 1, 'un', 0.22, 'S', None), ('Touca', 1, 'un', 0.15, 'S', None), ('Pente', 0, 'un', 8.00, 'T', PENTE)]
FOT = [('Minoxidil', 1, 'mL', 2.40, 'S', None), ('Dutasterida', 1, 'mL', 8.48, 'S', None), ('Clobetasol', 1, 'mL', 0.75, 'S', '2 a 20 gotas; frasco de R$ 22,27 com 30 mL'),
       ('Luvas P', 1, 'par', 0.70, 'S', None), ('Seringa 3 mL com rosca', 1, 'un', 0.33, 'S', SER), ('Seringa 5 mL com rosca', 1, 'un', 0.50, 'S', None),
       ('Gaze', 4, 'un', 0.07, 'S', 'Pode gastar ou não (máximo 4)'), ('Lençol de maca', 1, 'un', 2.90, 'S', None), ('Máscara', 1, 'un', 0.22, 'S', None),
       ('Touca', 1, 'un', 0.15, 'S', None), ('Pente', 0, 'un', 8.00, 'T', PENTE)]
x = ws.cell(r, 1, 'Eletroporação'); x.font = F(13, True, PET); r += 1
r, t0 = ins_table(ws, r, ELE, 'TRI_ELE_S'); r = ins_sub(ws, r, 'Total da eletroporação', t0, 'INS_TRI_ELE'); r += 1
ws.row_breaks.append(Break(id=r - 1))
x = ws.cell(r, 1, 'Fototerapia'); x.font = F(13, True, PET); r += 1
r, t0 = ins_table(ws, r, FOT, 'TRI_FOT_S'); r = ins_sub(ws, r, 'Total da fototerapia', t0, 'INS_TRI_FOT'); r += 1
x = ws.cell(r, 1, 'Kit do paciente e sala de procedimento'); x.font = F(13, True, PET); r += 1
r, t0 = ins_table(ws, r, [('Kit SP', 1, 'kit', 135.50, 'T', 'Valor da folha de tricologia'), ('Sala de procedimento (30 min por sessão)', 1, 'sessão', 20.00, 'S', 'R$ 20,00 por sessão, conforme a folha'), ('Ajuste ao total da folha', 1, 'un', 1.00, 'T', 'Luvas da eletroporação: a folha soma R$ 3,10')], 'TRI_SESS'); r = ins_sub(ws, r, 'Total do kit e da sala', t0, 'INS_TRI_KIT'); r += 1
cell(ws, r, 1, 'INSUMOS DA TRICOLOGIA POR PACIENTE (tratamento completo)', bold=True, fill=F_CLI, size=14)
for c in range(2, 7): cell(ws, r, c, None, fill=F_CLI)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
cell(ws, r, 7, '=INS_TRI_ELE+INS_TRI_FOT+INS_TRI_KIT', BRL, True, F_CLI, 14); cell(ws, r, 8, 'Eletroporação + fototerapia + kit + sala', size=11, color='4A5F63', fill=F_CLI)
name('INS_TRI', f"Insumos!$G${r}"); ws.row_dimensions[r].height = 30; r += 1
# --- quadro 2: facial, pescoço e corporal
ws.row_breaks.append(Break(id=r - 1))
x = ws.cell(r, 1, 'Quadro 2 — Facial, pescoço (colo) e corporal (tratamento fechado)'); x.font = F(15, True, PET); ws.row_dimensions[r].height = 26; r += 1
r = kv(ws, r, 'Sessões com insumo', 8, '0', 'CFP_S', 'Aparelho de eletroporação. O mesmo protocolo vale para facial, pescoço e corporal'); r += 1
CFP = [('WCPR', 2, 'mL', 73.20, 'S', '0 a 2 mL por sessão (máximo)'), ('Lumicen', 2, 'mL', 38.31, 'S', '0 a 2 mL'), ('Mesolift', 2, 'mL', 24.00, 'S', '0 a 2 mL'),
       ('Silhouette', 2, 'mL', 23.14, 'S', '0 a 2 mL'), ('A-OX', 2, 'mL', 24.00, 'S', '0 a 2 mL'), ('Luvas P', 1, 'par', 0.70, 'S', None),
       ('Seringa 3 mL com rosca', 1, 'un', 0.33, 'S', 'A folha fala em até 4 seringas; a conta usou 1 de 3 mL + 1 de 5 mL'), ('Seringa 5 mL com rosca', 1, 'un', 0.50, 'S', None),
       ('Lençol de maca', 1, 'un', 2.90, 'S', None), ('Touca', 1, 'un', 0.15, 'S', None), ('Máscara', 1, 'un', 0.22, 'S', None),
       ('Lenços umedecidos — radiofrequência (Unyque)', 1, 'kit', 0.14, 'S', '3 lenços para limpeza do paciente (facial e colo)'),
       ('Lenços umedecidos — Hipro', 1, 'kit', 0.14, 'S', '3 lenços para limpeza do paciente (facial)')]
r, t0 = ins_table(ws, r, CFP, 'CFP_S'); r = ins_sub(ws, r, 'INSUMOS POR PACIENTE (facial, pescoço ou corporal)', t0, 'INS_CFP', F_CLI, 14)
note(ws, r, 'Campo eletromagnético (Supreme) e Velaryan, no corporal, não usam insumo extra.', 8, 11); r += 1

# ================= Custo da sala (cálculo detalhado, por hora)
ws = wb['Custo da sala']
setup(ws, [50, 18, 14, 14, 15, 18], 'Custo da sala — cálculo detalhado', 'Quanto custa usar uma sala da clínica de Rio Branco por hora. Os campos amarelos podem ser alterados; o resto se calcula sozinho.', landscape=True)
def obs(ws, r, text, fill=None, color='4A5F63', bold=False, size=11):
    for c in range(3, 7): cell(ws, r, c, text if c == 3 else None, fill=fill, size=size, color=color, bold=bold, wrap=True)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
def hd3(ws, r, a, b, c):
    head(ws, r, [a, b, c, None, None, None]); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
def step(ws, r, text):
    x = ws.cell(r, 1, text); x.font = F(15, True, PET); ws.row_dimensions[r].height = 26
HRS = '0.0" h"'; MIN = '0" min"'
r = 4
step(ws, r, 'Passo 1 — Quanto custa manter a clínica de Rio Branco funcionando por mês'); r += 1
note(ws, r, 'Valores de setembro/2026, do relatório de transações: despesas de estrutura de Rio Branco, uma a uma na aba "Despesas de RB" (lá os valores podem ser alterados). Aluguel de R$ 11.000,60 e condomínio de R$ 1.460,66 pagos em setembro.', 6, 11); ws.row_dimensions[r].height = 30; r += 1
hd3(ws, r, 'Despesa da unidade', 'R$ no mês', 'Observação'); r += 1; c0 = r
for g in GRP_USED:
    cell(ws, r, 1, g); cell(ws, r, 2, f'=SUMIFS(UNI_VAL,UNI_GRP,A{r},UNI_ENT,"SIM")', BRL); obs(ws, r, CF_NOTE.get(g)); ws.row_dimensions[r].height = 20; r += 1
cell(ws, r, 1, 'TOTAL — custo mensal de Rio Branco', bold=True, fill=F_TOT); cell(ws, r, 2, f'=SUM(B{c0}:B{r-1})', BRL, True, F_TOT); obs(ws, r, None, F_TOT)
name('SALA_CF', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 28; r += 2
ws.row_breaks.append(Break(id=r - 1)); step(ws, r, 'Passo 2 — Salas produtivas e metragem (inventário da clínica)'); r += 1
note(ws, r, 'Só as salas que geram receita carregam o custo da clínica: recepção, espera, negociação, administração, copa, DML e circulação servem a todas elas. Área útil total da clínica: 177,61 m².', 6, 11); ws.row_dimensions[r].height = 30; r += 1
M2 = '0.00" m²"'; SALA_ROW = {}
hd3(ws, r, 'Sala produtiva', 'Área', 'Para que serve'); r += 1; a0 = r
for cod, nome, area, fin in SALAS:
    cell(ws, r, 1, nome); edit(ws, r, 2, area, M2); obs(ws, r, fin); ws.row_dimensions[r].height = 30 if len(fin or '') > 70 else 22; SALA_ROW[cod] = r; r += 1
cell(ws, r, 1, 'ÁREA PRODUTIVA TOTAL', bold=True, fill=F_TOT); cell(ws, r, 2, f'=SUM(B{a0}:B{r-1})', M2, True, F_TOT); obs(ws, r, f'=COUNT(B{a0}:B{r-1})&" salas produtivas"', F_TOT)
name('AREA_PROD', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 28; NSALAS = f'COUNT(B{a0}:B{r-1})'; r += 2
hd3(ws, r, 'Tempo de uso de cada sala', 'Quantidade', 'Observação'); r += 1
cell(ws, r, 1, 'Dias de funcionamento por mês'); edit(ws, r, 2, 22, '0'); obs(ws, r, 'Dias úteis (planilha de precificação)'); name('SALA_DIAS', f"{q('Custo da sala')}!$B${r}"); r += 1
cell(ws, r, 1, 'Horas de funcionamento por dia'); edit(ws, r, 2, 10, HRS); obs(ws, r, 'Planilha de precificação'); name('SALA_H', f"{q('Custo da sala')}!$B${r}"); r += 1
cell(ws, r, 1, 'Taxa de ocupação considerada'); edit(ws, r, 2, 0.65, '0%')
obs(ws, r, 'Mesma taxa de Cruzeiro do Sul e da planilha de precificação: as salas ficam ocupadas 65% do tempo e o custo do tempo vazio é dividido entre as horas usadas.')
ws.row_dimensions[r].height = 46; name('SALA_OCUP', f"{q('Custo da sala')}!$B${r}"); r += 1
cell(ws, r, 1, 'Horas ocupadas de cada sala por mês', bold=True, fill=F_TOT); cell(ws, r, 2, '=SALA_DIAS*SALA_H*SALA_OCUP', HRS, True, F_TOT); obs(ws, r, 'Dias × horas por dia × ocupação', F_TOT)
name('SALA_HM', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 28; r += 1
cell(ws, r, 1, 'Metro quadrado × horas ocupadas no mês', bold=True, fill=F_TOT); cell(ws, r, 2, '=AREA_PROD*SALA_HM', '#,##0.0', True, F_TOT); obs(ws, r, 'Área produtiva × horas ocupadas de cada sala', F_TOT)
name('SALA_HD', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 28; r += 2
step(ws, r, 'Passo 3 — Custo de uma sala por hora (pela metragem)'); r += 1
hd3(ws, r, 'Cálculo', 'R$', 'Como é feito'); r += 1
cell(ws, r, 1, 'Custo de 1 m² por hora'); cell(ws, r, 2, '=ROUND(SALA_CF/SALA_HD,2)', BRL); obs(ws, r, 'Custo mensal ÷ (área produtiva × horas ocupadas de cada sala)'); name('SALA_M2H', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 26; r += 1
cell(ws, r, 1, 'Área da sala usada pela Dra. Patrícia'); edit(ws, r, 2, 9.45, M2); obs(ws, r, 'Sala de Procedimento (9,45 m²), que aceita todas as tecnologias. Para outra sala, troque a área.')
name('AREA_USO', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 30; r += 1
cell(ws, r, 1, 'CUSTO REAL de 1 sala por hora', bold=True, fill=F_TOT, size=14); cell(ws, r, 2, '=ROUND(SALA_M2H*AREA_USO,2)', BRL, True, F_TOT, 14); obs(ws, r, 'Custo de 1 m² por hora × área da sala', F_TOT)
name('SALA_HORA', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 30; r += 1
cell(ws, r, 1, 'Custo real de 1 sala por dia inteiro (para comparar)'); cell(ws, r, 2, '=ROUND(SALA_HORA*SALA_H,2)', BRL); obs(ws, r, 'Custo por hora × horas de funcionamento por dia'); r += 1
cell(ws, r, 1, 'CUSTO NEGOCIADO de 1 sala por hora', bold=True, fill=F_CLI, size=14); edit(ws, r, 2, 40.0, BRL); ws.cell(r, 2).font = F(14, True, '1F4FBF')
obs(ws, r, 'Mesmo valor combinado pela diretoria para esta bateria de procedimentos (Cruzeiro do Sul e Rio Branco). É este valor que entra na conta.', F_CLI)
name('SALA_NEG', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 36; r += 1
cell(ws, r, 1, 'Desconto por hora concedido na negociação'); cell(ws, r, 2, '=SALA_HORA-SALA_NEG', BRL); obs(ws, r, '=IFERROR("Custo real − custo negociado: "&FIXED((1-SALA_NEG/SALA_HORA)*100,1)&"% abaixo do custo real","")'); r += 2
ws.row_breaks.append(Break(id=r - 1)); step(ws, r, 'Passo 4 — Consultas da Dra. Patrícia'); r += 1
note(ws, r, 'Na consulta a Dra. Patrícia avalia o paciente e define o tratamento. Cada consulta ocupa uma sala por 1 hora.', 6, 12); ws.row_dimensions[r].height = 22; r += 1
hd3(ws, r, 'Item', 'Quantidade', 'Observação'); r += 1
cell(ws, r, 1, 'Tempo de cada consulta', bold=True); edit(ws, r, 2, 1, HRS); obs(ws, r, '1 hora por consulta'); name('CONS_H', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 26; r += 1
cell(ws, r, 1, 'Consultas realizadas'); cell(ws, r, 2, f'=COUNTIF(Receitas!B{REC_F0}:B{REC_F1},">0")', '0', True); obs(ws, r, 'Pacientes com consulta na aba Receitas (contados sozinhos)')
name('CONS_N', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 26; r += 1
cell(ws, r, 1, 'Horas de sala nas consultas', bold=True, fill=F_TOT); cell(ws, r, 2, '=CONS_N*CONS_H', HRS, True, F_TOT); obs(ws, r, 'Consultas × tempo de cada consulta', F_TOT)
name('CONS_HS', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 28; r += 2
step(ws, r, 'Passo 5 — Visitas para os procedimentos (tratamento fechado)'); r += 1
note(ws, r, 'Depois da consulta, cada paciente volta à clínica para as visitas do tratamento fechado. Cada visita ocupa uma sala. Quem fez dois tratamentos (ex.: facial e pescoço) faz as visitas dos dois.', 6, 12); ws.row_dimensions[r].height = 36; r += 1
head(ws, r, ['Tratamento', 'Pacientes', 'Visitas por paciente', 'Total de visitas', 'Tempo de cada visita', 'Horas de sala'], 40); r += 1; v0 = r
VIS = [('Tricologia', 'C', 11, 30, 'TRI'), ('Pescoço', 'E', 8, 30, 'PES'), ('Facial', 'D', 8, 30, 'FAC'), ('Corporal', 'F', 16, 60, 'COR')]
for lab, col, nv, mn, k in VIS:
    cell(ws, r, 1, lab, bold=True); cell(ws, r, 2, f'=COUNTIF(Receitas!{col}{REC_F0}:{col}{REC_F1},">0")', '0', h='center')
    edit(ws, r, 3, nv, '0'); ws.cell(r, 3).alignment = Alignment(horizontal='center', vertical='center'); name('VIS_' + k, f"{q('Custo da sala')}!$C${r}")
    cell(ws, r, 4, f'=B{r}*C{r}', '0', h='center')
    edit(ws, r, 5, mn, MIN); ws.cell(r, 5).alignment = Alignment(horizontal='center', vertical='center'); name('MIN_' + k, f"{q('Custo da sala')}!$E${r}")
    cell(ws, r, 6, f'=ROUND(D{r}*E{r}/60,2)', HRS, True, h='center'); ws.row_dimensions[r].height = 26; r += 1
cell(ws, r, 1, 'TOTAL', bold=True, fill=F_TOT); cell(ws, r, 2, None, fill=F_TOT); cell(ws, r, 3, None, fill=F_TOT)
cell(ws, r, 4, f'=SUM(D{v0}:D{r-1})', '0', True, F_TOT, h='center'); cell(ws, r, 5, None, fill=F_TOT); cell(ws, r, 6, f'=SUM(F{v0}:F{r-1})', HRS, True, F_TOT, h='center')
name('AVAL_HS', f"{q('Custo da sala')}!$F${r}"); name('VIS_TOT', f"{q('Custo da sala')}!$D${r}"); ws.row_dimensions[r].height = 28; r += 1
note(ws, r, 'Tricologia, pescoço e facial: 30 minutos por visita. Corporal: 16 visitas de 1 hora. Pacientes contados sozinhos na aba Receitas.', 6, 11); r += 2
ws.row_breaks.append(Break(id=r - 1)); x = ws.cell(r, 1, 'Horas de sala de cada paciente (consulta + visitas)'); x.font = F(13, True, PET); ws.row_dimensions[r].height = 24; r += 1
head(ws, r, ['Paciente', 'Tratamento', 'Consulta', 'Visitas', 'Horas das visitas', 'Total de horas'], 40); r += 1; w0 = r
for i in range(REC_F0, REC_F1 + 1):
    R = lambda c: f'Receitas!{c}{i}'
    cell(ws, r, 1, f'={R("A")}', bold=True)
    trat = (f'=IF({R("C")}+{R("D")}+{R("E")}+{R("F")}=0,"Só consulta",MID(IF({R("C")}>0,", Tricologia","")&IF({R("E")}>0,", Pescoço","")'
            f'&IF({R("D")}>0,", Facial","")&IF({R("F")}>0,", Corporal",""),3,80))')
    cell(ws, r, 2, trat, size=12, wrap=True)
    cell(ws, r, 3, f'=IF({R("B")}>0,CONS_H,0)', HRS, h='center')
    cell(ws, r, 4, f'=IF({R("C")}>0,VIS_TRI,0)+IF({R("E")}>0,VIS_PES,0)+IF({R("D")}>0,VIS_FAC,0)+IF({R("F")}>0,VIS_COR,0)', '0', h='center')
    cell(ws, r, 5, f'=ROUND((IF({R("C")}>0,VIS_TRI*MIN_TRI,0)+IF({R("E")}>0,VIS_PES*MIN_PES,0)+IF({R("D")}>0,VIS_FAC*MIN_FAC,0)+IF({R("F")}>0,VIS_COR*MIN_COR,0))/60,2)', HRS, h='center')
    cell(ws, r, 6, f'=C{r}+E{r}', HRS, True, h='center'); ws.row_dimensions[r].height = 26; r += 1
cell(ws, r, 1, 'TOTAL', bold=True, fill=F_TOT); cell(ws, r, 2, None, fill=F_TOT)
for c in (3, 4, 5, 6): cell(ws, r, c, f'=SUM({L(c)}{w0}:{L(c)}{r-1})', '0' if c == 4 else HRS, True, F_TOT, h='center')
name('PAC_HS', f"{q('Custo da sala')}!$F${r}"); ws.row_dimensions[r].height = 28; r += 1
note(ws, r, 'Consulta: 1 hora para quem tem consulta na aba Receitas. Visitas: conforme o tratamento fechado (quadro acima).', 6, 11); r += 2
ws.row_breaks.append(Break(id=r - 1)); step(ws, r, 'Conta final — taxa de ocupação da sala'); r += 1
hd3(ws, r, 'Cálculo', 'Valor', 'Como é feito'); r += 1
cell(ws, r, 1, 'Horas de sala — consultas'); cell(ws, r, 2, '=CONS_HS', HRS); obs(ws, r, '="Passo 4: "&CONS_N&" consultas × "&FIXED(CONS_H,1)&" h"'); r += 1
cell(ws, r, 1, 'Horas de sala — visitas dos procedimentos'); cell(ws, r, 2, '=AVAL_HS', HRS); obs(ws, r, '="Passo 5: "&VIS_TOT&" visitas"'); r += 1
cell(ws, r, 1, 'Total de horas de sala', bold=True, fill=F_TOT); cell(ws, r, 2, '=CONS_HS+AVAL_HS', HRS, True, F_TOT); obs(ws, r, None, F_TOT); name('SALA_HS', f"{q('Custo da sala')}!$B${r}"); r += 1
cell(ws, r, 1, 'Custo negociado de 1 sala por hora'); cell(ws, r, 2, '=SALA_NEG', BRL); obs(ws, r, 'Passo 3 (valor combinado para esta bateria)'); r += 1
cell(ws, r, 1, 'Ocupação — consultas'); cell(ws, r, 2, '=ROUND(SALA_NEG*CONS_HS,2)', BRL); obs(ws, r, 'Horas das consultas × custo negociado'); name('SALA_AT', f"{q('Custo da sala')}!$B${r}"); r += 1
cell(ws, r, 1, 'Ocupação — visitas dos procedimentos'); cell(ws, r, 2, '=ROUND(SALA_NEG*AVAL_HS,2)', BRL); obs(ws, r, 'Horas das visitas × custo negociado'); name('SALA_AV', f"{q('Custo da sala')}!$B${r}"); r += 1
cell(ws, r, 1, 'TAXA DE OCUPAÇÃO DA SALA (negociada)', bold=True, fill=F_CLI, size=14); cell(ws, r, 2, '=SALA_AT+SALA_AV', BRL, True, F_CLI, 14)
obs(ws, r, 'Consultas + visitas · é este valor que vai para o Resumo', F_CLI); name('SALA_TAXA', f"{q('Custo da sala')}!$B${r}"); ws.row_dimensions[r].height = 32; r += 1
GR = '7A8B8E'
cell(ws, r, 1, 'Pelo custo real seria (para comparar)', color=GR, italic=True); cell(ws, r, 2, '=ROUND(SALA_HORA*CONS_HS,2)+ROUND(SALA_HORA*AVAL_HS,2)', BRL, color=GR, italic=True)
obs(ws, r, '=FIXED(SALA_HS,1)&" h × custo real de R$ "&FIXED(SALA_HORA,2)&" por hora"', color=GR); name('SALA_REAL', f"{q('Custo da sala')}!$B${r}"); r += 1
cell(ws, r, 1, 'Desconto concedido pela clínica nesta negociação', bold=True, color='0F7F59'); cell(ws, r, 2, '=SALA_REAL-SALA_TAXA', BRL, True, color='0F7F59')
obs(ws, r, 'Custo real − custo negociado. A clínica absorve essa diferença.', color='0F7F59'); name('SALA_DESC', f"{q('Custo da sala')}!$B${r}"); r += 2
ws.row_breaks.append(Break(id=r - 1)); x = ws.cell(r, 1, 'Para comparar: custo real de 1 hora em cada sala produtiva'); x.font = F(13, True, PET); r += 1
head(ws, r, ['Sala', 'Área', 'Custo real por hora'], 34); r += 1
for cod, nome, area, fin in SALAS:
    cell(ws, r, 1, nome); cell(ws, r, 2, f'=B{SALA_ROW[cod]}', M2); cell(ws, r, 3, f'=ROUND(SALA_M2H*B{r},2)', BRL); ws.row_dimensions[r].height = 22; r += 1
cell(ws, r, 1, 'Média por sala (custo mensal ÷ salas ÷ horas ocupadas)', bold=True, fill=F_TOT); cell(ws, r, 2, None, fill=F_TOT)
cell(ws, r, 3, f'=ROUND(SALA_CF/({NSALAS}*SALA_HM),2)', BRL, True, F_TOT); ws.row_dimensions[r].height = 26; r += 2
note(ws, r, 'Mesmo método de Cruzeiro do Sul (custo mensal ÷ horas de sala ocupadas), agora pela metragem, porque as salas de Rio Branco têm tamanhos diferentes: a soroterapia (21,88 m², 5 lugares) custa mais por hora que uma sala de procedimento (9,45 m²). Custos de setembro/2026 de Rio Branco.', 6, 11)

# ================= Despesas de RB (lançamentos de setembro/2026)
ws = wb['Despesas de RB']
setup(ws, [14, 40, 44, 18, 14, 48], 'Despesas de Rio Branco — setembro/2026', 'Lançamentos do relatório de transações que formam o custo da sala. Os valores podem ser alterados; na coluna "Entra no custo?" pode trocar SIM por NÃO.', landscape=True)
ws.print_title_rows = '4:4'
head(ws, 4, ['Vencimento', 'Despesa', 'Grupo', 'Valor (R$)', 'Entra no custo?', 'Observação'], 34)
dv = DataValidation(type='list', formula1='"SIM,NÃO"', allow_blank=False); ws.add_data_validation(dv)
r = 5; z0 = r
for z in RBL:
    cell(ws, r, 1, z['d'], h='center'); cell(ws, r, 2, z['lab']); cell(ws, r, 3, z['g'], size=12, color='4A5F63'); edit(ws, r, 4, z['v'], BRL)
    edit(ws, r, 5, 'SIM'); ws.cell(r, 5).alignment = Alignment(horizontal='center', vertical='center'); dv.add(ws.cell(r, 5))
    cell(ws, r, 6, z['obs'] or None, size=11, color='4A5F63', wrap=True); ws.row_dimensions[r].height = 30 if len(z['obs']) > 45 else 22; r += 1
z1 = r - 1
name('UNI_VAL', f"{q('Despesas de RB')}!$D${z0}:$D${z1}"); name('UNI_GRP', f"{q('Despesas de RB')}!$C${z0}:$C${z1}"); name('UNI_ENT', f"{q('Despesas de RB')}!$E${z0}:$E${z1}")
cell(ws, r, 1, None, fill=F_TOT); cell(ws, r, 2, 'TOTAL QUE ENTRA NO CUSTO DA SALA', bold=True, fill=F_TOT); cell(ws, r, 3, None, fill=F_TOT)
cell(ws, r, 4, '=SUMIFS(UNI_VAL,UNI_ENT,"SIM")', BRL, True, F_TOT); cell(ws, r, 5, None, fill=F_TOT); cell(ws, r, 6, None, fill=F_TOT)
name('UNI_TOT', f"{q('Despesas de RB')}!$D${r}"); ws.row_dimensions[r].height = 28; r += 2
note(ws, r, 'Não entram no custo da sala: retiradas e pró-labore dos sócios, aluguel do depósito, empréstimos, parcelamentos e tributos, compras de equipamentos e materiais permanentes, medicamentos e insumos (já estão na despesa Insumos), marketing, consultorias, faturas de cartão de crédito, fretes, reembolsos e despesas de Epitaciolândia.', 6, 11)
ws.freeze_panes = 'A5'

# ================= Despesas (como uma conta razão)
ws = wb['Despesas']
setup(ws, [50, 20, 16, 18, 60], 'Despesas descontadas da receita bruta', 'Como uma conta razão: cada despesa com as linhas que a formam e de onde vem cada valor. Campos amarelos podem ser alterados.', landscape=True)
ws.print_title_rows = '4:4'
head(ws, 4, ['Despesa / detalhe', 'Base ou quantidade', 'Taxa ou preço', 'Valor (R$)', 'De onde vem'])
r = 5; SUBS = []; GR = '7A8B8E'
def conta(t):
    global r
    for c in range(1, 6): cell(ws, r, c, t if c == 1 else None, bold=True, fill=F_TOT, size=14, color=PET)
    ws.row_dimensions[r].height = 28; r += 1
def det(lab, base, bfmt, taxa, tfmt, val, origem, nm=None, edit_tx=False, info=False):
    global r
    col = GR if info else INK
    cell(ws, r, 1, '   ' + lab, size=12, color=col, italic=info, wrap=True); cell(ws, r, 2, base, bfmt, size=12, color=col)
    if edit_tx: edit(ws, r, 3, taxa, tfmt)
    else: cell(ws, r, 3, taxa, tfmt, size=12, color=col)
    cell(ws, r, 4, val.format(r=r) if val else None, BRL, size=12, color=col, italic=info); cell(ws, r, 5, origem, size=11, color='4A5F63', wrap=True)
    if nm: name(nm, f"Despesas!$D${r}")
    ws.row_dimensions[r].height = 28; r += 1; return r - 1
def subtotal(lab, rows, nm):
    global r
    cell(ws, r, 1, 'Total — ' + lab, bold=True); cell(ws, r, 2, None); cell(ws, r, 3, None)
    cell(ws, r, 4, '=' + '+'.join(f'D{x}' for x in rows), BRL, True); cell(ws, r, 5, None)
    name(nm, f"Despesas!$D${r}"); SUBS.append(r); ws.row_dimensions[r].height = 26; r += 2
VAL = '=ROUND(B{r}*C{r},2)'
conta('1. PIS e COFINS')
a = det('PIS', '=FAT_TOT', BRL, 0.0065, PCT, VAL, 'Receita bruta (aba Receitas) × 0,65% (Lucro Presumido)', 'C_PIS', True)
b = det('COFINS', '=FAT_TOT', BRL, 0.03, PCT, VAL, 'Receita bruta (aba Receitas) × 3% (Lucro Presumido)', 'C_COF', True)
subtotal('PIS e COFINS', [a, b], 'S_IMP')
conta('2. Taxa do cartão (maquininha)')
a = det('Vendas no cartão de crédito', '=FAT_CRED', BRL, 0.0192, PCT, VAL, 'Aba Receitas. Mesma taxa de Cruzeiro do Sul; em Rio Branco o relatório mostra R$ 643,73 sobre R$ 33.700,00 com taxa registrada (1,91%)', 'C_MDR', True)
subtotal('taxa do cartão', [a], 'S_MDR')
conta('3. Antecipação do cartão')
a = det('Vendas no crédito antecipadas', '=FAT_CRED', BRL, 0.0874, PCT, VAL, 'Mesma regra de Cruzeiro do Sul: todas as vendas no crédito antecipadas (taxa informada pela diretoria)', 'C_ANT', True)
subtotal('antecipação', [a], 'S_ANT')
conta('4. Insumos (material de cada tratamento fechado)')
rows = []
for lab, col, cp, ob in (('Tricologia', 'C', 'INS_TRI', 'Aba Insumos, quadro 1: tratamento completo (eletroporação + fototerapia + kit + sala)'), ('Pescoço (colo)', 'E', 'INS_CFP', 'Aba Insumos, quadro 2'),
                         ('Facial', 'D', 'INS_CFP', 'Aba Insumos, quadro 2'), ('Corporal', 'F', 'INS_CFP', 'Aba Insumos, quadro 2')):
    rows.append(det(f'{lab}: pacientes × insumos por paciente', f'=COUNTIF(Receitas!{col}{REC_F0}:{col}{REC_F1},">0")', '0" pac."', f'={cp}', BRL, VAL, ob))
subtotal('insumos', rows, 'C_INS')
ws.row_breaks.append(Break(id=r - 1)); conta('5. Viagem (passagens, alimentação e táxi)')
a = det('Já descontada no fechamento de Cruzeiro do Sul', None, None, None, None, '=0', 'Passagens, alimentação e táxi da viagem entraram por inteiro em Cruzeiro do Sul', 'C_VIAG')
subtotal('viagem', [a], 'S_VIAG')
conta('6. Ocupação da sala (custo-hora negociado)')
a = det('Consultas: horas × custo negociado', '=CONS_HS', HRS, '=SALA_NEG', BRL, '=SALA_AT', '1 hora por consulta (aba Custo da sala, Passo 4)', 'C_SALA1')
b = det('Visitas dos procedimentos: horas × custo negociado', '=AVAL_HS', HRS, '=SALA_NEG', BRL, '=SALA_AV', 'Visitas do tratamento fechado (aba Custo da sala, Passo 5)', 'C_SALA2')
det('Para comparar: horas × custo real (não entra na conta)', '=SALA_HS', HRS, '=SALA_HORA', BRL, '=SALA_REAL', 'Custo real de 1 sala por hora (aba Custo da sala, Passo 3)', info=True)
det('Desconto concedido pela clínica (não entra na conta)', None, None, None, None, '=SALA_DESC', 'Custo real − custo negociado', info=True)
subtotal('ocupação da sala', [a, b], 'S_SALA')
conta('7. IRPJ e CSLL')
a = det('Sobre as consultas', '=FAT_CONS', BRL, 0.0768, PCT, VAL, 'Lucro Presumido: 32% × 24% = 7,68%', 'C_IRC', True)
b = det('Sobre os procedimentos', '=FAT_TRI+FAT_FAC+FAT_PES+FAT_COR', BRL, 0.0228, PCT, VAL, 'Lucro Presumido: 8% × 15% + 12% × 9% = 2,28%', 'C_IRP', True)
subtotal('IRPJ e CSLL', [a, b], 'S_IRC')
for c in range(1, 6): cell(ws, r, c, 'TOTAL DAS DESPESAS' if c == 1 else None, bold=True, fill=F_CLI, size=15)
cell(ws, r, 4, '=' + '+'.join(f'D{x}' for x in SUBS), BRL, True, F_CLI, 15); name('C_TOT', f"Despesas!$D${r}"); ws.row_dimensions[r].height = 32; r += 2
note(ws, r, 'Todas as despesas saem da receita bruta. Linhas em cinza são só para comparação e não entram na soma.', 5)

# ================= Resumo
ws = wb['Resumo']
setup(ws, [50, 22, 15, 52], 'Parceria de Tricologia — Rio Branco', 'Atendimentos de 08 de outubro de 2026 · Clínica Núcleo S e Dra. Patrícia Fabrini')
ws['A3'] = 'Da receita bruta saem só as 7 despesas abaixo. O que sobra é dividido meio a meio.'; ws['A3'].font = F(13, True, '4A5F63')
head(ws, 5, ['', 'Valor (R$)', '% da receita', 'O que é'])
def sec(ws, r, text):
    x = ws.cell(r, 1, text); x.font = F(14, True, PET); ws.row_dimensions[r].height = 28
def line(ws, r, lab, f, oq, kind=None):
    fill = {'t': F_TOT, 'res': F_CLI}.get(kind); big = kind in ('t', 'res'); sz = 15 if kind == 'res' else 14 if big else 13
    cell(ws, r, 1, lab, None, big, fill, sz); cell(ws, r, 2, f, BRL, big, fill, sz)
    cell(ws, r, 3, f'=IFERROR(B{r}/FAT_TOT,0)', PCT, big, fill); cell(ws, r, 4, oq, size=11, color='4A5F63', fill=fill, wrap=True)
    ws.row_dimensions[r].height = 32 if big else 28
r = 6; sec(ws, r, 'RECEITAS'); r += 1; r0 = r
REC = [('Consultas', 'Consultas da Dra. Patrícia (valor total de R$ 800 ou R$ 900)'), ('Tricologia', 'Tratamentos capilares'), ('Facial', 'Procedimento facial'),
       ('Pescoço', 'Procedimento de pescoço'), ('Corporal', 'Procedimentos corporais')]
for (lab, oq), t in zip(REC, TIPOS): line(ws, r, '   ' + lab, '=' + t[2], oq); r += 1
line(ws, r, 'RECEITA BRUTA TOTAL', f'=SUM(B{r0}:B{r-1})', 'Tudo o que os atendimentos faturaram', 't'); rb = r; r += 2
sec(ws, r, 'DESPESAS (descontadas da receita bruta)'); r += 1; d0 = r
DSP = [('1. PIS e COFINS', '=S_IMP', 'Impostos federais sobre a receita'), ('2. Taxa do cartão (maquininha)', '=S_MDR', 'Cobrada nas vendas no cartão de crédito'),
       ('3. Antecipação do cartão', '=S_ANT', 'Para receber já as vendas parceladas'), ('4. Insumos', '=C_INS', 'Material de cada tratamento fechado (aba Insumos)'),
       ('5. Viagem', '=S_VIAG', 'Já descontada no fechamento de Cruzeiro do Sul'),
       ('6. Ocupação da sala', '=S_SALA', '=FIXED(SALA_HS,1)&" h × R$ "&FIXED(SALA_NEG,2)&" (negociado; custo real R$ "&FIXED(SALA_HORA,2)&")"'),
       ('7. IRPJ e CSLL', '=S_IRC', 'Impostos sobre o lucro')]
for lab, f, oq in DSP: line(ws, r, '   ' + lab, f, oq); r += 1
line(ws, r, 'TOTAL DAS DESPESAS', f'=SUM(B{d0}:B{r-1})', 'Soma das 7 despesas', 't'); td = r; r += 2
line(ws, r, 'RESULTADO PARA DIVIDIR', f'=B{rb}-B{td}', 'Receita bruta total − total das despesas', 'res'); res = r; r += 2
sec(ws, r, 'DIVISÃO DO RESULTADO'); r += 1
cell(ws, r, 1, 'Parte da Dra. Patrícia'); edit(ws, r, 2, 0.5, '0%'); cell(ws, r, 3, None); cell(ws, r, 4, 'Meio a meio (50%)', size=11, color='4A5F63'); name('PARTE_PAT', f"Resumo!$B${r}"); ws.row_dimensions[r].height = 26; r += 1
cell(ws, r, 1, 'DRA. PATRÍCIA RECEBE', None, True, F_PAT, 16, '1F4FBF'); cell(ws, r, 2, f'=ROUND(B{res}*PARTE_PAT,2)', BRL, True, F_PAT, 16, '1F4FBF')
cell(ws, r, 3, f'=IFERROR(B{r}/FAT_TOT,0)', PCT, True, F_PAT); cell(ws, r, 4, 'Contra nota fiscal da empresa dela', size=11, color='4A5F63', fill=F_PAT); ws.row_dimensions[r].height = 36; rp = r; r += 1
cell(ws, r, 1, 'CLÍNICA NÚCLEO S RECEBE', None, True, F_CLI, 16, '0F7F59'); cell(ws, r, 2, f'=B{res}-B{rp}', BRL, True, F_CLI, 16, '0F7F59')
cell(ws, r, 3, f'=IFERROR(B{r}/FAT_TOT,0)', PCT, True, F_CLI); cell(ws, r, 4, 'O restante do resultado', size=11, color='4A5F63', fill=F_CLI); ws.row_dimensions[r].height = 36; rc = r; r += 2
cell(ws, r, 1, 'De cada R$ 100 de receita, a Dra. Patrícia recebe', bold=True); cell(ws, r, 2, f'=IFERROR(B{rp}/FAT_TOT*100,0)', BRL, True); ws.row_dimensions[r].height = 26; r += 1
cell(ws, r, 1, 'Desconto da sala concedido pela clínica', color='0F7F59'); cell(ws, r, 2, '=SALA_DESC', BRL, color='0F7F59')
cell(ws, r, 4, 'Custo real da sala − custo negociado de R$ 40 por hora (aba Custo da sala)', size=11, color='4A5F63', wrap=True); ws.row_dimensions[r].height = 26; r += 2
x = ws.cell(r, 1, 'Conferência'); x.font = F(11, True, '4A5F63'); r += 1
CHK = [('Receita bruta igual à aba Receitas', f'=IF(ABS(B{rb}-FAT_TOT)<0.005,"OK","VERIFICAR")'),
       ('Total das despesas igual à aba Despesas', f'=IF(ABS(B{td}-C_TOT)<0.005,"OK","VERIFICAR")'),
       ('Insumos iguais à aba Insumos', '=IF(ABS(C_INS-INS_TOTAL)<0.005,"OK","VERIFICAR")'),
       ('Taxa da sala igual à aba Custo da sala', '=IF(ABS(S_SALA-SALA_TAXA)<0.005,"OK","VERIFICAR")'),
       ('Custo mensal da sala igual à aba Despesas de RB', '=IF(ABS(SALA_CF-UNI_TOT)<0.005,"OK","VERIFICAR")'),
       ('Horas de sala: por paciente = consultas + visitas', '=IF(ABS(PAC_HS-CONS_HS-AVAL_HS)<0.005,"OK","VERIFICAR")'),
       ('Dra. Patrícia + clínica = resultado', f'=IF(ABS(B{rp}+B{rc}-B{res})<0.005,"OK","VERIFICAR")')]
for lab, f in CHK:
    a = ws.cell(r, 1, lab); a.font = F(11, False, '4A5F63'); bb = ws.cell(r, 2, f); bb.font = F(11, True, '0F7F59'); bb.alignment = Alignment(horizontal='center'); r += 1
ws.freeze_panes = 'A6'

# ================= Como foi feito
ws = wb['Como foi feito']
setup(ws, [6, 110], 'Como os números foram feitos', 'De onde vem cada valor desta planilha. Os campos amarelos de todas as abas podem ser alterados; o resto se recalcula sozinho.')
TXT = [('1', f'Receitas: relatório de transações de Rio Branco de 08/10/2026 ({N_LANC} lançamentos, R$ 61.300,00). Os lançamentos com mais de um serviço foram separados pelos detalhes da transação: facial R$ 8.000 + tricologia R$ 10.000; corporal, facial e pescoço R$ 7.200 cada. A consulta de R$ {FMTBR(V_ANT)} paga em 23/09 via PIX entra em outubro, como em Cruzeiro do Sul. O pescoço de R$ {FMTBR(V_ABERTO)} (eletroporação) está com pagamento em aberto, a cada 15 dias, e entra pelo valor do tratamento.'),
       ('2', 'Regra da divisão: da receita bruta saem só as despesas desta planilha. O que sobra é dividido 50% para a Dra. Patrícia e 50% para a Clínica Núcleo S. Não entram a taxa de estrutura nem a amortização do investimento.'),
       ('3', 'PIS (0,65%) e COFINS (3%): alíquotas do Lucro Presumido, sobre a receita bruta.'),
       ('4', 'Taxa do cartão (1,92%) e antecipação (8,74%): as mesmas de Cruzeiro do Sul, sobre as vendas no cartão de crédito (R$ 51.700,00). Em Rio Branco o relatório mostra R$ 643,73 de taxa sobre R$ 33.700,00 com taxa registrada (1,91%).'),
       ('5', 'Insumos: os mesmos de Cruzeiro do Sul, por tratamento fechado. Tricologia: R$ 1.232,08 por paciente (3 sessões de eletroporação + 8 de fototerapia + kit + sala de procedimento). Facial, pescoço e corporal: 8 sessões do protocolo de eletroporação.'),
       ('6', 'Viagem: passagens, alimentação e táxi foram descontados por inteiro no fechamento de Cruzeiro do Sul. Em Rio Branco não entram de novo.'),
       ('7', 'Ocupação da sala — custo real: despesas de estrutura de Rio Branco em setembro/2026 (relatório de transações, 44 lançamentos; aluguel de R$ 11.000,60 e condomínio de R$ 1.460,66 pagos em setembro; sem pró-labore, retiradas, empréstimos, aluguel do depósito, equipamentos, insumos e tributos), total de R$ ' + FMTBR(CF_TOT) + '. Esse total é dividido pela área das 6 salas produtivas do inventário (77,18 m²) × horas ocupadas de cada sala (22 dias × 10 horas × 65%), o que dá o custo de 1 m² por hora; × 9,45 m² da sala de procedimento = custo real de 1 sala por hora.'),
       ('8', 'Ocupação da sala — custo negociado: o mesmo de Cruzeiro do Sul para esta bateria de procedimentos, R$ 40,00 por hora de sala. Horas: 1 hora por consulta + visitas do tratamento fechado (tricologia 11, pescoço 8 e facial 8 visitas de 30 minutos; corporal 16 visitas de 1 hora). A diferença para o custo real é um desconto concedido pela clínica.'),
       ('9', 'IRPJ e CSLL: Lucro Presumido. Consultas: 7,68% (32% × 24%). Procedimentos: 2,28% (8% × 15% + 12% × 9%).')]
r = 4
for k, t in TXT:
    c = ws.cell(r, 1, k); c.font = F(14, True, PET); c.alignment = Alignment(vertical='top', horizontal='center')
    c = ws.cell(r, 2, t); c.font = F(13); c.alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[r].height = 20 * (len(t) // 100 + 1) + 6; r += 1
for s in wb.worksheets: s.sheet_properties.tabColor = {'Resumo': '0A6A70', 'Custo da sala': '1BAF7A', 'Insumos': '1BAF7A', 'Despesas de RB': '7FCFB0'}.get(s.title, '9FB3B6')
for s_ in wb.worksheets:
    for row in s_.iter_rows(min_row=3):
        if any(c.value not in (None, '') for c in row) and s_.row_dimensions[row[0].row].height is None: s_.row_dimensions[row[0].row].height = 24
wb.save(OUT); print('ok', OUT, 'CF', CF_TOT)
