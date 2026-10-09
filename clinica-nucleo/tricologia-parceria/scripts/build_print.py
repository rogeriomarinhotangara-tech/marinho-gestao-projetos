"""Impressão aba por aba da planilha simples (PDF com capa no estilo do relatório dos painéis) + JSON de blocos para o Word.
Uso: python3 -I build_print.py <scratch> <planilha_formatos.xlsx> <planilha_recalculada.xlsx> <saida_dir> <nome_base>"""
import sys, json, html, datetime, subprocess, os
import openpyxl
SP, FMT, VAL, OUTD, BASE = sys.argv[1:6]
os.makedirs(OUTD, exist_ok=True)
wf = openpyxl.load_workbook(FMT); wv = openpyxl.load_workbook(VAL, data_only=True)
DATA_REF = datetime.date.today().strftime('%d/%m/%Y')

def nfmt(v, d=2):
    s = f'{abs(v):,.{d}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return ('−' if v < 0 and round(abs(v), d) > 0 else '') + s
def fmtv(v, f):
    if v is None or v == '': return ''
    if isinstance(v, str): return v
    if isinstance(v, bool): return 'SIM' if v else 'NÃO'
    if isinstance(v, datetime.time): return v.strftime('%H:%M')
    if isinstance(v, datetime.datetime): return v.strftime('%H:%M') if 'h' in f else v.strftime('%d/%m/%Y')
    f = f or 'General'
    if f == 'hh:mm':
        m = round(float(v) * 24 * 60); return f'{m // 60:02d}:{m % 60:02d}'
    if '"R$"' in f:
        if abs(v) < 0.005 and '"–"' in f: return '–'
        return ('−' if v < -0.004 else '') + 'R$ ' + nfmt(abs(v))
    if f.endswith('%'):
        d = len(f.split('.')[1].rstrip('%')) if '.' in f else 0
        return nfmt(v * 100, d) + '%'
    if '" h"' in f: return nfmt(v, 1) + ' h'
    if '" min"' in f: return nfmt(v, 0) + ' min'
    if '" pac."' in f: return nfmt(v, 0) + ' pac.'
    if f == '0': return nfmt(v, 0)
    if f.startswith('#,##0.0'): return nfmt(v, 1)
    if float(v).is_integer(): return nfmt(v, 0)
    return nfmt(v, 2).rstrip('0').rstrip(',')
def rgb(c):
    try:
        x = c.rgb if c is not None else None
        return ('#' + x[-6:]) if isinstance(x, str) and len(x) >= 6 and x not in ('00000000',) else None
    except Exception: return None

ORDER = ['Resumo', 'Receitas', 'Despesas', 'Insumos', 'Custo da sala', 'Despesas de CZS', 'Viagem', 'Como foi feito']
SHEETS = []
for si, sn in enumerate(ORDER, 1):
    fs, vs = wf[sn], wv[sn]
    covered, span = set(), {}
    for mr in fs.merged_cells.ranges:
        span[(mr.min_row, mr.min_col)] = mr.max_col - mr.min_col + 1
        for rr in range(mr.min_row, mr.max_row + 1):
            for cc in range(mr.min_col, mr.max_col + 1):
                if (rr, cc) != (mr.min_row, mr.min_col): covered.add((rr, cc))
    widths = [fs.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 10 for c in range(1, 12)]
    breaks = {b.id + 1 for b in fs.row_breaks.brk}
    maxr = fs.max_row
    def bordered(r):
        return any(fs.cell(r, c).border.left.style or fs.cell(r, c).border.bottom.style for c in range(1, 12) if (r, c) not in covered)
    def cinfo(r, c):
        fc, v = fs.cell(r, c), vs.cell(r, c).value
        fill = rgb(fc.fill.fgColor) if fc.fill and fc.fill.fill_type == 'solid' else None
        txt = fmtv(v, fc.number_format)
        al = fc.alignment.horizontal or ('right' if isinstance(v, (int, float)) and not isinstance(v, bool) else 'left')
        return dict(t=txt, span=span.get((r, c), 1), b=bool(fc.font.b), i=bool(fc.font.i), sz=fc.font.sz or 13,
                    col=rgb(fc.font.color) or '#10272B', fill=fill, al=al, num=isinstance(v, (int, float)) and not isinstance(v, bool))
    blocks, tbl = [], None
    r = 3
    title = vs.cell(1, 1).value or sn; sub = vs.cell(2, 1).value or ''
    while r <= maxr:
        if r in breaks and blocks: blocks.append(dict(k='pb'))
        if bordered(r):
            rows = []
            while r <= maxr and bordered(r):
                if r in breaks and rows: break
                maxc = max(c + span.get((r, c), 1) - 1 for c in range(1, 12) if (fs.cell(r, c).border.left.style or fs.cell(r, c).border.bottom.style) and (r, c) not in covered)
                cells = []; c = 1
                while c <= maxc:
                    if (r, c) in covered: c += 1; continue
                    ci = cinfo(r, c); cells.append(ci); c += ci['span']
                hdr = any(x['fill'] and x['fill'].upper() == '#0A6A70' for x in cells)
                rows.append(dict(cells=cells, hdr=hdr, maxc=sum(x['span'] for x in cells), h=fs.row_dimensions[r].height or 22))
                r += 1
            nc = max(x['maxc'] for x in rows)
            for x in rows:  # completa linhas mais curtas
                used = sum(c['span'] for c in x['cells'])
                for _ in range(nc - used): x['cells'].append(dict(t='', span=1, b=False, i=False, sz=13, col=None, fill=None, al='left', num=False, blank=True))
            blocks.append(dict(k='table', w=widths[:nc], rows=rows))
            continue
        vals = [(c, vs.cell(r, c).value) for c in range(1, 12) if (r, c) not in covered and vs.cell(r, c).value not in (None, '')]
        if vals:
            fc = fs.cell(r, vals[0][0])
            if len(vals) >= 2:
                a = fmtv(vals[0][1], fs.cell(r, vals[0][0]).number_format); b = fmtv(vals[1][1], fs.cell(r, vals[1][0]).number_format)
                blocks.append(dict(k='kv', a=a, b=b, num=sn == 'Como foi feito'))
            else:
                t = fmtv(vals[0][1], fc.number_format); sz = fc.font.sz or 13
                if fc.font.b and sz >= 14: blocks.append(dict(k='h', t=t, lvl=1))
                elif fc.font.b and sz >= 13 and not fc.font.i: blocks.append(dict(k='h', t=t, lvl=2))
                elif fc.font.b: blocks.append(dict(k='p', t=t, b=True))
                else: blocks.append(dict(k='p', t=t, i=bool(fc.font.i)))
        r += 1
    SHEETS.append(dict(n=si, name=sn, title=title, sub=sub, land=fs.page_setup.orientation == 'landscape', blocks=blocks))

# ---------------- valores da capa (do Resumo)
def nm(n):
    s, c = list(wv.defined_names[n].destinations)[0]; return wv[s][c.replace('$', '')].value
rs = wv['Resumo']; K = {}
for row in rs.iter_rows(min_row=6):
    a, b = row[0].value, row[1].value
    if isinstance(a, str) and b is not None: K[a.strip()] = b
CAPA = dict(rec=K['RECEITA BRUTA TOTAL'], desp=K['TOTAL DAS DESPESAS'], res=K['RESULTADO PARA DIVIDIR'], pat=K['DRA. PATRÍCIA RECEBE'], cli=K['CLÍNICA NÚCLEO S RECEBE'],
            sala_real=nm('SALA_HORA'), sala_neg=nm('SALA_NEG'), desc=nm('SALA_DESC'), data=DATA_REF)
json.dump(dict(capa=CAPA, sheets=SHEETS), open(f'{OUTD}/{BASE}.json', 'w', encoding='utf-8'), ensure_ascii=False)

# ---------------- HTML → PDF
brl = lambda v: ('−' if v < 0 else '') + 'R$ ' + nfmt(abs(v))
esc = lambda s: html.escape(str(s))
PT = {11: 7.6, 12: 8.1, 13: 8.6, 14: 9.4, 15: 10.0, 16: 11.0, 20: 13}
CSS = open(SP + '/fonts/local.css', encoding='utf-8').read() + r'''
@page{size:A4;margin:15mm 13mm 17mm 13mm;
  @bottom-left{content:"Clínica Núcleo S · Parceria de Tricologia · Cruzeiro do Sul · outubro/2026";font:7.5pt "IBM Plex Sans",sans-serif;color:#6b7e82}
  @bottom-right{content:"Página " counter(page) " de " counter(pages);font:7.5pt "IBM Plex Sans",sans-serif;color:#6b7e82}}
@page:first{margin:0;@bottom-left{content:none}@bottom-right{content:none}}
@page land{size:A4 landscape;margin:12mm 13mm 15mm 13mm}
.land{page:land}
:root{--ink:#10272b;--ink2:#3d5357;--mut:#6b7e82;--line:#d9e3e3;--soft:#f3f7f7;--acc:#0a6a70}
*{box-sizing:border-box}html,body{margin:0;padding:0}
body{font:8.6pt/1.4 "IBM Plex Sans",sans-serif;color:var(--ink);-webkit-print-color-adjust:exact;print-color-adjust:exact;font-variant-numeric:tabular-nums}
h1,h2,h3,h4{font-family:"Manrope",sans-serif;margin:0;color:var(--ink)}
.sec{break-before:page}
.eyebrow{font:700 7.5pt "Manrope",sans-serif;letter-spacing:.12em;text-transform:uppercase;color:var(--acc);margin-bottom:1mm}
h2{font-size:15pt;font-weight:800;margin:0 0 1.5mm}
.lead{color:var(--ink2);margin:0 0 4mm;max-width:250mm}
h3{font-size:10.5pt;font-weight:800;color:var(--acc);margin:4mm 0 1.5mm;break-after:avoid}
h4{font-size:9.2pt;font-weight:700;color:var(--acc);margin:3mm 0 1.2mm;break-after:avoid}
p{margin:0 0 2mm}p.i{color:var(--mut);font-style:italic;font-size:7.8pt}p.b{font-weight:700}
.kv{display:flex;gap:3mm;padding:.5mm 0;border-bottom:1px solid var(--line);font-size:8pt;color:var(--ink2)}.kv b{color:#0f7f59;min-width:22mm}
.num{display:grid;grid-template-columns:8mm 1fr;gap:2mm;margin:0 0 3mm;font-size:9pt}.num b{font:800 11pt "Manrope";color:var(--acc)}
table{width:100%;border-collapse:collapse;margin:1mm 0 2.5mm;table-layout:fixed}
td,th{border:1px solid #c9d4d4;padding:.85mm 1.6mm;vertical-align:middle;overflow-wrap:anywhere;line-height:1.3}
td.blank{border:0;background:none}
.cmp td,.cmp th{padding:.45mm 1.4mm;line-height:1.2}
th{background:var(--soft);color:var(--ink2);font:600 7.8pt "IBM Plex Sans";text-align:center}
th:first-child{text-align:left}
tr{break-inside:avoid}
.pb{break-before:page}
.legend{display:inline-block;margin-top:2mm;font-size:7.6pt;color:var(--mut)}.legend i{display:inline-block;width:3mm;height:3mm;background:#fff2cc;border:1px solid #e0cf98;vertical-align:-.4mm;margin-right:1.2mm}
.cover{height:297mm;width:210mm;background:#0c2a2f;color:#e6f0f0;padding:26mm 20mm 18mm;display:flex;flex-direction:column}
.cover .mark{width:16mm;height:16mm;border-radius:4mm;background:#0a6a70;display:grid;place-items:center;font:800 15pt "Manrope";color:#fff}
.cover h1{color:#fff;font-size:27pt;font-weight:800;line-height:1.1;margin:13mm 0 4mm;letter-spacing:-.01em}
.cover .sub{font-size:11.5pt;color:#a9c4c6;max-width:160mm}
.cover .kp{display:grid;grid-template-columns:1fr 1fr;gap:5mm;margin-top:14mm}
.cover .kp div{border:1px solid #1d4a51;border-radius:3mm;padding:4.5mm 6mm;background:#10353b}
.cover .kp small{display:block;color:#8fb0b3;font-size:8.5pt}.cover .kp b{font:800 16pt "Manrope";color:#fff}
.cover .kp .pat b{color:#9cc6ff}.cover .kp .cli b{color:#8fe0bd}
.cover .nt{margin-top:6mm;font-size:9pt;color:#cfe0e0;border-left:2px solid #1baf7a;padding-left:4mm}
.cover .toc{margin-top:auto;border-top:1px solid #1d4a51;padding-top:6mm;columns:2;column-gap:10mm;font-size:9pt;color:#cfe0e0}
.cover .toc div{break-inside:avoid;margin-bottom:1.6mm}.cover .toc b{color:#fff}
.cover .foot{margin-top:6mm;font-size:8pt;color:#7f9ea1}
'''
H = ['<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>Parceria de Tricologia — Cruzeiro do Sul — planilha</title><style>' + CSS + '</style></head><body>']
C = CAPA
H.append(f'''<section class="cover"><div class="mark">NS</div>
<h1>Parceria de Tricologia<br>Fechamento de Cruzeiro do Sul</h1>
<div class="sub">Clínica Núcleo S × Dra. Patrícia Fabrini · atendimentos de 05 a 07/10/2026 · receitas, despesas, insumos, custo da sala, viagem e a base de cálculo de cada valor, aba por aba.</div>
<div class="kp"><div><small>Receita bruta</small><b>{brl(C['rec'])}</b></div><div><small>Total das despesas descontadas</small><b>{brl(C['desp'])}</b></div>
<div><small>Resultado para dividir</small><b>{brl(C['res'])}</b></div><div class="pat"><small>Dra. Patrícia · 50%</small><b>{brl(C['pat'])}</b></div></div>
<div class="nt">Sala: custo real de {brl(C['sala_real'])} por hora; nesta bateria de procedimentos foi negociado {brl(C['sala_neg'])} por hora (desconto de {brl(C['desc'])} concedido pela clínica).</div>
<div class="toc">{''.join(f'<div><b>{s["n"]}</b> · {esc(s["name"])}</div>' for s in SHEETS)}</div>
<div class="foot">Controladoria · Rogério Marinho · gerado em {C['data']} · fonte: relatórios de transações do sistema, folhas de procedimento e ajustes da diretoria. Em amarelo, os valores que podem ser alterados na planilha.</div></section>''')
def td(c, hdr, wsum, wcol):
    sty = [f'font-size:{PT.get(int(c["sz"]), 8.6)}pt', f'text-align:{ {"center": "center", "right": "right"}.get(c["al"], "left") }']
    if c['b']: sty.append('font-weight:700')
    if c['i']: sty.append('font-style:italic')
    if not hdr:
        if c['fill'] and c['fill'].upper() not in ('#FFFFFF',): sty.append(f'background:{c["fill"]}')
        if c['col'] and c['col'].upper() not in ('#10272B', '#000000'): sty.append(f'color:{c["col"]}')
    tag = 'th' if hdr else 'td'
    if c.get('blank'): return '<td class="blank"></td>'
    return f'<{tag}{" colspan=%d" % c["span"] if c["span"] > 1 else ""} style="{";".join(sty)}">{esc(c["t"]).replace(chr(10), "<br>")}</{tag}>'
for s in SHEETS:
    H.append(f'<section class="sec{" land" if s["land"] else ""}{" cmp" if s["name"] == "Despesas de CZS" else ""}"><div class="eyebrow">Aba {s["n"]} · {esc(s["name"])}</div><h2>{esc(s["title"])}</h2><p class="lead">{esc(s["sub"])}</p>')
    pend_pb = False
    for b in s['blocks']:
        cls = ' class="pb"' if pend_pb else ''
        if b['k'] == 'pb': pend_pb = True; continue
        if b['k'] == 'h': H.append(f'<h{3 if b["lvl"] == 1 else 4}{cls}>{esc(b["t"])}</h{3 if b["lvl"] == 1 else 4}>')
        elif b['k'] == 'p': H.append(f'<p class="{"b" if b.get("b") else "i" if b.get("i") else ""}{" pb" if pend_pb else ""}">{esc(b["t"])}</p>')
        elif b['k'] == 'kv':
            H.append(f'<div class="num{" pb" if pend_pb else ""}"><b>{esc(b["a"])}</b><div>{esc(b["b"])}</div></div>' if b['num'] else f'<div class="kv{" pb" if pend_pb else ""}"><span style="flex:1">{esc(b["a"])}</span><b>{esc(b["b"])}</b></div>')
        else:
            ws_ = sum(b['w']); cg = ''.join(f'<col style="width:{w / ws_ * 100:.2f}%">' for w in b['w'])
            H.append(f'<table{cls}><colgroup>{cg}</colgroup>' + ''.join('<tr>' + ''.join(td(c, row['hdr'], ws_, None) for c in row['cells']) + '</tr>' for row in b['rows']) + '</table>')
        pend_pb = False
    H.append('</section>')
H.append('</body></html>')
open(f'{OUTD}/{BASE}.html', 'w', encoding='utf-8').write(''.join(H))
pdf = f'{OUTD}/{BASE}.pdf'
subprocess.run(['/opt/pw-browsers/chromium', '--headless', '--no-sandbox', '--disable-gpu', '--no-pdf-header-footer', '--virtual-time-budget=8000',
                f'--print-to-pdf={pdf}', 'file://' + os.path.abspath(f'{OUTD}/{BASE}.html')], check=True, capture_output=True, timeout=180)
print('ok', pdf, len(SHEETS), 'abas')
