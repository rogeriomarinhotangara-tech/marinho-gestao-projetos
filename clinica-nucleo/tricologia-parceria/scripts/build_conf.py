import sys, datetime as dt, pandas as pd
sys.path.insert(0, sys.argv[2])
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
from xl_style import *
from rules2 import norm_desc, DEPARA_CATEGORIA, PALAVRAS
SP=sys.argv[1]; OUT=sys.argv[3]
b=pd.read_pickle(SP+'/ref/base_v2_out.pkl')
T=b[b.Tricologia=='SIM'].copy()
_pn=T['Descrição'].astype(str).str.extract(r'\((\d+)/(\d+)\)\s*$'); T['_n']=pd.to_numeric(_pn[0]).fillna(0)
T['_base']=T['Descrição'].astype(str).str.replace(r'\s*\(\d+/\d+\)\s*$','',regex=True)
T=T.sort_values(['Mes','R/D','Data de vencimento','_base','_n'])
MN={202604:'Abril',202605:'Maio',202606:'Junho',202607:'Julho',202608:'Agosto',202609:'Setembro'}
COMO={'R1 Consultas':'Receita da parceria — consulta','R2 Procedimentos de tricologia':'Receita da parceria — tricologia',
 'R3 Facial e corporal (Dra. Patrícia)':'Receita da parceria — facial/corporal (a confirmar se entra)','C1 Insumos e ativos':'Custo — insumo, desconta no mês',
 'D1 Viagens e deslocamento':'Custo — viagem/estadia, desconta no mês','D3 Outras despesas diretas':'Custo — despesa direta, desconta no mês',
 'X1 Equipamentos (CAPEX)':'Investimento — aparelho, volta em 24 parcelas','X2 Consultoria de implantação':'Investimento — consultoria, volta em 24 parcelas',
 'T2 Reembolsos a pacientes':'Dedução — reembolso','D2 Marketing da tricologia':'Custo — marketing, desconta no mês'}
P=pd.read_pickle(SP+'/ref/plano_v2.pkl').set_index('cod')
kwmot={k:m for k,c,m in PALAVRAS}
def reclass(x):
    if x['R/D']=='Receita':
        if x['Código']=='1.02.02.006': return 'Sim — no sistema é "Receita"; é estorno do microscópio comprado em 29/06'
        if x['Código']=='1.01.03.001' and str(x['Serviços']).strip().upper()=='CONSULTA': return 'Sim — serviço "Consulta" com Responsável Patrícia: consulta de tricologia'
        return 'Não'
    cat=x['Categorias']
    if not isinstance(cat,str): return 'Categoria não veio no export jan–jun — conferir no sistema'
    c0=DEPARA_CATEGORIA.get(' '.join(cat.split()).upper())
    if c0==x['Código']: return 'Não'
    nd=norm_desc(x['Descrição']); mot=[m for k,c,m in PALAVRAS if k in nd and c==x['Código']]
    return f'Sim — categoria do sistema "{cat}"; reclassificado: {mot[-1] if mot else P.at[x["Código"],"nome"]}'
wb=openpyxl.Workbook(); wb.active.title='Como conferir'
HEAD=['Nº','Vencimento','Baixa','Descrição (exata do sistema)','Serviço (sistema)','Categoria (sistema)','Responsável (sistema)','Conta bancária (sistema)','Forma de pagamento (sistema)','Valor no sistema (R$)',
      'Código (nosso plano)','Conta do plano de contas','Como entra na parceria','Reclassificado?','Confere? SIM/NÃO','Valor correto (se divergente)','Conta correta (se divergente)','Diferença (R$)','Observação (Viviane)']
WID=[5,11,11,44,26,30,17,30,22,14,12,40,38,46,11,15,24,13,40]
F_SYS=PatternFill('solid',fgColor='E3EDF8'); F_OUR=PatternFill('solid',fgColor='E1F2EA'); F_CHK=PatternFill('solid',fgColor='FFF2CC')
TOT={}; SHEETS=[]
dv_sn=None
def block(ws,r,df,label,tipo,mes):
    put(ws,r,1,label,font=FT_T2); r+=1
    for c0,c1,t,f in [(4,10,'COMO ESTÁ NO SISTEMA (relatório de transações)',F_SYS),(11,14,'COMO LANÇAMOS NA DRE DA PARCERIA (plano de contas)',F_OUR),(15,19,'CONFERÊNCIA — preencher',F_CHK)]:
        ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c1); c=put(ws,r,c0,t,font=FT_B,fill=f,align=CENTER)
    r+=1; hdr(ws,r,1,HEAD,None,height=34); r+=1; first=r
    if len(df)==0:
        put(ws,r,4,'Nenhum lançamento da parceria neste mês.',font=FT_NOTE); r+=1
    for i,(_,x) in enumerate(df.iterrows(),start=1):
        vals=[i,x['Data de vencimento'],x['Data de baixa'],x['Descrição'],x['Serviços'],x['Categorias'] if isinstance(x['Categorias'],str) else '— não veio no export (conferir)',
              x['Responsável'],x['Conta'],x['Método'],x['Valor'],x['Código'],x['Conta analítica'],COMO.get(x['Linha trico'],x['Linha trico']),reclass(x)]
        for j,v in enumerate(vals,start=1):
            if v is None or (isinstance(v,float) and v!=v): v=None
            c=put(ws,r,j,v,border=B_ALL)
            if j in (2,3) and v is not None: c.number_format=DATE
            if j==10: c.number_format=NUM
            if j in (4,13,14): c.alignment=Alignment(wrap_text=True,vertical='top')
            if j==14 and isinstance(v,str) and v.startswith('Sim'): c.font=FT_RED
        for j in (15,16,17,19): inp(ws,r,j,None,NUM if j==16 else None)
        put(ws,r,18,f'=IF(P{r}="",0,P{r}-J{r})',fmt=NUM,border=B_ALL)
        r+=1
    last=r-1
    put(ws,r,4,f'TOTAL {label.upper()}',font=FT_B,fill=F_TOT,border=B_ALL)
    put(ws,r,10,f'=SUM(J{first}:J{last})' if last>=first else 0,fmt=NUM,font=FT_B,fill=F_TOT,border=B_ALL)
    put(ws,r,18,f'=SUM(R{first}:R{last})' if last>=first else 0,fmt=NUM,font=FT_B,fill=F_TOT,border=B_ALL)
    put(ws,r,15,f'=COUNTIF(O{first}:O{last},"SIM")&" ok / "&COUNTIF(O{first}:O{last},"NÃO")&" div."' if last>=first else '',font=FT_B,fill=F_TOT,border=B_ALL)
    TOT[(mes,tipo)]=dict(sheet=ws.title,total=f'J{r}',dif=f'R{r}',rng=(first,last),n=len(df))
    if last>=first: dv_sn.add(f'O{first}:O{last}')
    return r+2
for mes in sorted(T.Mes.unique()):
    ws=wb.create_sheet(f'{str(mes)[-2:]} {MN[mes][:3]}'); SHEETS.append(ws.title)
    for i,w in enumerate(WID,start=1): ws.column_dimensions[L(i)].width=w
    title(ws,f'Parceria Tricologia — {MN[mes]}/2026: conferência com o sistema','Confira cada linha no sistema. Se estiver igual, escreva SIM na coluna O. Se não, NÃO e preencha o valor/conta corretos e a observação.')
    dv_sn=DataValidation(type='list',formula1='"SIM,NÃO"',allow_blank=True); ws.add_data_validation(dv_sn)
    d=T[T.Mes==mes]; r=4
    r=block(ws,r,d[d['R/D']=='Receita'],f'Receitas de {MN[mes].lower()}','R',mes)
    r=block(ws,r,d[d['R/D']=='Despesa'],f'Despesas de {MN[mes].lower()}','D',mes)
    ws.freeze_panes='E7'
# ---------------- Resumo
ws=wb.create_sheet('Resumo por mês',1)
title(ws,'Resumo da conferência — Parceria Tricologia (abr a set/2026)','Os totais vêm das abas de cada mês. A coluna "Diferença" soma o que a Sra. Viviane apontou como divergente.')
for col,w in zip('ABCDEFGHIJ',[12,16,10,16,10,18,14,14,14,40]): ws.column_dimensions[col].width=w
hdr(ws,4,1,['Mês','Receitas (R$)','Nº','Despesas (R$)','Nº','Conferidos (ok / div.)','Diferença receitas','Diferença despesas','Pendentes','Observação'])
r=5
for mes in sorted(T.Mes.unique()):
    put(ws,r,1,f'{MN[mes]}/26',border=B_ALL,font=FT_B)
    for tipo,cv,cn,cd in (('R',2,3,7),('D',4,5,8)):
        t=TOT.get((mes,tipo))
        if t and t['n']>0:
            put(ws,r,cv,f"='{t['sheet']}'!{t['total']}",fmt=NUM,border=B_ALL); put(ws,r,cn,t['n'],fmt=NUM0,border=B_ALL); put(ws,r,cd,f"='{t['sheet']}'!{t['dif']}",fmt=NUM,border=B_ALL)
        else:
            put(ws,r,cv,0,fmt=NUM,border=B_ALL); put(ws,r,cn,0,fmt=NUM0,border=B_ALL); put(ws,r,cd,0,fmt=NUM,border=B_ALL)
    rngs=[TOT[(mes,tp)] for tp in ('R','D') if (mes,tp) in TOT and TOT[(mes,tp)]['n']>0]
    sh=rngs[0]['sheet'] if rngs else None
    ok='+'.join([f"COUNTIF('{t['sheet']}'!O{t['rng'][0]}:O{t['rng'][1]},\"SIM\")" for t in rngs]) or '0'
    no='+'.join([f"COUNTIF('{t['sheet']}'!O{t['rng'][0]}:O{t['rng'][1]},\"NÃO\")" for t in rngs]) or '0'
    put(ws,r,6,f'=({ok})&" / "&({no})',border=B_ALL); put(ws,r,9,f'=C{r}+E{r}-({ok})-({no})',fmt=NUM0,border=B_ALL); r+=1
put(ws,r,1,'Total',font=FT_B,fill=F_TOT,border=B_ALL)
for c in (2,3,4,5,7,8,9): put(ws,r,c,f'=SUM({L(c)}5:{L(c)}{r-1})',fmt=NUM if c in (2,4,7,8) else NUM0,font=FT_B,fill=F_TOT,border=B_ALL)
r+=2
put(ws,r,1,'Por linha da DRE da parceria (valores como estão hoje)',font=FT_T2); r+=1
lines=[l for l in ['R1 Consultas','R2 Procedimentos de tricologia','R3 Facial e corporal (Dra. Patrícia)','C1 Insumos e ativos','D1 Viagens e deslocamento','D3 Outras despesas diretas','X1 Equipamentos (CAPEX)','X2 Consultoria de implantação'] if l in set(T['Linha trico'])]
meses=sorted(T.Mes.unique())
hdr(ws,r,1,['Linha']+[MN[m][:3]+'/26' for m in meses]+['Total']); r+=1; l0=r
for l in lines:
    put(ws,r,1,COMO.get(l,l),border=B_ALL); ws.cell(r,1).alignment=Alignment(wrap_text=True)
    for k,m in enumerate(meses):
        v=T[(T['Linha trico']==l)&(T.Mes==m)]; v=(v[v['R/D']=='Despesa'].Valor.sum()-v[v['R/D']=='Receita'].Valor.sum()) if l.startswith('X1') else v.Valor.sum()
        put(ws,r,2+k,round(v,2),fmt=NUM,border=B_ALL)
    put(ws,r,2+len(meses),f'=SUM(B{r}:{L(1+len(meses))}{r})',fmt=NUM,border=B_ALL,font=FT_B); ws.row_dimensions[r].height=30; r+=1
ws.column_dimensions['A'].width=46
put(ws,r,1,'Investimento já descontado o estorno do microscópio (R$ 275,92).',font=FT_NOTE)
ws.freeze_panes='A5'
# ---------------- Para localizar
ws=wb.create_sheet('Para localizar',2)
title(ws,'Lançamentos para localizar no sistema — podem ser da parceria e não estão marcados','Regra atual (ajustada em 08/10): passagem aérea, hotel, alimentação e transporte local da Dra. Patrícia são pagos pela clínica e descontados da receita da tricologia.')
for col,w in zip('ABCDEFGHI',[11,11,40,30,30,12,16,16,40]): ws.column_dimensions[col].width=w
r=4
put(ws,r,1,'1) Visitas da Dra. Patrícia identificadas pelas datas de atendimento',font=FT_T2); r+=1
hdr(ws,r,1,['De','Até','Visita','Hotel lançado (R$)','Alimentação/transporte lançados (R$)','','','','O que conferir']); r+=1
for de,ate,lab,hot,ali,obs in [(dt.date(2026,8,10),dt.date(2026,8,19),'1ª visita — agosto (tricologia, facial, corporal e consultas)',1170.00,70.00,'10 dias de atendimento e só R$ 1.170 de hotel e R$ 70 de refeições, sem passagem. Conferir passagens (ida e volta), diárias de hotel, refeições e transporte de todos os dias.'),
                               (dt.date(2026,9,2),dt.date(2026,9,4),'2ª visita? — início de setembro (consultas 02 e 04/09, tricologia 02/09)',0,0,'Nenhuma passagem, hotel ou refeição marcada como da parceria. Confirmar se a Dra. Patrícia esteve na clínica e onde estão os custos.'),
                               (dt.date(2026,9,22),dt.date(2026,9,28),'3ª visita? — fim de setembro (tricologia 22–23/09, consultas 23 e 28/09)',0,0,'Nenhuma passagem nem hotel marcados. Há refeições em CZS nesses dias (abaixo) que podem ser dela.')]:
    put(ws,r,1,de,fmt=DATE,border=B_ALL); put(ws,r,2,ate,fmt=DATE,border=B_ALL); put(ws,r,3,lab,border=B_ALL,align=WRAP); put(ws,r,4,hot,fmt=NUM,border=B_ALL); put(ws,r,5,ali,fmt=NUM,border=B_ALL)
    ws.merge_cells(start_row=r,start_column=9,end_row=r,end_column=9); put(ws,r,9,obs,border=B_ALL,align=WRAP); ws.row_dimensions[r].height=48; r+=1
r+=1
put(ws,r,1,'2) Despesas nesses dias que podem ser da Dra. Patrícia (hoje NÃO estão na parceria)',font=FT_T2); r+=1
hdr(ws,r,1,['Vencimento','Baixa','Descrição (sistema)','Categoria (sistema)','Conta bancária','Valor (R$)','É da Dra. Patrícia? SIM/NÃO','Valor que é dela (R$)','Observação']); r+=1; c0=r
J=b[(b.Mes>=202607)&(b['R/D']=='Despesa')&(b.Tricologia!='SIM')].copy(); J['nd']=J['Descrição'].map(norm_desc)
cand=J[J.nd.isin(['ALMOÇO EQUIPE 09/08','CAFE EQUIPE','COMPRA CZS','COMPRA ALIMENTAÇÃO E LIMPEZA CZS','TAXA EMBARQUE','COMBUSTIVEL','COMPRA REFEIÇÃO CZS','COMPRAS MERCADO CZS','REFEIÇÃO','REFEIÇÃO CZS','ALIMENTAÇAO CZS'])]
cand=cand[(cand['Data de vencimento']>=dt.date(2026,8,7))&(cand['Data de vencimento']<=dt.date(2026,9,30))].sort_values('Data de vencimento')
dv2=DataValidation(type='list',formula1='"SIM,NÃO"',allow_blank=True); ws.add_data_validation(dv2)
for _,x in cand.iterrows():
    put(ws,r,1,x['Data de vencimento'],fmt=DATE,border=B_ALL); put(ws,r,2,x['Data de baixa'],fmt=DATE,border=B_ALL); put(ws,r,3,x['Descrição'],border=B_ALL)
    put(ws,r,4,x['Categorias'],border=B_ALL); put(ws,r,5,x['Conta'],border=B_ALL); put(ws,r,6,x.Valor,fmt=NUM,border=B_ALL)
    inp(ws,r,7,None); inp(ws,r,8,None,NUM); inp(ws,r,9,None); r+=1
dv2.add(f'G{c0}:G{r-1}')
put(ws,r,3,'Total',font=FT_B,fill=F_TOT,border=B_ALL); put(ws,r,6,f'=SUM(F{c0}:F{r-1})',fmt=NUM,font=FT_B,fill=F_TOT,border=B_ALL)
put(ws,r,7,'Confirmado da parceria:',font=FT_B,fill=F_TOT,border=B_ALL); put(ws,r,8,f'=SUM(H{c0}:H{r-1})',fmt=NUM,font=FT_B,fill=F_TOT,border=B_ALL); r+=2
put(ws,r,1,'3) Outros pontos para conferir no sistema',font=FT_T2); r+=1
hdr(ws,r,1,['Item','','O que procurar','','','Valor de referência (R$)','Encontrado? SIM/NÃO','Valor encontrado (R$)','Observação']); r+=1; c1=r
ITENS=[('Passagem, hotel e alimentação de setembro','Passagens, diárias e refeições da Dra. Patrícia nas visitas de 02–04/09 e 22–28/09. Podem estar pagas em dinheiro ou no cartão.',None),
 ('Faturas de cartão (jul–set)','Itens da tricologia dentro das faturas: passagens, hotel, refeição, transporte, produtos. Faturas sem detalhamento: jul R$ 63.372, ago R$ 34.692, set R$ 20.261.',118325.95),
 ('Produtos do facial e corporal','Compra dos produtos usados nos procedimentos facial/corporal da Dra. Patrícia (R$ 67,2 mil de receita em ago–set sem custo identificado).',None),
 ('Passagens aéreas da Dra. Patrícia','Pela regra atual, a clínica paga e desconta da receita da tricologia. Localizar as passagens de ida e volta de cada visita (ago e set). Hoje nenhuma está lançada na parceria — podem estar na fatura do cartão.',None),
 ('MATERIAL TRICOLOGIA R$ 2.744,67','Aparece 3 vezes (27/07, 30/07 e 10/08). Confirmar se são parcelas da mesma compra ou lançamento repetido.',2744.67*3),
 ('CONSULTORIA TRICOLOGIA','Quem recebe os R$ 5.714,28/mês e se a 7ª parcela (out/26) já foi lançada.',40000),
 ('TAXA EMBARQUE R$ 250 (12/08)','De quem é a viagem. Se for da Dra. Patrícia, entra na parceria junto com a passagem.',250),
 ('PENTES TRICOLOGIA R$ 27,06','Quantos pentes vieram (para o custo por sessão da ficha técnica).',27.06)]
dv3=DataValidation(type='list',formula1='"SIM,NÃO"',allow_blank=True); ws.add_data_validation(dv3)
for it,oq,v in ITENS:
    put(ws,r,1,it,border=B_ALL,font=FT_B,align=WRAP); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2)
    put(ws,r,3,oq,border=B_ALL,align=WRAP); ws.merge_cells(start_row=r,start_column=3,end_row=r,end_column=5)
    put(ws,r,6,v,fmt=NUM,border=B_ALL); inp(ws,r,7,None); inp(ws,r,8,None,NUM); inp(ws,r,9,None); ws.row_dimensions[r].height=44; r+=1
dv3.add(f'G{c1}:G{r-1}')
ws.freeze_panes='A4'
# ---------------- Como conferir
ws=wb['Como conferir']; title(ws,'Conferência da Parceria de Tricologia — para a Sra. Viviane (assistente financeira)','Clínica Núcleo · Controladoria (Rogério Marinho) · 08/10/2026')
ws.column_dimensions['A'].width=4; ws.column_dimensions['B'].width=110
r=4
for t in ['Viviane, esta planilha reúne tudo o que lançamos como tricologia. Precisamos que você confira cada linha no sistema antes de apresentarmos a conta à Dra. Patrícia.',
 'Por que conferir: a partilha com a Dra. Patrícia começa nas receitas de OUTUBRO/26. De abril a setembro a receita é 100% da clínica, mas são esses meses que fixam a base de cálculo (insumos, viagens, despesas diretas e o investimento). Se um lançamento estiver errado aqui, o percentual que vamos cobrar também fica errado.',
 'O que é: todos os lançamentos de receita e despesa que entraram na conta da parceria de tricologia (Dra. Patrícia), de abril a setembro de 2026, um por linha, separados por mês.',
 'Cada linha mostra, lado a lado: (azul) como o lançamento está no sistema, exatamente como no relatório de transações; (verde) como ele foi classificado no plano de contas da clínica para a DRE da parceria; (amarelo) as colunas para você preencher.',
 'Passo 1 — Abra a aba do mês (04 Abr a 09 Set). Comece pelas receitas e depois as despesas.',
 'Passo 2 — Procure cada lançamento no sistema pela descrição, data e valor. Se estiver igual, escreva SIM na coluna O (Confere?).',
 'Passo 3 — Se estiver diferente, escreva NÃO e preencha: valor correto (coluna P), conta correta (coluna Q) e o motivo na observação (coluna S). A diferença aparece sozinha na coluna R.',
 'Passo 4 — Se encontrar no sistema um lançamento da tricologia que NÃO está nesta planilha, anote na aba "Para localizar" (seção 3) ou no fim da aba do mês.',
 'Passo 5 — Abra a aba "Para localizar": confirme as despesas das visitas da Dra. Patrícia (hotel, refeições, transporte) e os outros pontos. A despesa de viagem está baixa e precisa ser confirmada.',
 'A coluna "Reclassificado?" mostra quando a conta que usamos é diferente da categoria do sistema. Exemplo: "HOTEL DR PATRICIA" está no sistema como "DESPESA HOSPEDAGEM, ALIMENTAÇÃO E COMBUSTIVEL" e entrou na parceria como hospedagem da parceira (2.02.06.002).',
 'Abril a junho: o relatório de transações dessa época não trazia a coluna de categoria. Nesses meses, confira a categoria direto no sistema.',
 'Regra de viagem (ajustada em 08/10): passagem aérea, hotel, alimentação e transporte local da Dra. Patrícia são pagos pela clínica e descontados da receita da tricologia. Todos entram na conta da parceria.',
 'A partir de outubro, a mesma conferência é feita todo mês, até o dia 10, antes do repasse à Dra. Patrícia.',
 'Ao terminar, a aba "Resumo por mês" mostra quantos lançamentos foram conferidos, quantos divergem e o valor das diferenças. Devolva a planilha preenchida.']:
    put(ws,r,2,t,align=WRAP); ws.row_dimensions[r].height=max(34,15*(len(t)//100+1)+4); r+=1
r+=1; put(ws,r,2,'Legenda',font=FT_T2); r+=1
put(ws,r,2,'Azul = como está no sistema',fill=F_SYS,border=B_ALL); r+=1
put(ws,r,2,'Verde = como lançamos no plano de contas',fill=F_OUR,border=B_ALL); r+=1
put(ws,r,2,'Amarelo = preencher',fill=F_CHK,border=B_ALL,font=FT_IN); r+=1
put(ws,r,2,'Vermelho na coluna "Reclassificado?" = a conta é diferente da categoria do sistema (ver motivo)',font=FT_RED)
for s in wb.worksheets: s.sheet_view.showGridLines=False
wb.save(OUT); print('ok', {k:v['n'] for k,v in TOT.items()}, len(cand))
