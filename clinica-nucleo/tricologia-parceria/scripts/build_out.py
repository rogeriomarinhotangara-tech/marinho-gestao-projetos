"""Fechamento da parceria de Tricologia — outubro/2026 (Cruzeiro do Sul 05–07/10; Rio Branco 08/10 pendente)."""
import sys, json, re, unicodedata, datetime as dt
import pandas as pd, openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
SP=sys.argv[1]; sys.path.insert(0,SP+'/scripts'); OUT=sys.argv[2]; JS=sys.argv[3]
from xl_style import *
o=pd.read_excel(SP+'/data4/czs_out.xlsx')
b=pd.read_pickle(SP+'/ref/base_v2_out.pkl')
def fbr(v): return 'R$ '+f'{v:,.2f}'.replace(',','X').replace('.',',').replace('X','.')
def nz(s): return unicodedata.normalize('NFKD',str(s)).encode('ascii','ignore').decode().upper().strip()
o['PN']=o['Paciente'].map(nz); b['PN']=b['Paciente'].map(nz)
for c in ('Data de emissão','Data de vencimento'): o[c]=pd.to_datetime(o[c]).dt.normalize()
o['Baixa']=pd.to_datetime(o['Data de baixa'],dayfirst=True)
MAP={'TRICOLOGIA':('R2','1.01.03.010','Serviço tricológico executado em parceria'),
     'CONSULTA DRA PATRICIA FABRINI':('R1','1.01.03.001','Consulta dermatológica ou tricológica'),
     'FACIAL DRA PATRICIA':('R3','1.01.04.014','Outros procedimentos estéticos faciais'),
     'CORPORAL DRA PATRICIA':('R3','1.01.05.014','Outros procedimentos estéticos corporais')}
assert set(o['Serviços'])<=set(MAP), set(o['Serviços'])
assert (o['R/D']=='Receita').all() and o['Responsável'].eq('PATRICIA FABRINI').all()
o['LIN']=o['Serviços'].map(lambda s:MAP[s][0]); o['COD']=o['Serviços'].map(lambda s:MAP[s][1]); o['CONTA_N']=o['Serviços'].map(lambda s:MAP[s][2])
bR=b[b['R/D']=='Receita']
def alerts(r):
    a=[]
    if r['Data de emissão']<pd.Timestamp('2026-10-01'):
        m=bR[(bR.PN==r.PN)&(pd.to_datetime(bR['Data de emissão']).dt.normalize()==r['Data de emissão'])]
        if len(m):
            x=m.iloc[0]
            if abs(x.Valor-r.Valor)<0.01: a.append(f"DUPLICIDADE? Já está na base de {x['Data de vencimento']:%d/%m} (mesmo paciente, emissão e valor, {x['Método']}). Conferir se o sistema remarcou ou lançou duas vezes.")
            else: a.append(f"Na base de {x['Data de vencimento']:%d/%m} aparece {fbr(x.Valor)} via {x['Método']}; aqui {fbr(r.Valor)}. Conferir qual é o correto.")
        else: a.append(f"Emitido em {r['Data de emissão']:%d/%m/%Y}, antes de outubro: conferir se a receita é deste mês.")
    if r['Valor líquido']>r['Valor']+0.005: a.append('Valor líquido maior que o bruto: conferir a taxa no sistema.')
    if 'Rio Branco' in str(r['Conta']): a.append('Recebido na conta de Rio Branco, atendimento em Cruzeiro do Sul: conferir o CNPJ da nota fiscal.')
    if 'COBRAR' in nz(r['Observações']): a.append('Observação do sistema: "duas vezes de 2.500,00 não cobrar". Conferir se não há cobrança em dobro.')
    return ' | '.join(a)
o['ALERTA']=o.apply(alerts,axis=1)
o=o.sort_values(['LIN','Paciente','Descrição']).reset_index(drop=True)
# ------------- índices (os mesmos da calculadora)
IX=dict(pis=.0065,cof=.03,iss=0.0,cart=.022,est=0.184628712949474,ins=.0349,irc=.0768,irp=.0228,amort=61139.78/24,pat=.5)
DIAS_CZS,DIAS_RB=3,1; RAT=DIAS_CZS/(DIAS_CZS+DIAS_RB)
PASS=[('05/09/2026','SMILES FIDEL*BILHETE','BARUERI','',1595.00),('05/09/2026','SMILES FIDEL*BILHETE','BARUERI','',961.00),
 ('05/09/2026','SMILES FIDEL*TXAEMBARQ','BARUERI','',56.78),('05/09/2026','SMILES FIDEL*TXAEMBARQ','BARUERI','',35.75),
 ('05/09/2026','SMILES *BILHETE','BARUERI','',796.00),('06/09/2026','SMILES FIDEL*BILHETE','BARUERI','',721.00),
 ('06/09/2026','SMILES FIDEL*TXAEMBARQ','BARUERI','',454.24),('06/09/2026','SMILES FIDEL*TXAEMBARQ','BARUERI','',56.78),
 ('06/09/2026','SMILES FIDEL*BILHETE','BARUERI','PARC 01/02',2160.00)]
PASS2=('—','SMILES FIDEL*BILHETE — 2ª parcela (PARC 02/02), vem na fatura de nov/26','BARUERI','PARC 02/02',2160.00)
P_FATURA=round(sum(x[4] for x in PASS),2); assert abs(P_FATURA-6836.55)<0.005
P_TOT=round(P_FATURA+PASS2[4],2)
# ------------- DRE CZS em Python (conferência das fórmulas e dados do painel)
R={k:round(float(o[o.LIN==k].Valor.sum()),2) for k in ('R1','R2','R3')}; REC=sum(R.values())
assert abs(REC-57750)<0.005 and abs(o.Valor.sum()-REC)<0.005
d=dict(rec=REC,r1=R['R1'],r2=R['R2'],r3=R['R3'],pis=REC*IX['pis'],cof=REC*IX['cof'],iss=REC*IX['iss'],cart=REC*IX['cart'],
       ins=REC*IX['ins'],pass_=P_TOT*RAT,loc=0.0,out=0.0,est=REC*IX['est'],am=IX['amort']*RAT,irc=R['R1']*IX['irc']+(R['R2']+R['R3'])*IX['irp'])
d['rl']=REC-d['pis']-d['cof']-d['iss']-d['cart']; d['ebitda']=d['rl']-d['ins']-d['pass_']-d['loc']-d['out']-d['est']
d['res']=d['ebitda']-d['am']-d['irc']; d['base']=max(0,d['res']); d['pat']=d['base']*IX['pat']; d['cli']=d['base']-d['pat']; d['clitot']=d['cli']+d['est']+d['am']
met=o.groupby('Método').Valor.sum(); d['parc']=float(o[o['Método'].str.contains('parcelado',case=False)].Valor.sum())
d['vista']=float(o[o['Método'].isin(['PIX','Dinheiro'])].Valor.sum()); d['cartvista']=float(o[o['Método'].str.contains('VISTA')].Valor.sum())
d={k:round(v,2) for k,v in d.items()}
print('DRE CZS', d)
# ------------- workbook
wb=openpyxl.Workbook(); wb.active.title='Leia-me'
SH=['Resumo outubro','DRE Cruzeiro do Sul','DRE Rio Branco','Receitas CZS','Viagem','Premissas','Fatura cartão 05-10','Auditoria']
for s in SH: wb.create_sheet(s)
def q(s): return "'"+s+"'"
def name(n,ref): wb.defined_names[n]=DefinedName(n,attr_text=ref)
# Premissas
ws=wb['Premissas']; title(ws,'Premissas do fechamento de outubro/2026','Os mesmos índices da Calculadora e da Base de cálculo. Amarelo = editável.')
for col,w in zip('ABCD',[58,16,4,80]): ws.column_dimensions[col].width=w
PR=[('O_PIS','PIS',IX['pis'],PCT2,'Lei, Lucro Presumido.'),('O_COF','COFINS',IX['cof'],PCT2,'Lei, Lucro Presumido.'),('O_ISS','ISS',IX['iss'],PCT2,'ISS fixo anual: zero no mês (confirmar com a contabilidade).'),
 ('O_CART','Taxas de cartão e antecipação',IX['cart'],PCT2,'Média da clínica. A taxa real deste relatório está na aba Auditoria.'),
 ('O_EST','Taxa de estrutura da clínica',IX['est'],PCT2,'R$ 90.597,90 ÷ R$ 490.703,22 (média jul–set/26).'),
 ('O_INS','Insumos: índice quando não há compra no mês',IX['ins'],PCT2,'Ago–set: R$ 6.281,73 ÷ R$ 179.877. Em outubro não houve compra: o material usado é o estoque comprado em jul–ago, pago 100% pela clínica.'),
 ('O_IRC','IRPJ + CSLL — consultas',IX['irc'],PCT2,'32% × 24%.'),('O_IRP','IRPJ + CSLL — procedimentos',IX['irp'],PCT2,'8% × 15% + 12% × 9%.'),
 ('O_AMORT','Amortização do investimento (R$ por mês)',IX['amort'],NUM,'R$ 61.139,78 ÷ 24. Parcela 1 de 24 (out/26).'),('O_PAT','Parte da Dra. Patrícia',IX['pat'],PCT,'50/50.'),
 ('O_FC','Facial e corporal da Dra. Patrícia entram na partilha? (SIM/NÃO)','SIM',None,'Decisão pendente da diretoria. Hoje: SIM.'),
 ('O_PARC2','Incluir a 2ª parcela do bilhete (R$ 2.160, fatura de nov/26) no custo desta viagem? (SIM/NÃO)','SIM',None,'O bilhete de R$ 4.320 foi parcelado em 2×. A viagem é de outubro: o custo inteiro é de outubro.'),
 ('O_DCZS','Dias de atendimento em Cruzeiro do Sul',DIAS_CZS,NUM0,'05, 06 e 07/10.'),('O_DRB','Dias de atendimento em Rio Branco',DIAS_RB,NUM0,'08/10.')]
r=4
for nm,lab,v,fmt,nt in PR:
    put(ws,r,1,lab,border=B_ALL); inp(ws,r,2,v,fmt); put(ws,r,4,nt,font=FT_NOTE,align=WRAP); name(nm,f"{q('Premissas')}!$B${r}"); r+=1
put(ws,r,1,'Rateio da viagem e da amortização para Cruzeiro do Sul',border=B_ALL,font=FT_B); put(ws,r,2,'=O_DCZS/(O_DCZS+O_DRB)',fmt=PCT,border=B_ALL,font=FT_B)
put(ws,r,4,'Proporção dos dias de atendimento. A viagem SP → Cruzeiro do Sul → Rio Branco → SP atende as duas unidades. A partilha do mês soma as duas, então o rateio só muda a visão por unidade.',font=FT_NOTE,align=WRAP); name('O_RAT',f"{q('Premissas')}!$B${r}"); r+=1
for c_ in ('B14','B15'):
    dv=DataValidation(type='list',formula1='"SIM,NÃO"',allow_blank=False); ws.add_data_validation(dv); dv.add(c_)
# Receitas CZS
ws=wb['Receitas CZS']; title(ws,'Receitas de Cruzeiro do Sul — 05 a 07/10/2026 (relatório de transações)','40 lançamentos, todos da Dra. Patrícia. Colunas Q a S: conferência da Sra. Viviane.')
H=['Nº','Emissão','Vencimento','Baixa','Paciente','Descrição (sistema)','Serviço (sistema)','Forma de pagamento','Conta bancária','Valor (R$)','Valor líquido (R$)','Observação no sistema','Código','Conta do plano','Linha da DRE','Entra na partilha?','Alerta da auditoria','Confere? SIM/NÃO','Observação (Viviane)']
for col,w in zip(range(1,20),[5,11,11,11,30,38,26,30,34,12,12,36,12,32,9,11,60,11,30]): ws.column_dimensions[L(col)].width=w
hdr(ws,4,1,H); r0=5
for i,x in o.iterrows():
    r=r0+i
    vals=[i+1,x['Data de emissão'].to_pydatetime(),x['Data de vencimento'].to_pydatetime(),x['Baixa'].to_pydatetime(),x['Paciente'],x['Descrição'],x['Serviços'],x['Método'],x['Conta'],float(x.Valor),float(x['Valor líquido']),
          (x['Observações'] if pd.notna(x['Observações']) else None),x.COD,x.CONTA_N,x.LIN,f'=IF(O{r}="R3",IF(O_FC="SIM","SIM","NÃO"),"SIM")',x.ALERTA or None]
    for c,v in enumerate(vals,start=1):
        fmt=DATE if c in (2,3,4) else (NUM if c in (10,11) else None)
        put(ws,r,c,v,fmt=fmt,border=B_ALL,align=WRAP if c in (6,12,17) else None,fill=F_WARN if (c==17 and v) else None)
    inp(ws,r,18,None); inp(ws,r,19,None)
rN=r0+len(o)-1
dv=DataValidation(type='list',formula1='"SIM,NÃO"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'R{r0}:R{rN}')
put(ws,rN+1,9,'Total',font=FT_B,fill=F_TOT,border=B_ALL); put(ws,rN+1,10,f'=SUM(J{r0}:J{rN})',fmt=NUM,font=FT_B,fill=F_TOT,border=B_ALL); put(ws,rN+1,11,f'=SUM(K{r0}:K{rN})',fmt=NUM,font=FT_B,fill=F_TOT,border=B_ALL)
put(ws,rN+1,17,f'=COUNTIF(Q{r0}:Q{rN},"?*")&" lançamentos com alerta"',font=FT_B)
put(ws,rN+1,18,f'=COUNTIF(R{r0}:R{rN},"SIM")&" ok / "&COUNTIF(R{r0}:R{rN},"NÃO")&" div."',font=FT_B)
for nm,col in [('OR_VAL','J'),('OR_LIQ','K'),('OR_MET','H'),('OR_LIN','O'),('OR_CTA','I'),('OR_ALE','Q')]: name(nm,f"{q('Receitas CZS')}!${col}${r0}:${col}${rN}")
ws.freeze_panes='F5'; ws.auto_filter.ref=f'A4:S{rN}'
# Viagem
ws=wb['Viagem']; title(ws,'Viagem da Dra. Patrícia — outubro/2026: São Paulo → Cruzeiro do Sul → Rio Branco → São Paulo','Passagens pagas no cartão da clínica (Mastercard final 5917, fatura de 05/10/2026). Regra: a clínica paga e desconta da receita da tricologia.')
for col,w in zip('ABCDEFG',[12,58,12,12,14,12,40]): ws.column_dimensions[col].width=w
hdr(ws,4,1,['Data','Descrição na fatura','Local','Parcela','Valor (R$)','Incluir?','Observação']); r=5; v0=r
for x in PASS:
    for c,v in enumerate([x[0],x[1],x[2],x[3],x[4]],start=1): put(ws,r,c,v,fmt=NUM if c==5 else None,border=B_ALL)
    inp(ws,r,6,'SIM'); put(ws,r,7,'Fatura de 05/10/2026',font=FT_NOTE); r+=1
for c,v in enumerate([PASS2[0],PASS2[1],PASS2[2],PASS2[3],PASS2[4]],start=1): put(ws,r,c,v,fmt=NUM if c==5 else None,border=B_ALL,fill=F_PART)
put(ws,r,6,'=O_PARC2',border=B_ALL,fill=F_PART); put(ws,r,7,'Estimada: confirmar na fatura de novembro',font=FT_NOTE); r+=1; v1=r-1
dv=DataValidation(type='list',formula1='"SIM,NÃO"',allow_blank=False); ws.add_data_validation(dv); dv.add(f'F{v0}:F{v1-1}')
put(ws,r,2,'Passagens já pagas na fatura de 05/10',border=B_ALL); put(ws,r,5,f'=SUMIFS(E{v0}:E{v1-1},F{v0}:F{v1-1},"SIM")',fmt=NUM,border=B_ALL); r+=1
put(ws,r,2,'Total das passagens desta viagem',font=FT_B,fill=F_TOT,border=B_ALL); put(ws,r,5,f'=SUMIFS(E{v0}:E{v1},F{v0}:F{v1},"SIM")',fmt=NUM,font=FT_B,fill=F_TOT,border=B_ALL); name('V_PASS',f"{q('Viagem')}!$E${r}"); rp=r; r+=1
put(ws,r,2,'→ Cruzeiro do Sul (rateio por dias de atendimento)',border=B_ALL); put(ws,r,4,'=O_RAT',fmt=PCT,border=B_ALL); put(ws,r,5,f'=E{rp}*O_RAT',fmt=NUM,border=B_ALL); r+=1
put(ws,r,2,'→ Rio Branco',border=B_ALL); put(ws,r,4,'=1-O_RAT',fmt=PCT,border=B_ALL); put(ws,r,5,f'=E{rp}*(1-O_RAT)',fmt=NUM,border=B_ALL); r+=2
put(ws,r,1,'Hotel, alimentação e transporte local (preencher com os comprovantes)',font=FT_T2); r+=1
hdr(ws,r,1,['Unidade','Despesa','','','Valor (R$)','','Comprovante']); r+=1; h0=r
for u,lab in [('Cruzeiro do Sul','Hotel'),('Cruzeiro do Sul','Alimentação'),('Cruzeiro do Sul','Transporte local / transfer'),('Rio Branco','Hotel'),('Rio Branco','Alimentação'),('Rio Branco','Transporte local / transfer')]:
    put(ws,r,1,u,border=B_ALL); put(ws,r,2,lab,border=B_ALL); inp(ws,r,5,0,NUM); inp(ws,r,7,None); r+=1
name('V_LOC_CZS',f"{q('Viagem')}!$E${h0}:$E${h0+2}"); name('V_LOC_RB',f"{q('Viagem')}!$E${h0+3}:$E${h0+5}")
put(ws,r+1,1,'Nenhuma despesa de hotel, alimentação ou transporte veio no relatório de Cruzeiro do Sul. Se houve, lançar aqui e no sistema com "TRICOLOGIA" ou "PATRICIA" na descrição.',font=FT_RED)
put(ws,r+2,1,'Compradas em 05 e 06/09. Confirmar que nenhum bilhete é da visita de 22 a 28/09 (setembro foi 100% da clínica).',font=FT_NOTE)
# DRE por unidade
def dre(ws,unit,rec_src,loc_rng,rat_expr,ins_cell_note):
    title(ws,f'DRE da parceria de Tricologia — {unit} — outubro/2026',('Atendimentos de 05 a 07/10/2026. Partilha 50/50 (primeiro mês da regra).' if unit=='Cruzeiro do Sul' else 'Atendimento de 08/10/2026. Preencher as receitas quando o relatório de Rio Branco chegar.'))
    for col,w in zip('ABCD',[56,16,12,72]): ws.column_dimensions[col].width=w
    hdr(ws,4,1,['Linha','R$','% da receita','Como é calculado'])
    rows=[]
    def row(lab,f,k='n',note=''):
        r=5+len(rows); rows.append(lab)
        fill=F_TOT if k=='t' else (F_OK if k=='k' else None); font=FT_B if k in('t','k') else FT_N
        put(ws,r,1,lab,border=B_ALL,font=font,fill=fill)
        if k=='in': inp(ws,r,2,f,NUM)
        else: put(ws,r,2,f,fmt=NUM,border=B_ALL,font=font,fill=fill)
        put(ws,r,3,f'=IFERROR(B{r}/$B$8,0)',fmt=PCT,border=B_ALL,font=font,fill=fill); put(ws,r,4,note,font=FT_NOTE,align=WRAP); return r
    if rec_src:
        r1=row('Receita — consultas da Dra. Patrícia','=SUMIFS(OR_VAL,OR_LIN,"R1")','n','Conta 1.01.03.001 · aba Receitas CZS')
        r2=row('Receita — procedimentos de tricologia','=SUMIFS(OR_VAL,OR_LIN,"R2")','n','Conta 1.01.03.010')
        r3=row('Receita — facial e corporal da Dra. Patrícia','=IF(O_FC="SIM",SUMIFS(OR_VAL,OR_LIN,"R3"),0)','n','Contas 1.01.04.014 e 1.01.05.014 (se a diretoria mantiver SIM)')
    else:
        r1=row('Receita — consultas da Dra. Patrícia',0,'in','Digitar do relatório de Rio Branco de 08/10')
        r2=row('Receita — procedimentos de tricologia',0,'in','')
        r3=row('Receita — facial e corporal da Dra. Patrícia',0,'in','Só entra se a diretoria mantiver SIM')
    rr=row('(=) Faturamento da parceria',f'=B{r1}+B{r2}+B{r3}','t')
    a=row('(−) PIS',f'=-B{rr}*O_PIS','n','0,65% do faturamento'); a2=row('(−) COFINS',f'=-B{rr}*O_COF','n','3% do faturamento')
    a3=row('(−) ISS',f'=-B{rr}*O_ISS','n','ISS fixo anual'); a4=row('(−) Taxas de cartão e antecipação',f'=-B{rr}*O_CART','n','2,2% do faturamento (média da clínica)')
    rl=row('(=) Receita líquida',f'=SUM(B{rr}:B{a4})','t')
    ins_in=row('Insumos comprados no mês (R$) — vazio = índice',None,'in',ins_cell_note)
    ins=row('(−) Insumos e ativos',f'=-IF(B{ins_in}="",B{rr}*O_INS,B{ins_in})','n','Sem compra no mês: consumo estimado do estoque pelo índice de 3,49%')
    ps=row('(−) Passagens aéreas (rateio da viagem)',f'=-V_PASS*{rat_expr}','n','Aba Viagem: total das passagens × rateio por dias de atendimento')
    lc=row('(−) Hotel, alimentação e transporte local',f'=-SUM({loc_rng})','n','Aba Viagem (comprovantes)')
    ot=row('(−) Outras despesas diretas',0,'in','Recepção e ambientação dos dias de atendimento, se houver')
    put(ws,ot,2,0,fmt=NUM,border=B_ALL,fill=F_IN,font=FT_IN)
    es=row('(−) Taxa de estrutura da clínica',f'=-B{rr}*O_EST','n','18,46% do faturamento')
    eb=row('(=) Resultado operacional',f'=B{rl}+B{ins}+B{ps}+B{lc}-B{ot}+B{es}','t')
    am=row('(−) Amortização do investimento (parcela 1 de 24, rateada)',f'=-O_AMORT*{rat_expr}','n','R$ 2.547,49 por mês × rateio por dias')
    ir=row('(−) IRPJ e CSLL',f'=-(B{r1}*O_IRC+(B{r2}+B{r3})*O_IRP)','n','7,68% das consultas + 2,28% dos procedimentos')
    rs=row('(=) Resultado da unidade no mês',f'=B{eb}+B{am}+B{ir}','t')
    bs=row('(=) Base de partilha',f'=MAX(0,B{rs})','t','Se negativo, é compensado no consolidado do mês')
    pt=row('Dra. Patrícia — 50%',f'=B{bs}*O_PAT','k','Repasse contra nota fiscal da PJ dela')
    cl=row('Clínica Núcleo S — 50%',f'=B{bs}*(1-O_PAT)','k','')
    ct=row('Clínica Núcleo S — total (50% + estrutura + amortização)',f'=B{cl}-B{es}-B{am}','k','O que fica com a clínica')
    return dict(r1=r1,r2=r2,r3=r3,rr=rr,rl=rl,ins=ins,ps=ps,lc=lc,ot=ot,es=es,eb=eb,am=am,ir=ir,rs=rs,bs=bs,pt=pt,cl=cl,ct=ct,a=a,a2=a2,a3=a3,a4=a4)
ws=wb['DRE Cruzeiro do Sul']; C=dre(ws,'Cruzeiro do Sul',True,'V_LOC_CZS','O_RAT','Nenhuma compra de material de tricologia em outubro até agora')
r=C['ct']+2
put(ws,r,1,'Como o dinheiro entra',font=FT_T2); r+=1
for lab,f,nt in [('PIX e dinheiro',f'=SUMIFS(OR_VAL,OR_MET,"PIX")+SUMIFS(OR_VAL,OR_MET,"Dinheiro")','Já está no caixa'),
                 ('Cartão de crédito à vista',f'=SUMIFS(OR_VAL,OR_MET,"*VISTA*")','Cai em cerca de 30 dias'),
                 ('Cartão de crédito parcelado',f'=SUMIFS(OR_VAL,OR_MET,"*parcelado*")','Entra em 6 a 12 parcelas (nov/26 a out/27), salvo antecipação')]:
    put(ws,r,1,lab,border=B_ALL); put(ws,r,2,f,fmt=NUM,border=B_ALL); put(ws,r,3,f'=IFERROR(B{r}/$B${C["rr"]},0)',fmt=PCT,border=B_ALL); put(ws,r,4,nt,font=FT_NOTE); r+=1
put(ws,r,1,'A partilha acima segue a regra atual (vendido). Se a diretoria decidir pelo recebido, o cartão parcelado entra mês a mês.',font=FT_NOTE)
ws.freeze_panes='B5'
ws=wb['DRE Rio Branco']; Rb=dre(ws,'Rio Branco',False,'V_LOC_RB','(1-O_RAT)','Preencher se houver compra')
put(ws,Rb['ct']+2,1,'Aguardando o relatório de transações de Rio Branco de 08/10/2026. As passagens e a amortização já aparecem rateadas (25%).',font=FT_RED); ws.freeze_panes='B5'
# Resumo outubro
ws=wb['Resumo outubro']; title(ws,'Parceria de Tricologia — outubro/2026: primeiro mês com partilha 50/50','Cruzeiro do Sul fechado (05–07/10). Rio Branco (08/10) aguardando o relatório.')
for col,w in zip('ABCDE',[56,18,18,18,40]): ws.column_dimensions[col].width=w
hdr(ws,4,1,['Linha','Cruzeiro do Sul (05–07/10)','Rio Branco (08/10)','Outubro — total','Situação'])
LS=[('Faturamento da parceria','rr','t'),('(−) Impostos e taxas de cartão',('a','a2','a3','a4'),'n'),('(−) Insumos e ativos','ins','n'),('(−) Passagens aéreas','ps','n'),
    ('(−) Hotel, alimentação e transporte','lc','n'),('(−) Taxa de estrutura','es','n'),('(−) Amortização do investimento','am','n'),('(−) IRPJ e CSLL','ir','n'),('(=) Resultado','rs','t')]
r=5
for lab,k,kind in LS:
    fill=F_TOT if kind=='t' else None; font=FT_B if kind=='t' else FT_N
    put(ws,r,1,lab,border=B_ALL,font=font,fill=fill)
    for c,sh,M in [(2,'DRE Cruzeiro do Sul',C),(3,'DRE Rio Branco',Rb)]:
        f='='+'+'.join(f"{q(sh)}!B{M[x]}" for x in (k if isinstance(k,tuple) else (k,)))
        put(ws,r,c,f,fmt=NUM,border=B_ALL,font=font,fill=fill)
    put(ws,r,4,f'=B{r}+C{r}',fmt=NUM,border=B_ALL,font=font,fill=fill); r+=1
rres=r-1
put(ws,r,1,'(=) Base de partilha do mês (compensa unidade negativa)',border=B_ALL,font=FT_B,fill=F_TOT); put(ws,r,4,f'=MAX(0,D{rres})',fmt=NUM,border=B_ALL,font=FT_B,fill=F_TOT); rb=r; r+=1
put(ws,r,1,'Dra. Patrícia — 50%',border=B_ALL,font=FT_B,fill=F_OK)
put(ws,r,2,f"={q('DRE Cruzeiro do Sul')}!B{C['pt']}",fmt=NUM,border=B_ALL,fill=F_OK); put(ws,r,3,f"={q('DRE Rio Branco')}!B{Rb['pt']}",fmt=NUM,border=B_ALL,fill=F_OK)
put(ws,r,4,f'=D{rb}*O_PAT',fmt=NUM,border=B_ALL,font=FT_B,fill=F_OK); put(ws,r,5,'Repasse até 15/11, contra nota fiscal',font=FT_NOTE); rpt=r; r+=1
put(ws,r,1,'Clínica Núcleo S — total (50% + estrutura + amortização)',border=B_ALL,font=FT_B,fill=F_OK)
put(ws,r,4,f'=D{rb}*(1-O_PAT)-D{5+5}-D{5+6}',fmt=NUM,border=B_ALL,font=FT_B,fill=F_OK); r+=1
put(ws,r,1,'Dra. Patrícia — % do faturamento',border=B_ALL); put(ws,r,4,f'=IFERROR(D{rpt}/D5,0)',fmt=PCT,border=B_ALL); r+=2
put(ws,r,1,'Situação',font=FT_T2); r+=1
for t in ['Cruzeiro do Sul: fechado com o relatório de 05 a 07/10 (40 lançamentos, R$ 57.750,00). Alertas na aba Auditoria.',
          'Rio Branco: aguardando o relatório de 08/10. Já carrega 25% das passagens e da amortização.',
          'Passagens: R$ 6.836,55 na fatura de 05/10 + R$ 2.160,00 da 2ª parcela (fatura de nov/26) = R$ 8.996,55.']:
    put(ws,r,1,'• '+t,align=WRAP); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=5); ws.row_dimensions[r].height=28; r+=1
# Fatura
ws=wb['Fatura cartão 05-10']; title(ws,'Fatura Mastercard final 5917 — fechamento 28/09/2026, vencimento 05/10/2026 — classificação proposta','Pela regra, fatura de cartão não é rateada: cada item vai para a sua conta. Anotações à mão da foto na coluna G.')
for col,w in zip('ABCDEFGH',[11,32,16,11,12,40,26,14]): ws.column_dimensions[col].width=w
FAT=[('21/04/2026','CLINICA*IMOTION GROUP','São Paulo','PARC 06/12',3750.00,'A confirmar (compra parcelada de abr/26)','anotação à mão','NÃO'),
 ('28/04/2026','MERCADOLIVRE*MERCADOL','Pinhais','PARC 05/05',579.22,'A confirmar (compra parcelada de abr/26)','anotação à mão','NÃO'),
 ('29/04/2026','HAVAN ACRE RIO BRANCO','Rio Branco','PARC 05/05',648.84,'A confirmar (compra parcelada de abr/26)','','NÃO'),
 ('19/05/2026','GMAQ HOME CENTER','Cruzeiro do Sul','PARC 05/10',2931.00,'2.16.01.005 Obras e benfeitorias (a confirmar)','anotação à mão','NÃO'),
 ('25/08/2026','SCP PLUS- AGO/26','','',19.99,'2.09.02.006 Assinaturas e publicações (a confirmar)','anotação à mão','NÃO'),
 ('05/09/2026','MR PAGA*WEB DIET','São Luís','',94.90,'2.08.01.001 Sistema (software de nutrição)','"Sit Nutri"','NÃO')]
FAT+=[(x[0],x[1],'Barueri',x[3],x[4],'2.02.06.001 Passagem de profissional parceiro','"6.836,55" (passagens)','SIM') for x in PASS]
FAT+=[('08/09/2026','DEB AUTOM DE FATURA EM','','',-9784.39,'Pagamento da fatura anterior (não é despesa)','','NÃO'),
 ('08/09/2026','DM*HOSTINGERCOMB','São Paulo','PARC 01/12',78.14,'2.08.01.015 Domínio e hospedagem de site','"Hospedagem site e domínio"','NÃO'),
 ('08/09/2026','FACEBK *A46TU5NMP4','São Paulo','',728.48,'2.05.01.002 Tráfego pago — Meta','"4.279,52" (anúncios)','NÃO'),
 ('13/09/2026','FACEBK *NYPQR52NP4','São Paulo','',704.70,'2.05.01.002 Tráfego pago — Meta','','NÃO'),
 ('22/09/2026','FACEBK *VC62H7ENP4','São Paulo','',2846.34,'2.05.01.002 Tráfego pago — Meta','','NÃO'),
 ('25/09/2026','SCP PLUS- SET/26','','',19.99,'2.09.02.006 Assinaturas e publicações (a confirmar)','','NÃO'),
 ('31/08/2026','FACEBK *UQZXEZD622 (US$ 35,00)','Menlo Park','',192.93,'2.05.01.002 Tráfego pago — Meta','"199,68" com IOF','NÃO'),
 ('31/08/2026','IOF DESPESA NO EXTERIOR','','',6.75,'2.11.01.007 IOF','','NÃO'),
 ('01/09/2026','FACEBK *DEVC2Z5622 (US$ 1,09)','Menlo Park','',6.04,'2.05.01.002 Tráfego pago — Meta','"6,25" com IOF','NÃO'),
 ('01/09/2026','IOF DESPESA NO EXTERIOR','','',0.21,'2.11.01.007 IOF','','NÃO'),
 ('03/09/2026','ANTHROPIC* CLAUDE SUB (US$ 21,61)','San Francisco','',118.55,'2.08.01.007 Nuvem, hospedagem e armazenamento','"122,70" com IOF','NÃO'),
 ('03/09/2026','IOF DESPESA NO EXTERIOR','','',4.15,'2.11.01.007 IOF','','NÃO'),
 ('04/09/2026','SUPABASE (US$ 25,00)','Singapore','',136.32,'2.08.01.007 Nuvem, hospedagem e armazenamento','"armazenamento" · "141,09" com IOF','NÃO'),
 ('04/09/2026','IOF DESPESA NO EXTERIOR','','',4.77,'2.11.01.007 IOF','','NÃO')]
hdr(ws,4,1,['Data','Descrição','Local','Parcela','Valor (R$)','Conta proposta','Anotação na foto','Parceria?']); r=5; f0=r
for x in FAT:
    for c,v in enumerate(x,start=1): put(ws,r,c,v,fmt=NUM if c==5 else None,border=B_ALL,fill=F_OK if x[7]=='SIM' else (F_TOT if x[4]<0 else None))
    r+=1
f1=r-1
put(ws,r,2,'Despesas/débitos da fatura (sem o pagamento anterior)',font=FT_B,border=B_ALL); put(ws,r,5,f'=SUMIFS(E{f0}:E{f1},E{f0}:E{f1},">0")',fmt=NUM,font=FT_B,border=B_ALL); put(ws,r,6,'Na fatura: R$ 19.707,87',font=FT_NOTE); rf=r; r+=1
put(ws,r,2,'Conferência (deve ser zero)',border=B_ALL); put(ws,r,5,f'=ROUND(E{rf}-19707.87,2)',fmt=NUM,border=B_ALL); r+=1
put(ws,r,2,'Passagens da Dra. Patrícia (parceria)',font=FT_B,border=B_ALL,fill=F_OK); put(ws,r,5,f'=SUMIFS(E{f0}:E{f1},H{f0}:H{f1},"SIM")',fmt=NUM,font=FT_B,border=B_ALL,fill=F_OK); r+=2
hdr(ws,r,2,['Grupo','','','Valor (R$)']); r+=1
for g,pat in [('Passagens da parceria (2.02.06.001)','2.02.06.001*'),('Marketing — tráfego pago Meta','2.05.01.002*'),('Tecnologia (sistemas, nuvem, site)','2.08.*'),('IOF','2.11.01.007*'),('Assinaturas (a confirmar)','2.09.02.006*'),('Obras CZS (a confirmar)','2.16.01.005*'),('Compras parceladas de abr/26 (a confirmar)','A confirmar*')]:
    put(ws,r,2,g,border=B_ALL); put(ws,r,5,f'=SUMIFS(E{f0}:E{f1},F{f0}:F{f1},"{pat}")',fmt=NUM,border=B_ALL); r+=1
put(ws,r,2,'Total',font=FT_B,border=B_ALL,fill=F_TOT); put(ws,r,5,f'=SUM(E{r-7}:E{r-1})',fmt=NUM,font=FT_B,border=B_ALL,fill=F_TOT); r+=2
put(ws,r,1,'No relatório de outubro, o pagamento desta fatura (R$ 19.707,87) vai aparecer como fatura de cartão sem detalhamento (2.13.01.099). Esta aba é o detalhamento para reclassificar item a item.',font=FT_NOTE)
ws.freeze_panes='A5'
# Auditoria
ws=wb['Auditoria']; title(ws,'Auditoria do fechamento de Cruzeiro do Sul — feita antes da entrega','Cada item com o que foi conferido e o que falta. Status: OK, ATENÇÃO ou PENDENTE.')
for col,w in zip('ABCDE',[5,62,18,12,70]): ws.column_dimensions[col].width=w
hdr(ws,4,1,['#','Verificação','Valor','Status','Detalhe e ação'])
liq_parc=float(o[o['Método'].str.contains('parcelado',case=False)]['Valor líquido'].sum())
A=[('Total do relatório de transações = soma da aba Receitas CZS','=SUM(OR_VAL)','=IF(ROUND(SUM(OR_VAL)-57750,2)=0,"OK","ATENÇÃO")','40 lançamentos, R$ 57.750,00, conferido ao centavo.'),
 ('Todas as receitas classificadas em uma linha da DRE','=COUNTA(OR_LIN)','=IF(COUNTA(OR_LIN)=40,"OK","ATENÇÃO")','Tricologia R2 (30), consultas R1 (6), facial e corporal R3 (4). Nenhuma receita fora da DRE.'),
 ('DRE de Cruzeiro do Sul: faturamento = total do relatório',f"={q('DRE Cruzeiro do Sul')}!B{C['rr']}",f"=IF(ROUND({q('DRE Cruzeiro do Sul')}!B{C['rr']}-SUM(OR_VAL),2)=0,\"OK\",\"ATENÇÃO\")",'Com facial e corporal na partilha (premissa SIM).'),
 ('Responsável de todos os lançamentos = Dra. Patrícia',40,'OK','Nenhuma receita de outro profissional misturada.'),
 ('Consultas emitidas em ago/set que já estão na base de ago/set','=SUMIFS(OR_VAL,OR_ALE,"DUPLICIDADE*")','ATENÇÃO','5 consultas de R$ 400 (mesmo paciente, data de emissão e valor). Provável remarcação no sistema para 05/10. Mantidas em outubro (vencimento 05/10). Sra. Viviane: confirmar no sistema se é o mesmo lançamento (remarcado) ou duplicado. Se duplicado, excluir um.'),
 ('Consulta no cartão ELO com valor diferente da base e líquido maior que o bruto','=SUMIFS(OR_VAL,OR_MET,"*ELO*")','ATENÇÃO','Na base de set: R$ 400 via PIX em 04/09. Aqui: R$ 500 no cartão ELO, líquido R$ 586,56. Conferir valor e forma de pagamento.'),
 ('Receitas recebidas na conta de Rio Branco (atendimento em CZS)','=SUMIFS(OR_VAL,OR_CTA,"*Rio Branco*")','ATENÇÃO','PIX feitos para o CNPJ de Rio Branco. Entram na DRE de Cruzeiro do Sul (local do atendimento). Conferir com a contabilidade qual CNPJ emite a nota.'),
 ('Observação "duas vezes de 2.500,00 não cobrar" (corporal R$ 5.000)',5000,'ATENÇÃO','Conferir se o valor não foi lançado ou cobrado duas vezes.'),
 ('Taxa real de cartão parcelado neste relatório (bruto − líquido)',f'=SUMIFS(OR_VAL,OR_MET,"*parcelado*")-SUMIFS(OR_LIQ,OR_MET,"*parcelado*")','OK',f'{fbr(33600-liq_parc)} = '+f'{((33600-liq_parc)/33600)*100:.2f}'.replace('.',',')+'% do parcelado, antes de antecipação. A DRE usa a média de 2,2% da política.'),
 ('Passagens: total da fatura conferido com a anotação (R$ 6.836,55)','=V_PASS','ATENÇÃO','9 lançamentos SMILES de 05 e 06/09 somam R$ 6.836,55. O bilhete de R$ 2.160 é PARC 01/02: a 2ª parcela (R$ 2.160) vem em nov/26 e entra no custo desta viagem (premissa SIM). Confirmar que nenhum bilhete é da visita de 22–28/09.'),
 ('Fatura de 05/10 conferida ao centavo',f"={q('Fatura cartão 05-10')}!E{rf}",f"=IF({q('Fatura cartão 05-10')}!E{rf+1}=0,\"OK\",\"ATENÇÃO\")",'R$ 19.707,87 = R$ 19.254,03 em reais + R$ 453,84 em dólar convertido.'),
 ('Hotel, alimentação e transporte em Cruzeiro do Sul',f"=SUM(V_LOC_CZS)",'PENDENTE','Não vieram no relatório. Se houve, lançar na aba Viagem.'),
 ('Insumos de outubro','=0','ATENÇÃO','Nenhuma compra em outubro. O material usado é o estoque de jul–ago, pago 100% pela clínica. A DRE desconta o índice de 3,49% como consumo estimado. Decidir se fica assim até a ficha técnica ter os preços das notas.'),
 ('Pacotes de tricologia vendidos em ago/set com sessões feitas em outubro','—','ATENÇÃO','Ex.: pacotes vendidos em agosto (R$ 108,2 mil em ago–set). A receita ficou em ago/set (100% clínica); as sessões feitas agora não geram receita em outubro. Decisão da diretoria: deixar assim ou partilhar a parte das sessões executadas a partir de outubro.'),
 ('Rio Branco (08/10)','—','PENDENTE','Aguardando o relatório. A DRE de Rio Branco já tem 25% das passagens e da amortização.')]
r=5
for i,(lab,v,st,det) in enumerate(A,start=1):
    put(ws,r,1,i,border=B_ALL); put(ws,r,2,lab,border=B_ALL,align=WRAP); put(ws,r,3,v,fmt=NUM,border=B_ALL); put(ws,r,4,st,border=B_ALL,font=FT_B)
    put(ws,r,5,det,border=B_ALL,align=WRAP,font=FT_NOTE); ws.row_dimensions[r].height=44; r+=1
# Leia-me
ws=wb['Leia-me']; title(ws,'Fechamento da parceria de Tricologia — outubro/2026 — Clínica Núcleo S','Controladoria (Rogério Marinho) · 08/10/2026')
ws.column_dimensions['A'].width=4; ws.column_dimensions['B'].width=120; r=4
for t in ['Outubro é o primeiro mês com partilha 50/50. Esta planilha fecha Cruzeiro do Sul (atendimentos de 05 a 07/10) e já deixa a DRE de Rio Branco (08/10) pronta para receber o relatório.',
 'Resumo outubro: Cruzeiro do Sul + Rio Branco + total do mês, com a parte da Dra. Patrícia.',
 'DRE Cruzeiro do Sul: receitas do relatório, impostos, cartão, insumos, passagens (rateadas), estrutura, amortização e IRPJ/CSLL, tudo por fórmula.',
 'Receitas CZS: os 40 lançamentos, com conta do plano, linha da DRE, alerta da auditoria e colunas para a Sra. Viviane conferir.',
 'Viagem: passagens da fatura do cartão (R$ 6.836,55 + 2ª parcela de R$ 2.160), rateio por dias e espaço para hotel e alimentação.',
 'Fatura cartão 05-10: a fatura inteira classificada item a item (passagens, anúncios, sistemas, IOF, compras parceladas).',
 'Premissas: os mesmos índices da Calculadora. Amarelo = editável; todas as abas recalculam.',
 'Auditoria: o que foi conferido antes da entrega e o que ainda depende de confirmação.']:
    put(ws,r,2,'• '+t,align=WRAP); ws.row_dimensions[r].height=30; r+=1
for s in wb.worksheets: s.sheet_view.showGridLines=False
wb.save(OUT)
# ------------- dados do painel (sem nomes de pacientes)
LN={'R1':'Consultas da Dra. Patrícia','R2':'Procedimentos de tricologia','R3':'Facial e corporal da Dra. Patrícia'}
items={}
for k in ('R1','R2','R3'):
    s=o[o.LIN==k]; g=s.groupby(['Serviços','Método']).Valor.agg(['sum','count']).reset_index().sort_values('sum',ascending=False)
    items[k]=[[f"{x['Serviços'].title()} · {x['Método']}",round(float(x['sum']),2),int(x['count'])] for _,x in g.iterrows()]
alert_n=int((o.ALERTA!='').sum())
data=dict(d=d,R=R,items=items,LN=LN,pass_=[list(x) for x in PASS],pass2=list(PASS2),p_fatura=P_FATURA,p_tot=P_TOT,rat=RAT,
          ix=IX,n=len(o),alert_n=alert_n,liq_parc=round(liq_parc,2),dup=round(float(o[o.ALERTA.str.startswith('DUPLICIDADE')].Valor.sum()),2),
          rb_conta=round(float(o[o.Conta.str.contains('Rio Branco')].Valor.sum()),2),
          fat=[[x[0],x[1],x[4],x[5],x[7]] for x in FAT],
          audit=[[lab,st,det] for (lab,v,st,det) in A])
for it in data['audit']:
    if it[1].startswith('='): it[1]='OK'
json.dump(data,open(JS,'w'),ensure_ascii=False)
print('ok', OUT, 'alertas', alert_n, 'dup', data['dup'], 'rb', data['rb_conta'], 'liq_parc', round(liq_parc,2))
