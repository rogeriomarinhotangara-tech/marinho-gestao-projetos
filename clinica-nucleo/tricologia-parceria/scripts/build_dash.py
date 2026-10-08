"""Gera os dados agregados do painel (sem nomes de pacientes, sem detalhe de folha)."""
import sys, json, re
import pandas as pd
SP=sys.argv[1]; OUT=sys.argv[2]
b=pd.read_pickle(SP+'/ref/base_v2_out.pkl')
old=json.load(open(SP+'/out/dash_data.json'))
MESES=[202601,202602,202603,202604,202605,202606,202607,202608,202609]
b=b[b.Mes.isin(MESES)].copy()
b['T']=b['R/D']
def mv(df):
    s=df.groupby('Mes')['Valor'].sum(); return [round(float(s.get(m,0)),2) for m in MESES]
def g(grp,tipo): return mv(b[(b.Grupo==grp)&(b['T']==tipo)])
OPEX=['Pessoal','Comercial e Relacionamento','Marketing','Ocupação e Infraestrutura','Equipamentos e Compliance Clínico','Tecnologia',
      'Administrativo e Profissionais','Taxas e Obrigações Regulatórias','Rateios e Overhead Compartilhado','Despesas sem identificação']
add=lambda *xs:[round(sum(v),2) for v in zip(*xs)]
neg=lambda x:[-v for v in x]
L=[]
rob=g('Receita Operacional Bruta','Receita'); oro=g('Outras Receitas Operacionais','Receita'); ded=neg(g('Deduções da Receita','Despesa'))
rl=add(rob,oro,ded); cd=neg(g('Custos Diretos','Despesa')); mc=add(rl,cd)
L+=[dict(k='rob',l='Receita operacional bruta',t='n',g='Receita Operacional Bruta',v=rob),
    dict(k='oro',l='Outras receitas operacionais',t='n',g='Outras Receitas Operacionais',v=oro),
    dict(k='ded',l='(−) Deduções da receita (impostos sobre faturamento)',t='n',g='Deduções da Receita',v=ded),
    dict(k='rl',l='(=) Receita líquida',t='t',v=rl),
    dict(k='cd',l='(−) Custos diretos (CMV: medicamentos, insumos, honorários, comissões)',t='n',g='Custos Diretos',v=cd),
    dict(k='mc',l='(=) Margem de contribuição',t='t',v=mc)]
op=[]
for gname in OPEX:
    v=neg(g(gname,'Despesa')); op.append(v)
    lab={'Despesas sem identificação':'Faturas de cartão sem detalhamento','Pessoal':'Pessoal (despesa trabalhista)'}.get(gname,gname)
    L.append(dict(k='op_'+gname,l='(−) '+lab,t='n',g=gname,v=v))
ebitda=add(mc,*op)
fin_r=g('Receitas Financeiras','Receita'); fin_d=neg(g('Despesas Financeiras','Despesa'))
nor=g('Outras Receitas / Resultado Não Operacional','Receita'); nod=neg(g('Outras Despesas / Não Operacionais','Despesa'))
da=neg(g('Depreciação e Amortização','Despesa')); irc=neg(g('Tributos sobre Lucro','Despesa'))
lair=add(ebitda,da,fin_r,fin_d,nor,nod); ll=add(lair,irc)
L+=[dict(k='ebitda',l='(=) EBITDA',t='t',v=ebitda),
    dict(k='fin',l='(±) Resultado financeiro',t='n',g='Despesas Financeiras',v=add(fin_r,fin_d)),
    dict(k='nop',l='(±) Não operacional',t='n',v=add(nor,nod)),
    dict(k='irc',l='(−) IRPJ e CSLL',t='n',g='Tributos sobre Lucro',v=irc),
    dict(k='ll',l='(=) Resultado líquido',t='t',v=ll)]
FORA=[('Investimentos - Não DRE','Investimentos (equipamentos, obras)'),('Financiamentos - Não DRE','Empréstimos e parcelamentos (principal)'),
      ('Patrimônio - Não DRE','Retiradas de sócio'),('Compras para Estoque - Não DRE','Compras para estoque'),
      ('Adiantamentos e Ativos - Não DRE','Adiantamentos'),('Contas de Passagem - Não DRE','Transferências entre contas (saídas)')]
fora=[dict(k='f'+str(i),l=lab,g=gg,v=g(gg,'Despesa')) for i,(gg,lab) in enumerate(FORA)]
# conferência ao centavo
tot_d=mv(b[b['T']=='Despesa'])
soma=[round(-(ded[i]+cd[i]+sum(o[i] for o in op)+da[i]+fin_d[i]+nod[i]+irc[i])+sum(f['v'][i] for f in fora),2) for i in range(9)]
assert all(abs(tot_d[i]-soma[i])<0.02 for i in range(9)), (tot_d,soma)
# ------------- razão
SENS_G={'Pessoal','Patrimônio - Não DRE','Adiantamentos e Ativos - Não DRE'}
SENS_C=('2.02.04','2.02.08','2.09.01.004','2.09.01.010','2.17')
def sens(code,grp): return grp in SENS_G or str(code).startswith(SENS_C)
def clean(s):
    s=str(s); s=re.sub(r'^\s*Transação Recorrente:\s*','',s,flags=re.I); s=re.sub(r'\s*\(\d+/\d+\)\s*$','',s)
    return re.sub(r'\s+',' ',s).strip().upper()
def det_key(r):
    if r['T']=='Receita':
        if str(r['Grupo'])=='Receita Operacional Bruta': return 'Serviço: '+(str(r['Serviços']) if pd.notna(r['Serviços']) else 'sem serviço')
        return clean(r['Descrição'])
    if sens(r['Código'],r['Grupo']): return None
    return clean(r['Descrição'])
b['DK']=b.apply(det_key,axis=1)
GORD=['Receita Operacional Bruta','Outras Receitas Operacionais','Deduções da Receita','Custos Diretos']+OPEX+['Despesas Financeiras','Tributos sobre Lucro']+[f[0] for f in FORA]
def detail(cs,code,gg,tipo):
    det={}
    for m,ms in cs.groupby('Mes'):
        if sens(code,gg) and tipo=='Despesa':
            det[str(m)]=dict(mask=1,n=int(len(ms))); continue
        a=ms.groupby('DK')['SV'].agg(['sum','count']).sort_values('sum',ascending=False)
        top=[[k,round(float(v['sum']),2),int(v['count'])] for k,v in a.head(8).iterrows()]
        rest=a.iloc[8:]
        if len(rest): top.append(['Outros ('+str(int(rest['count'].sum()))+' lançamentos)',round(float(rest['sum'].sum()),2),int(rest['count'].sum())])
        det[str(m)]=dict(items=top,n=int(len(ms)))
    return det
def mvs(df):
    s_=df.groupby('Mes')['SV'].sum(); return [round(float(s_.get(m,0)),2) for m in MESES]
def conta_obj(code,cs,gg,tipo):
    return dict(c=str(code),n=str(cs['Conta analítica'].iloc[0]),v=mvs(cs),est=mvs(cs[cs.Estrutura=='SIM']),d=detail(cs,code,gg,tipo))
b['SV']=b['Valor']
raz=[]
for gg in GORD:
    sub=b[b.Grupo==gg]
    if sub.empty: continue
    tipo='Receita' if gg in ('Receita Operacional Bruta','Outras Receitas Operacionais') else 'Despesa'
    sub=sub[sub['T']==tipo]
    contas=[conta_obj(code,cs,gg,tipo) for code,cs in sub.groupby('Código')]
    contas.sort(key=lambda x:-sum(x['v']))
    raz.append(dict(g=gg,tipo=tipo,dre=not gg.endswith('Não DRE'),v=mv(sub),contas=contas))
# receita por unidade × conta
uni_c={}
for u in ['Rio Branco','Cruzeiro do Sul','Epitaciolândia']:
    sub=b[(b.Grupo=='Receita Operacional Bruta')&(b['T']=='Receita')&(b.Unidade==u)]
    cl=[conta_obj(code,cs,'Receita Operacional Bruta','Receita') for code,cs in sub.groupby('Código')]
    uni_c[u]=sorted(cl,key=lambda x:-sum(x['v']))
# razão da tricologia por linha da DRE da parceria
TL=[('R1 Consultas','Receita','Consultas da Dra. Patrícia'),('R2 Procedimentos de tricologia','Receita','Procedimentos de tricologia'),
    ('R3 Facial e corporal (Dra. Patrícia)','Receita','Facial e corporal da Dra. Patrícia'),('C1 Insumos e ativos','Despesa','Insumos e ativos'),
    ('D1 Viagens e deslocamento','Despesa','Viagens e estadia'),('D2 Marketing da tricologia','Despesa','Marketing da tricologia'),
    ('D3 Outras despesas diretas','Despesa','Outras despesas diretas'),('X1 Equipamentos (CAPEX)','Despesa','Aparelhos (investimento)'),
    ('X2 Consultoria de implantação','Despesa','Consultoria de implantação')]
tri={}
tb=b[b.Tricologia=='SIM'].copy()
for key,tipo,lab in TL:
    sub=tb[tb['Linha trico']==key].copy()
    sub['SV']=sub['Valor'].where(sub['T']==tipo,-sub['Valor'])
    cl=[conta_obj(code,cs,'Tricologia',tipo) for code,cs in sub.groupby('Código')]
    tri[key[:2]]=dict(l=lab,tipo=tipo,contas=sorted(cl,key=lambda x:-sum(x['v'])))
assert abs(sum(sum(c['v']) for c in tri['X1']['contas'])-21139.82)<0.02
assert abs(sum(sum(c['v']) for c in tri['C1']['contas'])-19898.07)<0.02
# ------------- estrutura
est=b[(b['T']=='Despesa')&(b.Estrutura=='SIM')]
est_g={gg:mv(est[est.Grupo==gg]) for gg in est.Grupo.unique()}
est_tot=mv(est)
chk=[68434.82,68936.24,62243.14,64590.13,62254.37,74608.27,88590.18,94179.76,89023.77]
assert all(abs(est_tot[i]-chk[i])<0.02 for i in range(9)), est_tot
# ------------- unidades
uni={u:mv(b[(b.Grupo=='Receita Operacional Bruta')&(b['T']=='Receita')&(b.Unidade==u)]) for u in ['Rio Branco','Cruzeiro do Sul','Epitaciolândia']}
# ------------- parceria
import openpyxl
wbv=openpyxl.load_workbook(SP+'/test/v3_calc.xlsx',data_only=True); t=wbv['03 DRE Tricologia']
P={}
for r in range(9,44):
    lab=t.cell(r,1).value
    if lab and t.cell(r,2).value is not None:
        P[str(r)]=[lab]+[(round(t.cell(r,c).value,2) if isinstance(t.cell(r,c).value,(int,float)) else t.cell(r,c).value) for c in range(2,6)]
assert abs(P['35'][4]-109897.12)<0.01, P['35']
data=dict(meses=MESES,dre=L,fora=fora,razao=raz,est_g=est_g,est_tot=est_tot,uni=uni,uni_c=uni_c,tri=tri,
          parc=dict(dre=P,cash=old['cash'],rev=old['rev']),
          checks=dict(ebitda=ebitda,rob=rob))
json.dump(data,open(OUT,'w'),ensure_ascii=False,separators=(',',':'))
print('ok', len(json.dumps(data,ensure_ascii=False))//1024,'KB')
print('rob',rob); print('ebitda',ebitda); print('ll',ll)
print('est %',[round(est_tot[i]/rob[i]*100,2) for i in range(9)], round(sum(est_tot[6:])/sum(rob[6:])*100,4))
print('est grupos',{k:round(sum(v[6:])/3,2) for k,v in est_g.items()})
