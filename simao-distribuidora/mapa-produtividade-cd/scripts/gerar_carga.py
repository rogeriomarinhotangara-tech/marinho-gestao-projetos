"""Gera os documentos da carga inicial do Mapa de Produtividade — CD Simão.

Lê a aba BASE da planilha de análise (janeiro a junho) e as extrações mensais do
Kaizen (julho, agosto e setembro), confere os totais com a tabela de fechamento e
grava um JSON por documento em carga-inicial/<colecao>/<id>.json, além de
carga-inicial/seed.json (todos os documentos, usado no teste local).

Uso:
  python3 scripts/gerar_carga.py <planilha_analise.xlsx> <julho.xlsx> <agosto.xlsx> <setembro.xlsx>
"""
import json
import os
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, "carga-inicial")
MESES = ["janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho", "agosto",
         "setembro", "outubro", "novembro", "dezembro"]
NOME_MES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto",
            "Setembro", "Outubro", "Novembro", "Dezembro"]
FUNCOES = ["recebimento", "separacao", "expedicao"]
NOME_FUNCAO = {"recebimento": "Recebimento", "separacao": "Separação", "expedicao": "Conferência de Expedição"}
ANO = 2026
AGORA = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")

DIAS_OPERACAO = {1: 26, 2: 24, 3: 26, 4: 26, 5: 26, 6: 26, 7: 27, 8: 26, 9: 26}
ESPERADO = {
    (1, "recebimento"): (3, 31, 57, 565, 144737), (1, "separacao"): (9, 151, 1034, 9536, 267763), (1, "expedicao"): (5, 64, 1023, 9531, 267523),
    (2, "recebimento"): (4, 37, 121, 1192, 240605), (2, "separacao"): (8, 119, 902, 7480, 218294), (2, "expedicao"): (4, 65, 891, 7372, 214725),
    (3, "recebimento"): (6, 48, 183, 2354, 350112), (3, "separacao"): (10, 142, 1121, 12857, 335300), (3, "expedicao"): (6, 73, 1110, 12724, 332281),
    (4, "recebimento"): (2, 40, 138, 1643, 269230), (4, "separacao"): (12, 148, 1428, 11984, 295669), (4, "expedicao"): (4, 70, 1422, 12021, 297414),
    (5, "recebimento"): (2, 40, 142, 2402, 359063), (5, "separacao"): (10, 173, 1741, 14157, 334661), (5, "expedicao"): (3, 71, 1728, 14025, 331892),
    (6, "recebimento"): (2, 41, 146, 2656, 403011), (6, "separacao"): (12, 178, 1853, 16616, 401894), (6, "expedicao"): (4, 69, 1827, 16716, 402802),
    (7, "recebimento"): (2, 49, 147, 2935, 455497), (7, "separacao"): (9, 164, 1602, 11803, 263452), (7, "expedicao"): (3, 70, 1598, 11647, 260740),
    (8, "recebimento"): (3, 41, 107, 1500, 309729), (8, "separacao"): (10, 147, 1547, 12084, 327862), (8, "expedicao"): (4, 68, 1543, 12174, 329454),
    (9, "recebimento"): (3, 48, 177, 2623, 455601), (9, "separacao"): (13, 188, 2046, 16876, 411397), (9, "expedicao"): (6, 85, 2043, 16795, 410325),
}


def norm(s):
    s = unicodedata.normalize("NFD", str(s).replace("³", "3")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def funcao(texto):
    t = norm(texto)
    if "receb" in t:
        return "recebimento"
    if "separa" in t:
        return "separacao"
    if "exped" in t:
        return "expedicao"
    raise ValueError(f"função desconhecida: {texto!r}")


def operador(texto):
    s = re.sub(r"\s+", " ", str(texto)).strip()
    i = s.index("-")
    return s[:i].strip(), s[i + 1:].strip()


def dias(texto):
    if texto is None:
        return []
    if isinstance(texto, (int, float)):
        return [int(texto)]
    return [int(x) for x in re.findall(r"\d+", str(texto))]


def numero_sufixo(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"(kg|m³|m3)$", "", str(v).strip(), flags=re.I)
    return float(s)


PREP = {"da", "de", "do", "das", "dos", "e"}
ACENTOS = {"joao": "João", "deoclecio": "Deoclécio"}


def cap(p):
    k = norm(p)
    if k in ACENTOS:
        return ACENTOS[k]
    low = p.lower()
    return low[:1].upper() + low[1:]


def nome_curto(nome):
    partes = [p for p in nome.split() if norm(p) not in PREP]
    if len(partes) <= 1:
        return cap(partes[0]) if partes else ""
    return f"{cap(partes[0])} {cap(partes[-1])}"


def ler_base(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb["BASE"]
    cab = [norm(c.value) if c.value else "" for c in ws[3]]
    idx = {h: i for i, h in enumerate(cab)}
    out = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row[idx["operador"]] or not row[idx["atividade"]]:
            continue
        mes = MESES.index(norm(row[idx["mes"]])) + 1
        if mes > 6:
            continue  # julho vem do arquivo de julho
        cod, nome = operador(row[idx["operador"]])
        out.append(dict(mes=mes, funcao=funcao(row[idx["atividade"]]), codigo=cod, nomeKaizen=nome,
                        diasLista=dias(row[idx["dias efetivos de trabalho no mes"]]), diasFuncao=int(row[idx["total de dias trabalhados"]]),
                        pedidos=row[idx["qtde pedido om"]], nfs=row[idx["qt conf recebimento nf"]], skus=row[idx["qtde skus"]],
                        mediaKaizen=row[idx["media skus dia informada kaizen"]], unidades=row[idx["qtde unidades"]],
                        pesoBruto=None, volumeBruto=None,
                        arquivo="2026_ANALISE_DE_DESEMPENHO_CD_SIMÃO.xlsx (aba BASE)"))
    cadastro = []
    wp = wb["PARAMETROS"]
    for r in range(34, 59):
        v = wp.cell(r, 1).value
        if v:
            cadastro.append(operador(v))
    return out, cadastro


def ler_kaizen(caminho, nome_arquivo):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb.worksheets[0]
    cab = [norm(c.value) if c.value is not None else "" for c in ws[1]]
    idx = {h: i for i, h in enumerate(cab)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        op, at = row[idx["operador"]], row[idx["atividade"]]
        if op in (None, 0, "0", "") or at in (None, 0, "0", ""):
            continue
        cod, nome = operador(op)
        out.append(dict(mes=MESES.index(norm(row[idx["mes"]])) + 1, funcao=funcao(at), codigo=cod, nomeKaizen=nome,
                        diasLista=dias(row[idx["dias"]]), diasFuncao=int(row[idx["qtde dias"]]),
                        pedidos=row[idx["qtde pedido"]], nfs=row[idx["qt conf recebimento"]], skus=row[idx["qtde skus"]],
                        mediaKaizen=row[idx["media skus dia"]], unidades=row[idx["qtde und"]],
                        pesoBruto=numero_sufixo(row[idx["peso kg"]]), volumeBruto=numero_sufixo(row[idx["volume m3"]]),
                        arquivo=nome_arquivo))
    return out


def inteiro_se_possivel(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def main():
    analise, julho, agosto, setembro = sys.argv[1:5]
    base, cadastro = ler_base(analise)
    lancs = base + ler_kaizen(julho, "PRODUTIVIDADE_JULHO.xlsx") + ler_kaizen(agosto, "PRODUTIVIDADE_AGOSTO.xlsx") + \
        ler_kaizen(setembro, "Produtividade_Setembro_Logistica.xlsx")
    assert len(lancs) == 159, len(lancs)
    # conferência com a tabela de fechamento
    for (m, f), esp in ESPERADO.items():
        rs = [l for l in lancs if l["mes"] == m and l["funcao"] == f]
        calc = (len(rs), sum(l["diasFuncao"] for l in rs), sum(l["pedidos"] + l["nfs"] for l in rs),
                sum(l["skus"] for l in rs), sum(l["unidades"] for l in rs))
        assert calc[:4] == esp[:4] and abs(calc[4] - esp[4]) <= 0.5, (m, f, calc, esp)
    for m, d in DIAS_OPERACAO.items():
        dd = set()
        for l in lancs:
            if l["mes"] == m:
                dd |= set(l["diasLista"])
        assert len(dd) == d, (m, len(dd), d)

    docs = defaultdict(dict)
    # duplicidade A1
    suspeitos = set()
    for m in range(1, 10):
        rs = [l for l in lancs if l["mes"] == m]
        for i, a in enumerate(rs):
            for b in rs[i + 1:]:
                if a["codigo"] == b["codigo"] and a["funcao"] != b["funcao"] and all(a[k] == b[k] for k in ("diasLista", "diasFuncao", "pedidos", "nfs", "skus", "unidades")):
                    suspeitos |= {(m, a["codigo"], a["funcao"]), (m, b["codigo"], b["funcao"])}
    for l in lancs:
        lid = f"{ANO}-{l['mes']:02d}-{l['codigo']}-{l['funcao']}"
        susp = (l["mes"], l["codigo"], l["funcao"]) in suspeitos
        docs["lancamentos"][lid] = {
            "ano": ANO, "mes": l["mes"], "funcao": l["funcao"], "codigo": l["codigo"], "nomeKaizen": l["nomeKaizen"],
            "diasLista": l["diasLista"], "diasFuncao": l["diasFuncao"],
            "pedidos": inteiro_se_possivel(l["pedidos"]), "nfs": inteiro_se_possivel(l["nfs"]), "skus": inteiro_se_possivel(l["skus"]),
            "unidades": inteiro_se_possivel(l["unidades"]), "mediaKaizen": inteiro_se_possivel(l["mediaKaizen"]),
            "pesoBruto": inteiro_se_possivel(l["pesoBruto"]), "volumeBruto": l["volumeBruto"],
            "divergencias": None, "destino": None, "observacao": "",
            "origem": "carga inicial", "arquivo": l["arquivo"], "importadoEm": AGORA, "editadoEm": None,
            "suspeito": susp,
            "motivoSuspeita": ("Achado A1: registro idêntico em Separação e em Conferência de Expedição no mesmo mês "
                               f"({l['diasFuncao']} dias, {l['pedidos']} pedidos, {l['skus']} itens, "
                               + f"{int(l['unidades']):,}".replace(",", ".")
                               + " unidades, mesma lista de dias). Confirmar com o suporte Kaizen e excluir um dos dois.") if susp else None,
        }

    # colaboradores: cadastro da planilha + novos de setembro
    nomes = {c: n for c, n in cadastro}
    primeiro = {}
    for l in lancs:
        nomes.setdefault(l["codigo"], l["nomeKaizen"])
        k = (l["mes"], min(l["diasLista"]) if l["diasLista"] else 1)
        primeiro[l["codigo"]] = min(primeiro.get(l["codigo"], k), k)
    dias_f = defaultdict(lambda: defaultdict(int))
    for l in lancs:
        dias_f[l["codigo"]][l["funcao"]] += l["diasFuncao"]
    novos_set = {c for c in nomes if c not in dict(cadastro)}
    for cod, nome in nomes.items():
        d = dias_f[cod]
        sugerida = max(FUNCOES, key=lambda f: (d.get(f, 0), -FUNCOES.index(f))) if d else None
        multi = len([f for f in FUNCOES if d.get(f)]) > 1
        dist = ", ".join(f"{NOME_FUNCAO[f]} {d[f]}" for f in FUNCOES if d.get(f))
        revisar, motivo, rotas = False, None, False
        if multi:
            revisar = True
            motivo = (f"Atua em {len([f for f in FUNCOES if d.get(f)])} funções (dias-função jan–set/2026: {dist}). "
                      "Função principal sugerida pelo maior número de dias-função; confirme.")
        if cod == "625":
            revisar = True
            motivo = f"Indicado para revisão pelo consultor; nos dados de jan–set/2026 aparece só em Separação ({dist} dias-função)."
        if cod in novos_set:
            m, dia = primeiro[cod]
            revisar = True
            rotas = True
            motivo = (f"Novo em {NOME_MES[m - 1].lower()}/{ANO} (primeiro registro em {dia:02d}/{m:02d}): confirme a função principal "
                      "e se atuava nas rotas externas (marcado como sugestão, pela entrada a partir de 22/09).")
        docs["colaboradores"][cod] = {
            "codigo": cod, "nome": nome, "nomeCurto": nome_curto(nome), "funcaoPrincipal": sugerida, "funcaoPrincipalSugerida": sugerida,
            "status": "ativo", "dataDesligamento": None, "observacao": "", "revisar": revisar, "motivoRevisao": motivo,
            "rotasExternas": rotas, "criadoEm": AGORA, "atualizadoEm": AGORA,
        }

    for m in range(1, 10):
        escopo = "Lojas Simão" if m < 9 else "Lojas Simão + rotas externas a partir de 22/09"
        obs = "Dias de operação apurados na planilha de análise (dias distintos com movimento)." if m <= 7 else \
            "Dias de operação calculados na carga inicial (dias distintos com movimento)."
        if m == 9:
            obs += " Entrada das rotas externas no coletor a partir de 22/09 (data a confirmar em Parâmetros)."
        docs["meses"][f"{ANO}-{m:02d}"] = {
            "ano": ANO, "mes": m, "diasOperacao": DIAS_OPERACAO[m], "diasOperacaoCalculado": DIAS_OPERACAO[m],
            "escopoColeta": escopo, "dataInicioRotasExternas": "2026-09-22" if m == 9 else None, "quebraSerie": m == 9,
            "situacao": "aberto", "observacao": obs, "atualizadoEm": AGORA,
        }

    docs["parametros"]["geral"] = {
        "minimoDias": 5, "pesoItens": 0.6, "pesoQuantidade": 0.4, "tolerancia": 0.85, "limiteConcentracao": 0.30, "limiteQueda": 0.20,
        "fatorMediana": 2, "mesesNovato": 3, "mesesSemMovimento": 2, "minimoDiasMeta": 10,
        "inicioRotasExternas": "2026-09-22", "inicioRotasConfirmado": False, "atualizadoEm": AGORA,
    }
    origem = ("Derivada do 3º quartil (P75) dos registros de jan–jul/2026 com 10 ou mais dias na função "
              "(planilha 2026_ANALISE_DE_DESEMPENHO_CD_SIMÃO). Substituir pela meta oficial quando definida.")
    for f, (it, un, dc) in {"recebimento": (60, 9000, 3.7), "separacao": (80, 2400, 12), "expedicao": (240, 6600, 24)}.items():
        docs["metas"][f"{f}-2026-01"] = {"funcao": f, "itensDia": it, "unidadesDia": un, "documentosDia": dc, "origem": origem,
                                         "vigenciaDesde": "2026-01", "criadoEm": AGORA, "atualizadoEm": AGORA}

    def totais(rs):
        t = {}
        for f in FUNCOES:
            x = [l for l in rs if l["funcao"] == f]
            t[f] = {"linhas": len(x), "dias": sum(l["diasFuncao"] for l in x), "documentos": sum(l["pedidos"] + l["nfs"] for l in x),
                    "itens": sum(l["skus"] for l in x), "unidades": round(sum(l["unidades"] for l in x), 3)}
        return t
    docs["importacoes"]["i-carga-base-2026"] = {
        "arquivo": "2026_ANALISE_DE_DESEMPENHO_CD_SIMÃO.xlsx (aba BASE)", "ano": ANO, "mes": None, "periodo": "Janeiro a junho/2026",
        "linhas": len(base), "modo": "carga inicial", "data": AGORA, "totais": totais(base), "novosColaboradores": [],
        "observacao": "Janeiro a junho vindos da aba BASE. As 14 linhas de julho da BASE são idênticas ao arquivo de julho (conferido campo a campo) e foram carregadas a partir do arquivo.",
        "autorId": None,
    }
    for m, arq in ((7, "PRODUTIVIDADE_JULHO.xlsx"), (8, "PRODUTIVIDADE_AGOSTO.xlsx"), (9, "Produtividade_Setembro_Logistica.xlsx")):
        rs = [l for l in lancs if l["mes"] == m]
        docs["importacoes"][f"i-carga-{ANO}-{m:02d}"] = {
            "arquivo": arq, "ano": ANO, "mes": m, "linhas": len(rs), "lidas": len(rs) + 1, "ignoradas": 1, "modo": "carga inicial", "data": AGORA,
            "totais": totais(rs), "novosColaboradores": sorted({l["codigo"] for l in rs} & novos_set, key=int),
            "diasOperacao": DIAS_OPERACAO[m], "autorId": None,
            "observacao": "Linha final de totais ignorada (sem operador e sem atividade)." + (" Idêntico às 14 linhas de julho da aba BASE." if m == 7 else ""),
        }
    docs["historico"]["h-carga-inicial"] = {
        "quando": AGORA, "autorId": None, "tipo": "carga inicial", "colecao": None, "alvo": None, "alvos": [], "mapa": False,
        "resumo": f"Carga inicial: {len(docs['lancamentos'])} lançamentos (jan–set/{ANO}), {len(docs['colaboradores'])} colaboradores, "
                  f"{len(docs['meses'])} meses, 3 metas e parâmetros gerais.",
        "antes": None, "depois": {k: len(v) for k, v in docs.items()},
    }

    os.makedirs(SAIDA, exist_ok=True)
    total = 0
    for col, ds in docs.items():
        os.makedirs(os.path.join(SAIDA, col), exist_ok=True)
        for i, d in ds.items():
            with open(os.path.join(SAIDA, col, f"{i}.json"), "w", encoding="utf-8") as fh:
                json.dump(d, fh, ensure_ascii=False, indent=1)
            total += 1
    with open(os.path.join(SAIDA, "seed.json"), "w", encoding="utf-8") as fh:
        json.dump(docs, fh, ensure_ascii=False)
    print(f"{total} documentos gravados em {SAIDA}")
    for col, ds in docs.items():
        print(f"  {col}: {len(ds)}")
    print("Colaboradores para revisar:")
    for c, d in sorted(docs["colaboradores"].items(), key=lambda x: int(x[0])):
        if d["revisar"]:
            print(f"  {c} {d['nomeCurto']}: {NOME_FUNCAO.get(d['funcaoPrincipal'])} — {d['motivoRevisao']}")


if __name__ == "__main__":
    main()
